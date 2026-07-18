from io import BytesIO

from docx import Document

from app.modules.replace.csv_parser import parse_replace_csv
from app.modules.replace.engine_docx import replace_docx


def _docx_bytes() -> bytes:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Привет ")
    p.add_run("мир")
    p.add_run(" и мирный")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "hello world"
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def test_parse_csv_supports_comments_bom() -> None:
    payload = "\ufeff#comment\nfoo;bar\n\nhello;world\n".encode("utf-8")
    rules = parse_replace_csv(payload)
    assert [r["from"] for r in rules] == ["foo", "hello"]


def test_parse_csv_preserves_quoted_semicolons() -> None:
    payload = '"foo;bar";baz\n'.encode("utf-8")
    rules = parse_replace_csv(payload)
    assert rules == [{"from": "foo;bar", "to": "baz", "flags": {}, "priority": 0}]


def test_replace_docx_handles_broken_runs_and_whole_word() -> None:
    data = _docx_bytes()
    replaced, hits = replace_docx(
        data,
        [{"from": "мир", "to": "world"}, {"from": "hello", "to": "hi"}],
        case_sensitive=False,
        whole_word=True,
    )
    text = "\n".join(p.text for p in Document(BytesIO(replaced)).paragraphs)
    assert "Привет world и мирный" in text
    assert any(h.rule_from == "мир" for h in hits)
