from pathlib import Path

from docx import Document

from backend.app.modules.files.extractors import extract_text_docx, normalize_text


def test_normalize_text_removes_extra_spaces_and_nuls() -> None:
    assert normalize_text("a\x00  b\n\n c") == "a b c"


def test_extract_text_docx_reads_body_tables_headers_footers(tmp_path: Path) -> None:
    file_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("body text")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "table text"
    section = doc.sections[0]
    section.header.paragraphs[0].text = "header text"
    section.footer.paragraphs[0].text = "footer text"
    doc.save(file_path)

    text = extract_text_docx(file_path)

    assert "body text" in text
    assert "table text" in text
    assert "header text" in text
    assert "footer text" in text
