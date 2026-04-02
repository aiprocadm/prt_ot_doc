from __future__ import annotations

from io import BytesIO

from docx import Document

from app.modules.templates.linter import lint_template
from app.modules.templates.passport import inject_passport
from app.modules.templates.render import render_docx


def _docx(text: str) -> bytes:
    doc = Document()
    p = doc.add_paragraph()
    # split placeholder across runs intentionally
    for part in text.split("|"):
        p.add_run(part)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_linter_extracts_placeholders_and_blocks() -> None:
    payload = _docx("{{ name }} {% if user.active %}{{ user.name }}{% endif %}")
    index = lint_template(payload)
    assert "name" in index["placeholders"]
    assert "user.active" in index["field_paths"]
    assert index["errors"] == []


def test_render_handles_split_runs_and_injects_passport() -> None:
    payload = _docx("{{ na|me| }}")
    rendered = render_docx(
        payload,
        {"name": "ALICE"},
        {"template_code": "contract", "template_version": 1},
    )
    output = Document(BytesIO(rendered))
    merged = "\n".join(p.text for p in output.paragraphs)
    assert "ALICE" in merged
    assert "PTD-PASSPORT" in merged


def test_passport_hidden_marker_in_document_xml() -> None:
    payload = _docx("Hello")
    out = inject_passport(payload, {"template_code": "x"}, visible=False)
    assert b"PTD-PASSPORT:" in out
    assert b"w:vanish" in out
