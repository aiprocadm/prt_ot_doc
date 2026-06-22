"""Ф4 чистый сборщик DOCX наряда-допуска: открыть результат python-docx и проверить содержимое."""

from __future__ import annotations

from io import BytesIO

from docx import Document

from app.domains.work_permits.print_form import (
    SignatureLine,
    StructuredSection,
    StructuredTable,
    WorkPermitPrintData,
    build_work_permit_docx,
)


def _all_text(doc_bytes: bytes) -> str:
    doc = Document(BytesIO(doc_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _sample() -> WorkPermitPrintData:
    return WorkPermitPrintData(
        number="НД-7",
        work_type_label="Работа на высоте",
        legal_reference="Приказ Минтруда России от 16.11.2020 № 782н",
        status_label="Выдан",
        org_header="ООО Ромашка",
        subdivision="Цех №1",
        planned_start="2026-06-20 08:00",
        planned_end="2026-06-20 18:00",
        zone_text="фасад корпуса А",
        content_text="монтаж",
        conditions_text="ясно",
        equipment_text="люлька",
        hazards_text="высота 12 м",
        structured_section=StructuredSection(
            title="Системы обеспечения безопасности (782н)",
            kv=[("Системы", "Удерживающие, Страховочные")],
            table=None,
        ),
        measures_before="ограждение",
        measures_during="наблюдение",
        special_conditions="—",
        ppe_text="каска, строп",
        members=[("Производитель работ", "Иванов И.И."), ("Допускающий", "Петров П.П.")],
        briefing={
            "conducted_by_fio": "Иванов И.И.",
            "conducted_at": "2026-06-20 07:45",
            "topics": "ТБ на высоте",
        },
        daily_admissions=[
            {
                "date": "2026-06-20",
                "start": "08:00",
                "end": "18:00",
                "admitted_by_fio": "Петров П.П.",
            }
        ],
        extensions=[
            {"old_end": "2026-06-20 18:00", "new_end": "2026-06-21 18:00", "at": "2026-06-20 17:30"}
        ],
        completion={"text": "работы окончены, место сдано", "recorded_at": "2026-06-21 17:00"},
        closed_at="2026-06-21 17:05",
        signatures=[
            SignatureLine(
                group="permit",
                role_label="Производитель работ",
                fio="Иванов И.И.",
                status_label="Подписано",
                signed_at="2026-06-20 08:05",
                mode="attested",
                hash_short="a1b2c3d4e5f6a7b8",
            ),
            SignatureLine(
                group="closing",
                role_label="Сдал",
                fio="Иванов И.И.",
                status_label="Подписано",
                signed_at="2026-06-21 17:00",
                mode="attested",
                hash_short="ffeeddccbbaa9988",
            ),
        ],
    )


def test_build_docx_contains_core_fields():
    text = _all_text(build_work_permit_docx(_sample()))
    assert "НД-7" in text
    assert "Иванов И.И." in text
    assert "№ 782н" in text
    assert "Удерживающие, Страховочные" in text
    assert "работы окончены, место сдано" in text
    assert "a1b2c3d4e5f6a7b8" in text
    assert "Сдал" in text


def test_build_docx_renders_structured_table():
    data = _sample()
    data.structured_section = StructuredSection(
        title="Анализ воздушной среды (902н)",
        kv=[("Вентиляция", "Принудительная")],
        table=StructuredTable(headers=["Параметр", "Значение"], rows=[["Кислород", "20.9"]]),
    )
    text = _all_text(build_work_permit_docx(data))
    assert "Анализ воздушной среды (902н)" in text
    assert "Вентиляция" in text and "Принудительная" in text
    assert "Кислород" in text and "20.9" in text


def test_build_docx_empty_sections_do_not_crash():
    data = WorkPermitPrintData(
        number="НД-8",
        work_type_label="Работа на высоте",
        legal_reference="Приказ Минтруда России от 16.11.2020 № 782н",
        status_label="Черновик",
        org_header="ООО Ромашка",
        subdivision=None,
        planned_start=None,
        planned_end=None,
        zone_text="зона",
        content_text=None,
        conditions_text=None,
        equipment_text=None,
        hazards_text=None,
        structured_section=None,
        measures_before=None,
        measures_during=None,
        special_conditions=None,
        ppe_text=None,
        members=[],
        briefing=None,
        daily_admissions=[],
        extensions=[],
        completion=None,
        closed_at=None,
        signatures=[],
    )
    out = build_work_permit_docx(data)
    assert isinstance(out, bytes) and len(out) > 0


def test_label_dicts_cover_lifecycle_vocab():
    from app.domains.work_permits import lifecycle as lc
    from app.domains.work_permits import print_form as pf

    assert set(pf.WORK_TYPE_LABELS) >= lc.WORK_TYPES
    assert set(pf.MEMBER_ROLE_LABELS) >= lc.MEMBER_ROLES
    assert set(pf.SAFETY_SYSTEM_LABELS) >= lc.SAFETY_SYSTEMS
    assert set(pf.STATUS_LABELS) >= lc.WORK_PERMIT_STATUSES
