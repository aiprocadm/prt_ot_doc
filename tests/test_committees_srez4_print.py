"""Committees срез-4 — печатная форма протокола: сервис-загрузчик + роут."""

from __future__ import annotations

from datetime import date, datetime, timezone
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document
from fastapi import HTTPException

from app.api.routes import committees as routes
from app.models.committees import (
    Committee,
    CommitteeAgendaItem,
    CommitteeDecision,
    CommitteeDecisionTask,
    CommitteeDecisionVote,
    CommitteeKind,
    CommitteeMeeting,
    CommitteeMeetingAttendance,
    CommitteeMeetingInvitation,
    DecisionTaskStatus,
    MeetingStatus,
    VoteChoice,
)
from app.models.master_data import Person
from app.services.committee_protocol_print import (
    PdfRendererUnavailable,
    render_committee_protocol,
)

_NOW = datetime(2026, 8, 3, 10, 0, tzinfo=timezone.utc)


def _tenant(tid="tenant-1"):
    return SimpleNamespace(id=tid, is_active=True, slug="t1", code="t1")


@pytest.fixture(autouse=True)
def _flag_on(monkeypatch):
    monkeypatch.setattr(routes, "is_module_enabled", AsyncMock(return_value=True))


def _all_text(docx_bytes: bytes) -> str:
    doc = Document(BytesIO(docx_bytes))
    parts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


async def _seed_full_meeting(session, tenant_id="tenant-1"):
    c = Committee(
        tenant_id=tenant_id,
        kind=CommitteeKind.OSMS,
        name="Комитет ОТ",
        is_active=True,
        quorum_threshold_pct=66,
    )
    session.add(c)
    await session.flush()
    m = CommitteeMeeting(
        tenant_id=tenant_id,
        committee_id=c.id,
        scheduled_at=_NOW,
        location="Зал 1",
        status=MeetingStatus.HELD,
        held_at=_NOW,
        protocol_seq=3,
        protocol_year=2026,
        members_total=4,
        present_count=3,
        quorum_met=True,
    )
    session.add(m)
    await session.flush()
    # ФИО есть только у p1; p2 без Person-строки → в печать уходит id.
    session.add(
        Person(
            id="p1",
            tenant_id=tenant_id,
            company_id="comp-1",
            first_name="Иван",
            last_name="Иванов",
            middle_name="Иванович",
        )
    )
    session.add(
        CommitteeMeetingAttendance(
            tenant_id=tenant_id, meeting_id=m.id, person_id="p1", present=True
        )
    )
    session.add(
        CommitteeMeetingAttendance(
            tenant_id=tenant_id, meeting_id=m.id, person_id="p2", present=False
        )
    )
    session.add(
        CommitteeMeetingInvitation(
            tenant_id=tenant_id, meeting_id=m.id, person_id="p1", invited_at=_NOW
        )
    )
    a = CommitteeAgendaItem(tenant_id=tenant_id, meeting_id=m.id, seq=1, title="О травматизме")
    session.add(a)
    await session.flush()
    d = CommitteeDecision(
        tenant_id=tenant_id, meeting_id=m.id, agenda_item_id=a.id, text="Усилить контроль"
    )
    session.add(d)
    await session.flush()
    session.add(
        CommitteeDecisionVote(
            tenant_id=tenant_id, decision_id=d.id, person_id="p1", choice=VoteChoice.FOR
        )
    )
    session.add(
        CommitteeDecisionTask(
            tenant_id=tenant_id,
            decision_id=d.id,
            assignee_person_id="p1",
            due_date=date(2026, 9, 1),
            status=DecisionTaskStatus.OPEN,
        )
    )
    await session.flush()
    return c, m


@pytest.mark.asyncio
async def test_render_docx_full_protocol(sessionmaker):
    async with sessionmaker() as session:
        _, m = await _seed_full_meeting(session)
        await session.commit()
        rendered = await render_committee_protocol(
            session, tenant_id="tenant-1", meeting=m, fmt="docx"
        )
    assert rendered.media_type.endswith("wordprocessingml.document")
    assert rendered.filename == "committee-protocol-3-2026.docx"
    text = _all_text(rendered.content)
    assert "ПРОТОКОЛ № 3/2026" in text
    assert "Комитет ОТ" in text
    assert "Иванов Иван Иванович" in text  # ФИО из Person
    assert "p2" in text  # фолбэк на id без Person-строки
    assert "порог 66%" in text
    assert "Усилить контроль" in text
    assert "За: 1" in text
    assert "2026-09-01" in text


@pytest.mark.asyncio
async def test_route_planned_meeting_409(sessionmaker):
    async with sessionmaker() as session:
        c = Committee(tenant_id="tenant-1", kind=CommitteeKind.OSMS, name="К", is_active=True)
        session.add(c)
        await session.flush()
        m = CommitteeMeeting(
            tenant_id="tenant-1",
            committee_id=c.id,
            scheduled_at=_NOW,
            status=MeetingStatus.PLANNED,
        )
        session.add(m)
        await session.flush()
        with pytest.raises(HTTPException) as exc:
            await routes.print_protocol(
                mid=m.id,
                tenant=_tenant(),
                session=session,
                access=SimpleNamespace(),
                fmt="docx",
            )
        assert exc.value.status_code == 409


@pytest.mark.asyncio
async def test_route_pdf_unavailable_503(sessionmaker, monkeypatch):
    async with sessionmaker() as session:
        _, m = await _seed_full_meeting(session)
        await session.commit()
        monkeypatch.setattr(
            routes,
            "render_committee_protocol",
            AsyncMock(side_effect=PdfRendererUnavailable("no soffice")),
        )
        with pytest.raises(HTTPException) as exc:
            await routes.print_protocol(
                mid=m.id,
                tenant=_tenant(),
                session=session,
                access=SimpleNamespace(),
                fmt="pdf",
            )
        assert exc.value.status_code == 503


@pytest.mark.asyncio
async def test_route_docx_response_headers(sessionmaker):
    async with sessionmaker() as session:
        _, m = await _seed_full_meeting(session)
        await session.commit()
        resp = await routes.print_protocol(
            mid=m.id,
            tenant=_tenant(),
            session=session,
            access=SimpleNamespace(),
            fmt="docx",
        )
    assert resp.headers["content-disposition"].startswith("attachment;")
    assert "committee-protocol-3-2026.docx" in resp.headers["content-disposition"]
