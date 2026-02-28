from __future__ import annotations

from pathlib import Path

import pytest
from sqlalchemy import select

from app.models.models import Tenant
from app.modules.files.extractors import extract_text_docx
from app.modules.files.models import FileContentIndex, FileContentIndexStatus, FileRecord, FileStatus
from app.modules.files.service import index_file_record
from app.modules.search.service import SearchService


def test_extract_text_docx_includes_paragraphs_and_tables(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Привет мир")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Ячейка 1"
    table.rows[0].cells[1].text = "Ячейка 2"
    doc.save(path)

    text = extract_text_docx(path)

    assert "Привет мир" in text
    assert "Ячейка 1" in text
    assert "Ячейка 2" in text


@pytest.mark.anyio
async def test_index_file_record_idempotent_by_sha256(sessionmaker) -> None:
    async with sessionmaker() as session:
        tenant = (await session.execute(select(Tenant).where(Tenant.slug == "test"))).scalar_one()
        tenant_id = str(tenant.id)
        file_rec = FileRecord(
            tenant_id=tenant_id,
            bucket="main",
            object_key="uploads/a.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            size_bytes=16,
            sha256="abc123",
            status=FileStatus.clean.value,
            av_result_json={},
            metadata_json={},
        )
        session.add(file_rec)
        await session.flush()
        index = FileContentIndex(
            tenant_id=tenant_id,
            file_id=file_rec.id,
            content_sha256="abc123",
            mime_type=file_rec.content_type,
            status=FileContentIndexStatus.indexed.value,
            attempts=2,
        )
        session.add(index)
        await session.commit()

        await index_file_record(session, tenant_id=tenant_id, file_id=file_rec.id)
        await session.flush()

        refreshed = (await session.execute(select(FileContentIndex).where(FileContentIndex.file_id == file_rec.id))).scalar_one()
        assert refreshed.attempts == 2
        assert refreshed.status == FileContentIndexStatus.indexed.value


def test_search_builds_entity_urls() -> None:
    assert SearchService._build_entity_url("documents", "1") == "/documents/1"
    assert SearchService._build_entity_url("risk", "2") == "/risk/2"
    assert SearchService._build_entity_url("unknown", "3") == "/unknown/3"
