"""Сервис рендера печатных форм 29н: собирает данные из БД и отдаёт DOCX-байты."""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from app.models.models import MedicalFactor, Position, PositionHazardLink
from app.models.risk import RiskHazard
from app.services.medical_print import render_contingent_register, render_named_list


def _docx_text(b: bytes) -> str:
    doc = Document(BytesIO(b))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


async def _seed(session, data_factory):
    tenant = await data_factory.ensure_tenant(session=session)
    company = await data_factory.create_company(tenant=tenant, session=session)
    hazard = RiskHazard(
        tenant_id=tenant.id, code="noise", title="Шум", medical_factor_code="4.4"
    )
    session.add(hazard)
    await session.flush()
    pos = Position(tenant_id=tenant.id, company_id=company.id, name="Сварщик")
    session.add(pos)
    await session.flush()
    session.add(PositionHazardLink(tenant_id=tenant.id, position_id=pos.id, hazard_id=hazard.id))
    await data_factory.create_person(
        tenant=tenant,
        company=company,
        session=session,
        position_id=pos.id,
        first_name="Иван",
        last_name="Петров",
    )
    session.add(
        MedicalFactor(
            tenant_id=tenant.id,
            code="4.4",
            name="Шум",
            exam_kinds=["periodic"],
            periodicity_months=12,
        )
    )
    await session.commit()
    return tenant


_DOCX_MEDIA = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.mark.asyncio
async def test_render_register_docx_includes_position_and_factor(sessionmaker, data_factory):
    async with sessionmaker() as session:
        tenant = await _seed(session, data_factory)
        rendered = await render_contingent_register(session, tenant=tenant, fmt="docx")
    assert rendered.media_type == _DOCX_MEDIA
    assert rendered.filename.endswith(".docx")
    assert rendered.content[:2] == b"PK"  # DOCX = zip
    text = _docx_text(rendered.content)
    assert "контингент" in text.lower()
    assert "Сварщик" in text
    assert "Шум" in text
    assert "Периодический" in text  # exam_kind переведён в метку


@pytest.mark.asyncio
async def test_render_named_list_docx_includes_person_and_status(sessionmaker, data_factory):
    async with sessionmaker() as session:
        tenant = await _seed(session, data_factory)
        rendered = await render_named_list(session, tenant=tenant, fmt="docx")
    assert rendered.filename.endswith(".docx")
    text = _docx_text(rendered.content)
    assert "поименный список" in text.lower()
    assert "Петров Иван" in text
    assert "Отсутствует" in text  # status="missing" → метка, не сырой код
    assert "missing" not in text


@pytest.mark.asyncio
async def test_render_empty_tenant_does_not_crash(sessionmaker, data_factory):
    async with sessionmaker() as session:
        tenant = await data_factory.ensure_tenant(session=session)
        await session.commit()
        reg = await render_contingent_register(session, tenant=tenant, fmt="docx")
        named = await render_named_list(session, tenant=tenant, fmt="docx")
    assert reg.content[:2] == b"PK"
    assert named.content[:2] == b"PK"
