from __future__ import annotations

from collections import UserDict
from io import BytesIO
from pathlib import Path

from docx import Document

from app.domains.templating.renderer import TemplateRenderer, _coerce_mapping, render_docx, render_docx_with_metadata


def _build_template(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_coerce_mapping_handles_non_dict() -> None:
    assert _coerce_mapping({"a": 1}) == {"a": 1}
    assert _coerce_mapping(UserDict({"b": 2})) == {"b": 2}
    assert _coerce_mapping(None) == {}


def test_render_docx_applies_context() -> None:
    template_bytes = _build_template("Value: {{ value }}")
    rendered = render_docx(template_bytes, {"value": "42"})
    doc = Document(BytesIO(rendered))
    assert "Value: 42" in doc.paragraphs[0].text


def test_render_docx_records_reproducibility_metadata_and_fallback_warning() -> None:
    template_bytes = _build_template("Hello {{ name }} {{ missing_field }}")
    rendered = render_docx_with_metadata(template_bytes, {"name": "Bob"})
    doc = Document(BytesIO(rendered.content))
    assert doc.paragraphs[0].text == "Hello Bob "
    assert rendered.metadata["engine"] == "docxtpl"
    assert rendered.metadata["reproducibility"]["context_bytes"] > 0
    assert any("fallback" in warning for warning in rendered.warnings)


def test_template_renderer_writes_to_path(tmp_path: Path) -> None:
    template_path = tmp_path / "template.docx"
    template_path.write_bytes(_build_template("Hello {{ name }}"))
    renderer = TemplateRenderer(template_path=str(template_path))
    output_path = tmp_path / "out.docx"
    result_path = Path(renderer.render({"name": "Bob"}, str(output_path)))
    assert result_path == output_path
    assert output_path.exists()
    doc = Document(output_path)
    assert doc.paragraphs[0].text == "Hello Bob"


def test_render_docx_metadata_counts_nested_context_keys() -> None:
    template_bytes = _build_template("{{ company.name }}")
    rendered = render_docx_with_metadata(template_bytes, {"company": {"name": "ACME"}})
    assert rendered.metadata["strict_mode"] is True
    assert rendered.metadata["reproducibility"]["context_key_count"] >= 2
    assert "company.name" in rendered.metadata["context_keys"]
