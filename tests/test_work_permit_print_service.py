"""Ф4 сервис рендера наряда: собирает данные из БД и отдаёт DOCX-байты."""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from app.domains.work_permits import lifecycle as lc
from app.domains.work_permits import service as svc
from app.domains.work_permits.signing import sign_closing
from app.services.work_permit_print import render_work_permit


async def _persons(data_factory, *names):
    tenant = await data_factory.ensure_tenant()
    company = await data_factory.create_company(tenant=tenant)
    people = []
    for i, nm in enumerate(names):
        people.append(
            await data_factory.create_person(
                tenant=tenant, company=company, first_name=nm, last_name=f"P{i}"
            )
        )
    return tenant, people


def _docx_text(b: bytes) -> str:
    doc = Document(BytesIO(b))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


@pytest.mark.asyncio
async def test_render_docx_includes_members_and_signature(sessionmaker, data_factory):
    tenant, (foreman, supervisor) = await _persons(data_factory, "Ivan", "Petr")
    tid = str(tenant.id)
    async with sessionmaker() as session:
        wp = await svc.create_work_permit(
            session, tenant_id=tid, work_type="height", zone_text="фасад", number="НД-7"
        )
        await svc.add_member(
            session, tenant_id=tid, work_permit_id=wp.id, person_id=foreman.id, role="foreman"
        )
        await svc.add_member(
            session, tenant_id=tid, work_permit_id=wp.id, person_id=supervisor.id, role="supervisor"
        )
        wp.status = lc.STATUS_ISSUED
        await session.flush()
        await svc.record_completion(
            session,
            tenant_id=tid,
            work_permit_id=wp.id,
            completion_text="работы окончены",
            actor_user_id="u1",
        )
        await sign_closing(
            session,
            tenant_id=tid,
            work_permit_id=wp.id,
            person_id=foreman.id,
            mode="attested",
            requested_by="u1",
        )
        rendered = await render_work_permit(
            session, tenant=tenant, permit_id=wp.id, fmt="docx", with_letterhead=False
        )
        assert (
            rendered.media_type
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        assert rendered.filename.endswith(".docx")
        text = _docx_text(rendered.content)
        assert "НД-7" in text
        assert "Ivan" in text  # ФИО члена бригады зарезолвлено
        assert "работы окончены" in text
        assert "Подписано" in text  # статус закрывающей подписи в таблице
        assert "Сдал" in text  # роль закрытия (handover) производителя работ


@pytest.mark.asyncio
async def test_render_unknown_permit_returns_none(sessionmaker, data_factory):
    tenant, _ = await _persons(data_factory, "Solo")
    async with sessionmaker() as session:
        rendered = await render_work_permit(
            session, tenant=tenant, permit_id="missing-id", fmt="docx", with_letterhead=False
        )
        assert rendered is None


@pytest.mark.asyncio
async def test_render_confined_space_uses_902n_and_gas_table(sessionmaker, data_factory):
    tenant, (foreman,) = await _persons(data_factory, "Sidor")
    tid = str(tenant.id)
    async with sessionmaker() as session:
        wp = await svc.create_work_permit(
            session,
            tenant_id=tid,
            work_type="confined_space",
            zone_text="колодец К-12",
            number="НД-ОЗП-1",
            type_specific={
                "gas_analysis": [{"parameter": "oxygen", "value": "20.9", "norm": "≥ 20"}],
                "ventilation": "forced",
            },
        )
        await svc.add_member(
            session, tenant_id=tid, work_permit_id=wp.id, person_id=foreman.id, role="foreman"
        )
        rendered = await render_work_permit(
            session, tenant=tenant, permit_id=wp.id, fmt="docx", with_letterhead=False
        )
        text = _docx_text(rendered.content)
        assert "902н" in text
        assert "Принудительная" in text
        assert "Кислород" in text and "20.9" in text
        assert "782н" not in text


@pytest.mark.asyncio
async def test_render_hot_work_uses_1479_and_fire_means(sessionmaker, data_factory):
    tenant, (foreman,) = await _persons(data_factory, "Fedor")
    tid = str(tenant.id)
    async with sessionmaker() as session:
        wp = await svc.create_work_permit(
            session,
            tenant_id=tid,
            work_type="hot_work",
            zone_text="эстакада №3",
            number="НД-ОГН-1",
            type_specific={
                "fire_fighting_means": ["extinguisher_powder", "sand"],
                "gas_analysis": [{"parameter": "flammable", "value": "0", "norm": "≤ 10 % НКПР"}],
            },
        )
        await svc.add_member(
            session, tenant_id=tid, work_permit_id=wp.id, person_id=foreman.id, role="foreman"
        )
        rendered = await render_work_permit(
            session, tenant=tenant, permit_id=wp.id, fmt="docx", with_letterhead=False
        )
        text = _docx_text(rendered.content)
        assert "1479" in text
        assert "Огнетушитель порошковый" in text and "Ящик с песком" in text
        assert "Горючие" in text and "0" in text
        assert "782н" not in text
