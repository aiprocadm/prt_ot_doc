from __future__ import annotations

import base64
from io import BytesIO

from docx import Document
from docxtpl import DocxTemplate, InlineImage

from app.services.docx import DocxService


def _build_docx(text: str = "Hello {{ name }}!") -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_render_template_substitutes_values() -> None:
    template_bytes = _build_docx("Greetings {{ name }}")
    rendered = DocxService.render_template(template_bytes, {"name": "Alice"})
    output = Document(BytesIO(rendered))
    assert "Greetings Alice" in "\n".join(p.text for p in output.paragraphs)


def test_prepare_context_converts_inline_images() -> None:
    template_bytes = _build_docx("Image placeholder {{ image }}")
    tpl = DocxTemplate(BytesIO(template_bytes))
    descriptor = {
        "_type": "inline_image",
        "data": base64.b64encode(b"data").decode(),
        "width_mm": 30,
    }
    prepared = DocxService._prepare_context(tpl, {"image": descriptor})
    assert "image" in prepared
    inline = prepared["image"]
    assert isinstance(inline, InlineImage)
    assert inline.width.mm == 30


def test_mass_replace_updates_multiple_sections() -> None:
    original = _build_docx("{{greeting}} {{greeting}}")
    updated = DocxService.mass_replace(original, {"{{greeting}}": "Привет"})
    doc = Document(BytesIO(updated))
    assert doc.paragraphs[0].text == "Привет Привет"


def test_set_headers_footers_handles_multiple_sections() -> None:
    doc = Document()
    doc.add_paragraph("Body")
    doc.sections[0].header.add_paragraph("Old header")
    doc.sections[0].footer.add_paragraph("Old footer")
    doc.add_section()
    buffer = BytesIO()
    doc.save(buffer)
    result = DocxService.set_headers_footers(
        buffer.getvalue(),
        header_texts=["Top"],
        footer_texts=["Bottom", "Tail"],
    )
    updated = Document(BytesIO(result))
    headers = [
        paragraph.text for section in updated.sections for paragraph in section.header.paragraphs
    ]
    footers = [
        paragraph.text for section in updated.sections for paragraph in section.footer.paragraphs
    ]
    assert headers == ["Top", "Top"]
    assert footers == ["Bottom", "Tail"]


def test_extract_placeholders_collects_unique_tokens() -> None:
    doc_bytes = _build_docx("{{ first }} and {{ second }} and {{ first }}")
    placeholders = DocxService.extract_placeholders(doc_bytes)
    assert placeholders == {"first", "second"}
