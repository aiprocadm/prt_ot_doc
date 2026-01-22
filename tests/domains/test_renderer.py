from __future__ import annotations

from io import BytesIO

from docx import Document as DocxDocument

from app.domains.templating.renderer import render_docx


def _build_template_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("Hello {{ name }}!")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_render_docx_renders_template_context() -> None:
    template_bytes = _build_template_bytes()

    rendered = render_docx(template_bytes, {"name": "World"})

    result = DocxDocument(BytesIO(rendered))
    paragraphs = [paragraph.text for paragraph in result.paragraphs]

    assert "Hello World!" in paragraphs
    assert all("{{" not in paragraph for paragraph in paragraphs)
