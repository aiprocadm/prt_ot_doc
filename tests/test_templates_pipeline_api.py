from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from httpx import AsyncClient

from app.api.v1.router import MAX_METADATA_JSON_BYTES
from app.core.payload_constraints import MAX_STRING_VALUE_LENGTH
from app.domains.packs.definitions import (
    PACK_CODE_INCIDENT,
    PACK_CODE_INSPECTION_PREP,
    PACK_CODE_SITE_ACCESS,
)
from app.services.file_storage import FileStorageService

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _build_template() -> bytes:
    doc = Document()
    doc.add_paragraph("Hello {{ name }}!")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _build_alt_template() -> bytes:
    doc = Document()
    doc.add_paragraph("Hello {{ name }}!")
    doc.add_paragraph("Version two")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _version_metadata_form() -> dict[str, str]:
    return {
        "document_type": "safety_doc",
        "required_fields_schema": '{"type":"object","properties":{"name":{"type":"string"}}}',
        "applicability_rules": "{}",
        "output_types": '["docx","pdf"]',
        "profile": "{}",
    }


@pytest.mark.anyio
async def test_template_creation_and_pipeline_execution(
    async_client: AsyncClient, make_auth_headers
) -> None:
    auth_headers = await make_auth_headers()
    headers = {**dict(async_client.headers), **auth_headers}
    template_bytes = _build_template()

    response = await async_client.post(
        "/api/v1/templates",
        files={"file": ("greeting.docx", template_bytes, DOCX_CONTENT_TYPE)},
        data={
            "name": "Greeting",
            "description": "Simple",
            "metadata": "{}",
            **_version_metadata_form(),
        },
        headers=headers,
    )
    assert response.status_code == 201
    template_id = response.json()["id"]

    payload = {
        "context": {"name": "World"},
        "replacements": {"Hello": "Hi"},
        "header_text": "Header",
        "footer_text": "Footer",
        "output_basename": "greeting",
        "idempotency_key": "test-key",
    }

    run_response = await async_client.post(
        f"/api/v1/pipelines/{template_id}/run",
        json=payload,
        params={"mode": "sync"},
        headers=headers,
    )
    assert run_response.status_code == 200
    body = run_response.json()

    assert body["status"] == "done"
    assert body["context"] == payload["context"]
    assert body["docx_storage_key"]
    assert body["pdf_storage_key"]

    storage = FileStorageService.default()
    assert storage.has(body["docx_storage_key"])
    assert storage.has(body["pdf_storage_key"])
    metadata = body["result_metadata"]
    assert metadata["result"]["replacements_applied"] == 1
    assert isinstance(metadata["result"]["pdf_fallback"], bool)
    assert metadata["request"]["header_text"] == "Header"

    list_response = await async_client.get("/api/v1/templates", headers=headers)
    assert list_response.status_code == 200
    listing = list_response.json()
    assert listing["total"] >= 1
    assert any(item["id"] == template_id for item in listing["items"])


@pytest.mark.anyio
async def test_pack_scenarios_listing_and_creation(
    async_client: AsyncClient, make_auth_headers
) -> None:
    headers = {**dict(async_client.headers), **await make_auth_headers()}

    list_response = await async_client.get("/api/v1/packs/scenarios", headers=headers)
    assert list_response.status_code == 200
    payload = list_response.json()
    codes = {item["code"] for item in payload["data"]}
    assert {PACK_CODE_SITE_ACCESS, PACK_CODE_INCIDENT, PACK_CODE_INSPECTION_PREP}.issubset(codes)

    create_response = await async_client.post(
        f"/api/v1/packs/scenarios/{PACK_CODE_SITE_ACCESS}",
        json={"name": "Выход на объект", "is_active": True},
        headers=headers,
    )
    assert create_response.status_code == 201
    body = create_response.json()
    assert body["code"] == PACK_CODE_SITE_ACCESS
    assert body["is_active"] is True


@pytest.mark.anyio
async def test_template_metadata_validation(async_client: AsyncClient, make_auth_headers) -> None:
    headers = {**dict(async_client.headers), **await make_auth_headers()}
    template_bytes = _build_template()

    response = await async_client.post(
        "/api/v1/templates",
        files={"file": ("greeting.docx", template_bytes, DOCX_CONTENT_TYPE)},
        data={"name": "Greeting", "metadata": "not-json", **_version_metadata_form()},
        headers=headers,
    )

    assert response.status_code == 400
    body = response.json()
    assert body["code"] == "BAD_REQUEST"
    assert body["message"] == "metadata must be a valid JSON object"
    assert body["trace_id"]


@pytest.mark.anyio
async def test_pipeline_rejects_non_string_replacements(
    async_client: AsyncClient, make_auth_headers
) -> None:
    headers = {**dict(async_client.headers), **await make_auth_headers()}
    template_bytes = _build_template()
    create = await async_client.post(
        "/api/v1/templates",
        files={"file": ("greeting.docx", template_bytes, DOCX_CONTENT_TYPE)},
        data={"name": "Greeting", "metadata": "{}", **_version_metadata_form()},
        headers=headers,
    )
    template_id = create.json()["id"]

    response = await async_client.post(
        f"/api/v1/pipelines/{template_id}/run",
        json={
            "context": {"name": "World"},
            "replacements": {"Hello": 123},
        },
        params={"mode": "sync"},
        headers=headers,
    )

    assert response.status_code == 400
    body = response.json()
    assert body["code"] == "BAD_REQUEST"
    assert body["message"] == "replacements values must be strings"
    assert body["trace_id"]


@pytest.mark.anyio
async def test_template_creation_rejects_large_metadata(
    async_client: AsyncClient, make_auth_headers
) -> None:
    headers = {**dict(async_client.headers), **await make_auth_headers()}
    template_bytes = _build_template()
    oversized_value = "x" * MAX_METADATA_JSON_BYTES
    large_metadata = '{"payload":"' + oversized_value + '"}'

    response = await async_client.post(
        "/api/v1/templates",
        files={"file": ("greeting.docx", template_bytes, DOCX_CONTENT_TYPE)},
        data={"name": "Greeting", "metadata": large_metadata, **_version_metadata_form()},
        headers=headers,
    )

    assert response.status_code == 413
    body = response.json()
    assert body["code"] == "PAYLOAD_TOO_LARGE"
    assert body["message"] == (f"metadata payload cannot exceed {MAX_METADATA_JSON_BYTES} bytes")
    assert body["trace_id"]


@pytest.mark.anyio
async def test_template_creation_conflict_when_checksum_changes(
    async_client: AsyncClient, make_auth_headers
) -> None:
    headers = {**dict(async_client.headers), **await make_auth_headers()}
    template_bytes = _build_template()

    response = await async_client.post(
        "/api/v1/templates",
        files={"file": ("greeting.docx", template_bytes, DOCX_CONTENT_TYPE)},
        data={"name": "Greeting", "metadata": "{}", **_version_metadata_form()},
        headers=headers,
    )
    assert response.status_code == 201

    conflict_response = await async_client.post(
        "/api/v1/templates",
        files={"file": ("greeting-v2.docx", _build_alt_template(), DOCX_CONTENT_TYPE)},
        data={"name": "Greeting", "metadata": "{}", **_version_metadata_form()},
        headers=headers,
    )

    assert conflict_response.status_code == 409
    body = conflict_response.json()
    assert body["code"] == "CONFLICT"
    assert body["message"] == "Template with this name already exists"
    assert body["trace_id"]


@pytest.mark.anyio
async def test_template_creation_rejects_missing_version_metadata(
    async_client: AsyncClient, make_auth_headers
) -> None:
    headers = {**dict(async_client.headers), **await make_auth_headers()}
    template_bytes = _build_template()

    response = await async_client.post(
        "/api/v1/templates",
        files={"file": ("greeting.docx", template_bytes, DOCX_CONTENT_TYPE)},
        data={"name": "Greeting", "metadata": "{}"},
        headers=headers,
    )

    assert response.status_code == 422


@pytest.mark.anyio
async def test_pipeline_rejects_context_value_exceeding_limit(
    async_client: AsyncClient, make_auth_headers
) -> None:
    headers = {**dict(async_client.headers), **await make_auth_headers()}
    template_bytes = _build_template()
    create = await async_client.post(
        "/api/v1/templates",
        files={"file": ("greeting.docx", template_bytes, DOCX_CONTENT_TYPE)},
        data={"name": "Greeting", "metadata": "{}", **_version_metadata_form()},
        headers=headers,
    )
    template_id = create.json()["id"]

    oversized_value = "x" * (MAX_STRING_VALUE_LENGTH + 1)
    response = await async_client.post(
        f"/api/v1/pipelines/{template_id}/run",
        json={
            "context": {"large": oversized_value},
        },
        params={"mode": "sync"},
        headers=headers,
    )

    assert response.status_code == 413
    body = response.json()
    assert body["code"] == "PAYLOAD_TOO_LARGE"
    assert body["message"] == (
        f"context values must not exceed {MAX_STRING_VALUE_LENGTH} characters"
    )
    assert body["trace_id"]


@pytest.mark.anyio
async def test_pipeline_async_enqueue(
    async_client: AsyncClient,
    monkeypatch: pytest.MonkeyPatch,
    make_auth_headers,
) -> None:
    headers = {**dict(async_client.headers), **await make_auth_headers()}
    template_bytes = _build_template()
    create = await async_client.post(
        "/api/v1/templates",
        files={"file": ("greeting.docx", template_bytes, DOCX_CONTENT_TYPE)},
        data={"name": "Greeting", "metadata": "{}", **_version_metadata_form()},
        headers=headers,
    )
    template_id = create.json()["id"]

    called: dict[str, tuple[str, str]] = {}

    class StubResult:
        id = "stub-task"

    def _fake_apply_async(
        *,
        args: list[str],
        kwargs: dict[str, str],
        task_id: str,
        headers: dict[str, str],
    ):
        called["args"] = (args[0], args[1])
        assert task_id == args[0]
        assert headers.get("trace_id")
        return StubResult()

    monkeypatch.setattr("app.api.v1.router.run_pipeline_task.apply_async", _fake_apply_async)

    payload = {"context": {"name": "Async"}, "output_basename": "async"}
    response = await async_client.post(
        f"/api/v1/pipelines/{template_id}/run",
        json=payload,
        headers=headers,
    )

    assert response.status_code == 202
    body = response.json()
    assert body["status"] == "queued"
    assert "request" in body["result_metadata"]
    assert called["args"][1] == "test"


@pytest.mark.anyio
async def test_tenant_listing(async_client: AsyncClient) -> None:
    response = await async_client.get("/api/v1/tenants", headers={"X-Tenant": "test"})
    assert response.status_code == 200
    payload = response.json()
    assert payload["total"] >= 1
    assert any(item["slug"] == "test" for item in payload["items"])


@pytest.mark.anyio
async def test_company_crud(async_client: AsyncClient, make_auth_headers) -> None:
    headers = {**dict(async_client.headers), **await make_auth_headers()}
    initial = await async_client.get("/api/v1/companies", headers=headers)
    assert initial.status_code == 200
    assert initial.json()["total"] == 0

    create = await async_client.post(
        "/api/v1/companies",
        json={
            "name": "Test LLC",
            "inn": "7701234567",
            "kpp": "770101001",
            "ogrn": "1027700132195",
            "legal_address": "Main",
            "actual_address": "Branch",
            "director": "Ivan Ivanov",
            "phone_numbers": ["+7 495 000-00-00"],
            "work_types": ["Manufacturing"],
            "hazardous_factors": ["Noise"],
        },
        headers=headers,
    )
    assert create.status_code == 201
    company_id = create.json()["id"]

    duplicate = await async_client.post(
        "/api/v1/companies",
        json={"name": "Test LLC"},
        headers=headers,
    )
    assert duplicate.status_code == 409

    listing = await async_client.get("/api/v1/companies", headers=headers)
    assert listing.status_code == 200
    data = listing.json()
    assert data["total"] == 1
    assert data["items"][0]["id"] == company_id
    assert data["items"][0]["inn"] == "7701234567"


@pytest.mark.anyio
async def test_employee_listing_empty(async_client: AsyncClient, make_auth_headers) -> None:
    headers = {**dict(async_client.headers), **await make_auth_headers()}
    response = await async_client.get("/api/v1/employees", headers=headers)
    assert response.status_code == 200
    assert response.json()["total"] == 0
