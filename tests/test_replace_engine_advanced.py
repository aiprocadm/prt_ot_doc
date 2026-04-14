from __future__ import annotations

from io import BytesIO

from docx import Document

from app.modules.replace.engine import ReplaceOptions, replace_docx_bytes


def _build_docx_bytes(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _extract_first_paragraph(docx_bytes: bytes) -> str:
    doc = Document(BytesIO(docx_bytes))
    return doc.paragraphs[0].text


def test_replace_engine_supports_regex_mode() -> None:
    source = _build_docx_bytes("Order #123 dated 2026-01-01")
    result = replace_docx_bytes(
        source,
        mapping={r"\d{4}-\d{2}-\d{2}": "2026-12-31"},
        options=ReplaceOptions(regex=True),
    )
    assert "2026-12-31" in _extract_first_paragraph(result.docx_bytes)
    assert result.hits


def test_replace_engine_supports_whole_word_mode() -> None:
    source = _build_docx_bytes("cat concatenate cat")
    result = replace_docx_bytes(
        source,
        mapping={"cat": "dog"},
        options=ReplaceOptions(whole_word=True),
    )
    assert _extract_first_paragraph(result.docx_bytes) == "dog concatenate dog"

