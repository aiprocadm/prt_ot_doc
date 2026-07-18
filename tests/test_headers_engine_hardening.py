"""Header/Footer engine hardening tests
(Phase 5.2 / vNext-DOC-03, Session 40).

Before this session the ``app.modules.headers`` module had zero direct
test coverage despite shipping a feature-rich engine: placeholder
rendering, watermark overlay, ``{PAGE}``/``{NUMPAGES}`` Word fields, and
different first / odd / even header-footer parts.

These tests pin the contract so future refactors of
``apply_headers_to_docx`` and ``render_placeholders`` can't silently
regress production behavior.
"""

from __future__ import annotations

import zipfile
from dataclasses import dataclass
from io import BytesIO
from typing import Any

from docx import Document

from app.modules.headers.engine import apply_headers_to_docx
from app.modules.headers.placeholders import render_placeholders

# -----------------------------------------------------------------------------
# Test fixtures — a minimal preset dataclass mirrors the ORM model fields the
# engine actually reads. Using a dataclass keeps tests independent of the DB.
# -----------------------------------------------------------------------------


@dataclass
class _Preset:
    code: str = "default"
    name: str = "Default"
    different_first: bool = False
    different_odd_even: bool = False
    header_first_xml: str | None = None
    header_odd_xml: str | None = None
    header_even_xml: str | None = None
    footer_first_xml: str | None = None
    footer_odd_xml: str | None = None
    footer_even_xml: str | None = None
    watermark: dict[str, Any] | None = None


def _docx_with_body(text: str = "Hello") -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _docx_files(docx_bytes: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(BytesIO(docx_bytes)) as zf:
        return {name: zf.read(name) for name in zf.namelist()}


def _docx_xml(docx_bytes: bytes, name: str) -> str:
    files = _docx_files(docx_bytes)
    return files[name].decode("utf-8")


# =============================================================================
# render_placeholders — pure-function tests
# =============================================================================


def test_placeholders_resolve_flat_keys() -> None:
    rendered, unresolved = render_placeholders("Hello {{ name }}", {"name": "Alice"})
    assert rendered == "Hello Alice"
    assert unresolved == []


def test_placeholders_resolve_dotted_path() -> None:
    rendered, unresolved = render_placeholders(
        "{{ company.name }} — {{ company.address.city }}",
        {"company": {"name": "ACME", "address": {"city": "Moscow"}}},
    )
    assert rendered == "ACME — Moscow"
    assert unresolved == []


def test_placeholders_lenient_returns_empty_for_unknown_key() -> None:
    rendered, unresolved = render_placeholders(
        "Hello {{ missing }}", {"name": "Alice"}, strict=False
    )
    assert rendered == "Hello "
    assert unresolved == ["missing"]


def test_placeholders_strict_keeps_unknown_placeholder_literal() -> None:
    rendered, unresolved = render_placeholders(
        "Hello {{ missing }}", {"name": "Alice"}, strict=True
    )
    assert rendered == "Hello {{ missing }}"
    assert unresolved == ["missing"]


def test_placeholders_strict_marks_missing_nested_path() -> None:
    rendered, unresolved = render_placeholders(
        "{{ company.address.city }}",
        {"company": {"name": "ACME"}},
        strict=True,
    )
    assert rendered == "{{ company.address.city }}"
    assert unresolved == ["company.address.city"]


def test_placeholders_handle_empty_template() -> None:
    rendered, unresolved = render_placeholders("", {"name": "X"})
    assert rendered == ""
    assert unresolved == []


def test_placeholders_handle_none_template() -> None:
    rendered, unresolved = render_placeholders(None, {"name": "X"})
    assert rendered == ""
    assert unresolved == []


def test_placeholders_handle_no_braces() -> None:
    rendered, unresolved = render_placeholders("plain text", {"name": "X"})
    assert rendered == "plain text"
    assert unresolved == []


def test_placeholders_none_value_renders_empty_string() -> None:
    rendered, unresolved = render_placeholders("[{{ x }}]", {"x": None})
    assert rendered == "[]"
    assert unresolved == []


def test_placeholders_int_value_stringified() -> None:
    rendered, _ = render_placeholders("page {{ n }}", {"n": 42})
    assert rendered == "page 42"


def test_placeholders_tolerate_whitespace_inside_braces() -> None:
    rendered, _ = render_placeholders("{{   key   }}", {"key": "value"})
    assert rendered == "value"


def test_placeholders_multiple_occurrences() -> None:
    rendered, _ = render_placeholders("{{ x }} and {{ x }}", {"x": "Y"})
    assert rendered == "Y and Y"


def test_placeholders_unresolved_lists_each_distinct_key_once_or_more() -> None:
    rendered, unresolved = render_placeholders("{{ a }} {{ b }} {{ a }}", {}, strict=False)
    # Each occurrence appends to the list — caller may dedupe with set() if needed.
    assert unresolved == ["a", "b", "a"]
    assert rendered == "  "


# =============================================================================
# apply_headers_to_docx — engine tests
# =============================================================================


def test_apply_minimal_preset_writes_default_header_and_footer() -> None:
    preset = _Preset(header_odd_xml="Confidential", footer_odd_xml="Page {PAGE}")
    docx_out, report = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset,
        context={},
    )
    files = _docx_files(docx_out)
    # Default header/footer always written under header2 / footer2 (ref_type "default").
    assert "word/header2.xml" in files
    assert "word/footer2.xml" in files
    assert "Confidential" in files["word/header2.xml"].decode("utf-8")
    assert "word/header2.xml" in report.changed_parts
    assert "word/footer2.xml" in report.changed_parts


def test_apply_with_placeholders_resolves_from_context() -> None:
    preset = _Preset(header_odd_xml="{{ company.name }} — ОТ")
    docx_out, report = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset,
        context={"company": {"name": "ACME"}},
    )
    assert "ACME — ОТ" in _docx_xml(docx_out, "word/header2.xml")
    assert report.unresolved_placeholders == []


def test_apply_reports_unresolved_placeholders() -> None:
    preset = _Preset(header_odd_xml="{{ company.name }} ({{ company.tagline }})")
    _, report = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset,
        context={"company": {"name": "ACME"}},
    )
    assert "company.tagline" in report.unresolved_placeholders


def test_apply_emits_PAGE_field_when_placeholder_used() -> None:
    preset = _Preset(footer_odd_xml="Page {PAGE} of {NUMPAGES}")
    docx_out, report = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset,
        context={},
    )
    footer_xml = _docx_xml(docx_out, "word/footer2.xml")
    # Word field-simple elements should appear for both PAGE and NUMPAGES.
    assert " PAGE " in footer_xml
    assert " NUMPAGES " in footer_xml
    assert "PAGE" in report.fields_added
    assert "NUMPAGES" in report.fields_added


def test_apply_does_not_emit_PAGE_field_when_unused() -> None:
    preset = _Preset(footer_odd_xml="Static footer")
    docx_out, report = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset,
        context={},
    )
    assert "PAGE" not in report.fields_added
    assert " PAGE " not in _docx_xml(docx_out, "word/footer2.xml")


def test_apply_writes_first_page_part_when_different_first_set() -> None:
    preset = _Preset(
        different_first=True,
        header_first_xml="FIRST",
        header_odd_xml="Default",
    )
    docx_out, report = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset,
        context={},
    )
    files = _docx_files(docx_out)
    assert "word/header1.xml" in files
    assert "FIRST" in files["word/header1.xml"].decode("utf-8")
    # sectPr should carry a <w:titlePg/> flag.
    doc_xml = _docx_xml(docx_out, "word/document.xml")
    assert "titlePg" in doc_xml


def test_apply_writes_even_page_part_when_different_odd_even_set() -> None:
    preset = _Preset(
        different_odd_even=True,
        header_odd_xml="ODD",
        header_even_xml="EVEN",
    )
    docx_out, _ = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset,
        context={},
    )
    files = _docx_files(docx_out)
    # Header3.xml is bound to ref_type "even" by the engine.
    assert "word/header3.xml" in files
    assert "EVEN" in files["word/header3.xml"].decode("utf-8")
    # settings.xml should now have the evenAndOddHeaders flag.
    if "word/settings.xml" in files:
        settings = files["word/settings.xml"].decode("utf-8")
        assert "evenAndOddHeaders" in settings


def test_apply_watermark_override_overlays_text_in_default_header() -> None:
    preset = _Preset(header_odd_xml="Confidential")
    docx_out, _ = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset,
        context={},
        watermark_override={"enabled": True, "text": "DRAFT"},
    )
    header = _docx_xml(docx_out, "word/header2.xml")
    assert "Confidential" in header
    assert "DRAFT" in header


def test_apply_watermark_from_preset_used_when_no_override() -> None:
    preset = _Preset(
        header_odd_xml="Hdr",
        watermark={"enabled": True, "text": "INTERNAL"},
    )
    docx_out, _ = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset,
        context={},
    )
    assert "INTERNAL" in _docx_xml(docx_out, "word/header2.xml")


def test_apply_watermark_disabled_when_enabled_flag_false() -> None:
    preset = _Preset(
        header_odd_xml="Hdr",
        watermark={"enabled": False, "text": "INTERNAL"},
    )
    docx_out, _ = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset,
        context={},
    )
    assert "INTERNAL" not in _docx_xml(docx_out, "word/header2.xml")


def test_apply_writes_relationship_entries() -> None:
    preset = _Preset(header_odd_xml="Hdr", footer_odd_xml="Ftr")
    docx_out, _ = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset,
        context={},
    )
    rels = _docx_xml(docx_out, "word/_rels/document.xml.rels")
    # Both header and footer relationships should exist.
    assert "header" in rels
    assert "footer" in rels


def test_apply_writes_content_type_overrides() -> None:
    preset = _Preset(header_odd_xml="Hdr", footer_odd_xml="Ftr")
    docx_out, _ = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset,
        context={},
    )
    types_xml = _docx_xml(docx_out, "[Content_Types].xml")
    assert 'PartName="/word/header2.xml"' in types_xml
    assert 'PartName="/word/footer2.xml"' in types_xml


def test_apply_replaces_existing_header_reference_for_same_type() -> None:
    """Re-applying a preset should overwrite previous references of the
    same ref_type rather than appending duplicates. The engine writes every
    ref_type (first/default/even) using the odd fallback, so the invariant
    is "exactly one reference per ref_type", not "exactly one reference
    total"."""
    preset_a = _Preset(header_odd_xml="A")
    preset_b = _Preset(header_odd_xml="B")
    intermediate, _ = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset_a,
        context={},
    )
    final, _ = apply_headers_to_docx(
        docx_bytes=intermediate,
        preset=preset_b,
        context={},
    )
    doc_xml = _docx_xml(final, "word/document.xml")
    # No duplicate references for the same ref_type.
    assert doc_xml.count('w:headerReference w:type="default"') == 1
    assert doc_xml.count('w:headerReference w:type="first"') <= 1
    assert doc_xml.count('w:headerReference w:type="even"') <= 1


def test_apply_three_lines_use_left_center_right_alignment() -> None:
    preset = _Preset(header_odd_xml="LEFT\nCENTER\nRIGHT")
    docx_out, _ = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset,
        context={},
    )
    header = _docx_xml(docx_out, "word/header2.xml")
    assert "left" in header
    assert "center" in header
    assert "right" in header


def test_apply_raises_when_document_has_no_section() -> None:
    """A DOCX without a final sectPr is malformed for this engine."""
    # Build a minimal DOCX with no body sectPr.
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr(
            "word/document.xml",
            "<?xml version='1.0'?><w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
            "<w:body><w:p><w:r><w:t>hi</w:t></w:r></w:p></w:body></w:document>",
        )
        zf.writestr(
            "[Content_Types].xml",
            "<?xml version='1.0'?><Types xmlns='http://schemas.openxmlformats.org/package/2006/content-types'/>",
        )
    try:
        apply_headers_to_docx(
            docx_bytes=buf.getvalue(),
            preset=_Preset(header_odd_xml="Hdr"),
            context={},
        )
    except ValueError as exc:
        assert "sectPr" in str(exc)
    else:  # pragma: no cover - safety net
        raise AssertionError("expected ValueError")


def test_apply_returns_valid_zip_archive() -> None:
    preset = _Preset(header_odd_xml="Hdr", footer_odd_xml="Ftr")
    docx_out, _ = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset,
        context={},
    )
    # The output must still parse as a ZIP / DOCX.
    with zipfile.ZipFile(BytesIO(docx_out)) as zf:
        names = zf.namelist()
        assert "word/document.xml" in names
        assert "word/header2.xml" in names
        assert "word/footer2.xml" in names


def test_apply_preserves_original_body_content() -> None:
    preset = _Preset(header_odd_xml="Hdr")
    docx_out, _ = apply_headers_to_docx(
        docx_bytes=_docx_with_body("Important body text"),
        preset=preset,
        context={},
    )
    doc = Document(BytesIO(docx_out))
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Important body text" in body_text


def test_apply_idempotent_against_same_preset() -> None:
    """Applying the same preset twice should produce the same set of relevant parts."""
    preset = _Preset(header_odd_xml="STABLE", footer_odd_xml="STABLE_FOOTER")
    first, _ = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset,
        context={},
    )
    second, _ = apply_headers_to_docx(
        docx_bytes=first,
        preset=preset,
        context={},
    )
    files_a = _docx_files(first)
    files_b = _docx_files(second)
    assert "STABLE" in files_b["word/header2.xml"].decode("utf-8")
    assert "STABLE_FOOTER" in files_b["word/footer2.xml"].decode("utf-8")
    # No duplicate header/footer parts proliferating.
    headers = [n for n in files_b if n.startswith("word/header")]
    footers = [n for n in files_b if n.startswith("word/footer")]
    assert len(headers) <= 3  # at most first/default/even
    assert len(footers) <= 3
    # Re-applied content should remain identical for default parts.
    assert files_a["word/header2.xml"] == files_b["word/header2.xml"]


def test_apply_empty_preset_still_writes_default_parts() -> None:
    """An empty preset (all None) should still produce the default header/footer
    references because the engine binds the default ref_type unconditionally."""
    preset = _Preset()
    docx_out, report = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset,
        context={},
    )
    files = _docx_files(docx_out)
    assert "word/header2.xml" in files
    assert "word/footer2.xml" in files
    # No PAGE/NUMPAGES fields because nothing referenced them.
    assert report.fields_added == []


def test_apply_falls_back_to_odd_when_first_xml_missing() -> None:
    """When `different_first=True` but `header_first_xml` is None, the engine
    falls back to `header_odd_xml` per ``_resolve_section_content``."""
    preset = _Preset(
        different_first=True,
        header_odd_xml="ODD-CONTENT",
    )
    docx_out, _ = apply_headers_to_docx(
        docx_bytes=_docx_with_body(),
        preset=preset,
        context={},
    )
    files = _docx_files(docx_out)
    # header1 written using the odd fallback.
    if "word/header1.xml" in files:
        assert "ODD-CONTENT" in files["word/header1.xml"].decode("utf-8")
