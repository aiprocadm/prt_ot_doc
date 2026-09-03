"""Контур ПромБез срез-42 (Доп. №1 разд. 54.2): документы и отчётность.

Требование: «Документы и отчётность: ПЛА (план ликвидации аварий), положения,
приказы, отчёты для Ростехнадзора».

СВЕРКА нашла то же, что у ГО-ЧС (56.1 срез-6) и БДД (56.2 срез-9) — и это уже
третий раз один и тот же рисунок:

1. **Часть пункта закрыта базовым комплектом.** OPO_COMPLIANCE печатает
   паспорт ОПО, ПЛА и матрицу обучения — «ПЛА» из формулировки ТЗ закрыт с
   первого среза.
2. **Не хватало ТРЁХ вещей: положения, приказа и самих ОТЧЁТОВ.** Положение
   о производственном контроле — документ, который ФЗ-116 ст. 11 требует
   иметь у каждой эксплуатирующей организации; приказ о назначении
   ответственного за ПК; сведения об организации ПК за прошедший год и
   сведения об инцидентах — то, что подают в Ростехнадзор.
3. **Способ выбран продуктом трижды** — комплект форм БЕЗ миграций.

ОТЛИЧИЕ ОТ ПРЕДЫДУЩИХ КОМПЛЕКТОВ: подсказки из реестров есть С ПЕРВОГО дня —
реестры ОПО, устройств, аттестаций и плана ПК к этому срезу уже полны, и
заводить комплект «без автосбора», чтобы потом дописывать подсказки отдельной
волной (как было у БДД и экологии), значило бы повторять чужой путь без
причины.

ГРАНИЦА: платформа НЕ решает, обязана ли организация подавать сведения и по
каким ОПО, и НЕ считает аварии: реестр происшествий не размечен дисциплиной.
"""

from __future__ import annotations

import re
from datetime import date, timedelta
from io import BytesIO

import pytest
from docx import Document
from sqlalchemy import select

from app.core.disciplines import Discipline
from app.models.industrial_safety import (
    DeviceWorkRecord,
    HazardousFacility,
    ProductionControlMeasure,
    ProductionControlPlan,
    TechnicalDevice,
)
from app.models.inspections import Attestation
from app.models.master_data import Company, Person
from app.models.models import Tenant
from app.modules.packs.context import enrich_context
from app.modules.packs.definitions import (
    PACK_CODE_ECO_REPORTS,
    PACK_CODE_GOCHS_REPORTS,
    PACK_CODE_OPO,
    PACK_CODE_OPO_REPORTS,
    PACK_DEFINITIONS_BY_CODE,
)
from app.modules.packs.fields import questions_for

pytestmark = pytest.mark.anyio

_PACK = PACK_DEFINITIONS_BY_CODE[PACK_CODE_OPO_REPORTS]
_FIELDS = "/api/v1/packs/scenarios"
_LAST_YEAR = date.today().year - 1
_TODAY = date.today()


def _template_text(builder) -> str:
    """Весь текст шаблона: тело, колонтитулы — всё, где есть подстановки."""

    doc = Document(BytesIO(builder()))
    parts = [p.text for p in doc.paragraphs]
    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(parts)


class TestКомплектОтчётности:
    def test_комплект_размечен_промбезопасностью(self) -> None:
        assert _PACK.disciplines == (Discipline.INDUSTRIAL_SAFETY,)
        assert _PACK.metadata["discipline"] == "Промышленная безопасность"

    def test_печатает_положение_приказ_и_отчёты(self) -> None:
        """Три недостающих вещи из формулировки ТЗ; отчётов — два."""

        by_code = {t.code: t for t in _PACK.templates}
        assert by_code["pack_opo_pc_regulation"].category == "regulation"
        assert by_code["pack_opo_pc_order"].category == "order"
        assert by_code["pack_opo_pc_report"].category == "report"
        assert by_code["pack_opo_incident_report"].category == "report"

    def test_базовый_комплект_не_тронут(self) -> None:
        """Сторож: срез добавляет комплект, а не переписывает существующий.

        ПЛА остаётся в базовом комплекте — его заводят при запуске контура,
        а отчёт готовят раз в год; смешивать поводы нельзя.
        """

        base = PACK_DEFINITIONS_BY_CODE[PACK_CODE_OPO]
        assert {t.code for t in base.templates} == {
            "pack_opo_passport",
            "pack_opo_emergency_plan",
            "pack_opo_training_matrix",
        }

    def test_комплекты_не_пересекаются_шаблонами(self) -> None:
        base = {t.code for t in PACK_DEFINITIONS_BY_CODE[PACK_CODE_OPO].templates}
        assert base & {t.code for t in _PACK.templates} == set()

    def test_шаблоны_печатают_заявленное(self) -> None:
        """Заголовок документа — то, что увидит инспектор, а не код шаблона."""

        by_code = {t.code: _template_text(t.builder) for t in _PACK.templates}
        assert "Положение о производственном контроле" in by_code["pack_opo_pc_regulation"]
        assert "Приказ о назначении ответственного" in by_code["pack_opo_pc_order"]
        assert "Сведения об организации производственного контроля" in by_code["pack_opo_pc_report"]
        assert "Отчётный год: {{ data.opo_report_year }}" in by_code["pack_opo_pc_report"]
        assert "Сведения об инцидентах" in by_code["pack_opo_incident_report"]


class TestВопросыМастера:
    def test_обязателен_только_составитель(self) -> None:
        """Как у ГО и ЧС, а НЕ как у экологии — и это сознательно.

        У экологии обязателен ещё и год: там ВСЕ формы годовые. Здесь
        положение и приказ отчётного года не имеют вовсе, и требовать год
        значило бы не дать напечатать приказ без отчёта.
        """

        required = {f.name for f in questions_for(PACK_CODE_OPO_REPORTS) if f.required}
        assert required == {"opo_report_author"}

        gochs_required = {f.name for f in questions_for(PACK_CODE_GOCHS_REPORTS) if f.required}
        assert gochs_required == {"gochs_report_author"}
        eco_required = {f.name for f in questions_for(PACK_CODE_ECO_REPORTS) if f.required}
        assert "eco_report_year" in eco_required

    def test_у_каждого_вопроса_есть_подпись(self) -> None:
        for field in questions_for(PACK_CODE_OPO_REPORTS):
            assert field.label.strip(), field.name

    def test_каждая_подстановка_шаблона_есть_среди_вопросов(self) -> None:
        """Плейсхолдер без вопроса — прочерк в документе, который никто не
        сможет заполнить."""

        asked = {f.name for f in questions_for(PACK_CODE_OPO_REPORTS)}
        for template in _PACK.templates:
            text = _template_text(template.builder)
            for key in re.findall(r"\{\{ data\.(\w+) \}\}", text):
                assert key in asked, f"{template.code}: {key}"


class TestЧестныеУмолчания:
    def test_невнесённое_число_не_становится_нулём(self) -> None:
        """САМОЕ ВАЖНОЕ В СРЕЗЕ.

        Подставь ноль в «аварий и инцидентов» или «просроченных ЭПБ» — и
        отчёт В РОСТЕХНАДЗОР заявит, что инцидентов не было и просрочек нет,
        хотя их просто не посчитали. Отвечает за это подписавший.
        """

        context = enrich_context(
            PACK_CODE_OPO_REPORTS,
            {"company": {"name": "Тест"}},
            {"opo_report_author": "Иванов"},
        )
        data = context["data"]
        for key in (
            "opo_incidents_count",
            "opo_epb_overdue",
            "opo_attestations_overdue",
            "opo_facilities_count",
            "opo_pc_measures_done",
        ):
            assert data[key] == "сведения не внесены", key
            assert data[key] != "0", key

    def test_молчание_о_нарушениях_не_выдаётся_за_благополучие(self) -> None:
        context = enrich_context(PACK_CODE_OPO_REPORTS, {"company": {"name": "Тест"}}, {})
        data = context["data"]
        assert data["opo_violations_note"] == "сведения не внесены"
        assert data["opo_pc_responsible"] == "Ответственный не назначен"
        assert data["opo_approved_by"] == "не утверждено"
        assert data["opo_report_year"] == "не указан"

    def test_внесённые_ответы_доходят_до_документа(self) -> None:
        context = enrich_context(
            PACK_CODE_OPO_REPORTS,
            {"company": {"name": "Тест"}},
            {
                "opo_report_author": "Иванов",
                "opo_report_year": str(_LAST_YEAR),
                "opo_incidents_count": "2",
                "opo_pc_tasks": ["проверки", "обучение"],
            },
        )
        data = context["data"]
        assert data["opo_report_author"] == "Иванов"
        assert data["opo_report_year"] == str(_LAST_YEAR)
        assert data["opo_incidents_count"] == "2"
        assert data["opo_pc_tasks"] == "- проверки\n- обучение"


# --- подсказки из реестров -------------------------------------------------


async def _tenant_id(sessionmaker) -> str:
    async with sessionmaker() as session:
        tenant = (await session.execute(select(Tenant).where(Tenant.slug == "test"))).scalar_one()
        return str(tenant.id)


async def _facility(
    sessionmaker, *, number: str, hazard_class: str = "III", status: str = "registered"
) -> str:
    async with sessionmaker() as session:
        facility = HazardousFacility(
            tenant_id=await _tenant_id(sessionmaker),
            name=f"ОПО {number}",
            register_number=number,
            hazard_class=hazard_class,
            status=status,
        )
        session.add(facility)
        await session.commit()
        return str(facility.id)


async def _device(
    sessionmaker,
    facility_id: str,
    *,
    name: str,
    status: str = "in_operation",
    epb_valid_until: date | None = None,
) -> str:
    async with sessionmaker() as session:
        device = TechnicalDevice(
            tenant_id=await _tenant_id(sessionmaker),
            facility_id=facility_id,
            kind="pressure_vessel",
            name=name,
            status=status,
            epb_valid_until=epb_valid_until,
        )
        session.add(device)
        await session.commit()
        return str(device.id)


async def _work(sessionmaker, device_id: str, *, kind: str, performed_on: date) -> None:
    async with sessionmaker() as session:
        session.add(
            DeviceWorkRecord(
                tenant_id=await _tenant_id(sessionmaker),
                device_id=device_id,
                kind=kind,
                performed_on=performed_on,
                result="passed",
            )
        )
        await session.commit()


async def _attestation(sessionmaker, *, area_code: str | None, expires_at: date | None) -> None:
    async with sessionmaker() as session:
        tid = await _tenant_id(sessionmaker)
        company = (
            (await session.execute(select(Company).where(Company.tenant_id == tid)))
            .scalars()
            .first()
        )
        if company is None:
            company = Company(tenant_id=tid, name="Головная компания")
            session.add(company)
            await session.flush()
        person = Person(
            tenant_id=tid, company_id=company.id, last_name="Аттестуемый", first_name="А"
        )
        session.add(person)
        await session.flush()
        session.add(
            Attestation(
                tenant_id=tid,
                person_id=person.id,
                name="Аттестация",
                area_code=area_code,
                expires_at=expires_at,
            )
        )
        await session.commit()


async def _plan(sessionmaker, *, year: int, statuses: tuple[str, ...]) -> None:
    async with sessionmaker() as session:
        tid = await _tenant_id(sessionmaker)
        plan = ProductionControlPlan(tenant_id=tid, year=year, title=f"План ПК {year}")
        session.add(plan)
        await session.flush()
        for index, status in enumerate(statuses):
            session.add(
                ProductionControlMeasure(
                    tenant_id=tid,
                    plan_id=plan.id,
                    section="inspections",
                    title=f"Мероприятие {index}",
                    due_on=date(year, 6, 1),
                    status=status,
                    completed_on=date(year, 6, 1) if status == "done" else None,
                )
            )
        await session.commit()


async def _fields(async_client, headers) -> dict:
    response = await async_client.get(f"{_FIELDS}/{PACK_CODE_OPO_REPORTS}/fields", headers=headers)
    assert response.status_code == 200, response.text
    body = response.json()
    return {f["name"]: f for f in body["fields"]} | {"__body__": body}


class TestПодсказкиСостояние:
    async def test_считаются_только_действующие_опо(
        self, async_client, make_auth_headers, sessionmaker
    ) -> None:
        """Исключённый из госреестра объект остаётся ради истории, но объектом
        надзора быть перестаёт — в сведениях его нет."""

        headers = await make_auth_headers()
        await _facility(sessionmaker, number="А01-00001", hazard_class="II")
        await _facility(sessionmaker, number="А01-00002", hazard_class="III")
        await _facility(sessionmaker, number="А01-00003", hazard_class="III", status="excluded")

        fields = await _fields(async_client, headers)
        assert fields["opo_facilities_count"]["suggested"] == "2"
        assert (
            fields["opo_facilities_by_class"]["suggested"]
            == "I класс — 0, II класс — 1, III класс — 1, IV класс — 0"
        )
        assert "сегодня" in fields["opo_facilities_count"]["suggested_source"]

    async def test_списанное_устройство_не_считается_и_не_просрочено(
        self, async_client, make_auth_headers, sessionmaker
    ) -> None:
        headers = await make_auth_headers()
        facility = await _facility(sessionmaker, number="А01-00001")
        await _device(
            sessionmaker, facility, name="Сосуд", epb_valid_until=_TODAY - timedelta(days=1)
        )
        await _device(
            sessionmaker, facility, name="Котёл", epb_valid_until=_TODAY + timedelta(days=400)
        )
        await _device(sessionmaker, facility, name="Без заключения")
        await _device(
            sessionmaker,
            facility,
            name="Списан",
            status="decommissioned",
            epb_valid_until=_TODAY - timedelta(days=100),
        )

        fields = await _fields(async_client, headers)
        assert fields["opo_devices_count"]["suggested"] == "3"
        # «заключения нет» — не просрочка (словарь состояний ЭПБ)
        assert fields["opo_epb_overdue"]["suggested"] == "1"

    async def test_аттестации_чужой_дисциплины_не_попадают(
        self, async_client, make_auth_headers, sessionmaker
    ) -> None:
        """Проверка знаний водителя по «ПДД» — не аттестация по
        промбезопасности, и в отчёт для Ростехнадзора ей нельзя (прецедент
        сводки контура)."""

        headers = await make_auth_headers()
        await _attestation(sessionmaker, area_code="Б.9", expires_at=_TODAY + timedelta(days=30))
        await _attestation(sessionmaker, area_code="А.1", expires_at=_TODAY - timedelta(days=1))
        await _attestation(sessionmaker, area_code="ПДД", expires_at=_TODAY - timedelta(days=1))
        await _attestation(sessionmaker, area_code=None, expires_at=_TODAY - timedelta(days=1))
        # срок не указан — ни действует, ни просрочена: неполные сведения
        await _attestation(sessionmaker, area_code="Б.8", expires_at=None)

        fields = await _fields(async_client, headers)
        assert fields["opo_attestations_count"]["suggested"] == "1"
        assert fields["opo_attestations_overdue"]["suggested"] == "1"


class TestПодсказкиЗаГод:
    async def test_экспертизы_считаются_за_прошлый_год(
        self, async_client, make_auth_headers, sessionmaker
    ) -> None:
        """Сведения подают ЗА ПРОШЕДШИЙ год; текущий, неполный, был бы
        заведомо неверным числом. Год назван в источнике."""

        headers = await make_auth_headers()
        facility = await _facility(sessionmaker, number="А01-00001")
        device = await _device(sessionmaker, facility, name="Сосуд")
        await _work(sessionmaker, device, kind="epb", performed_on=date(_LAST_YEAR, 3, 1))
        await _work(sessionmaker, device, kind="epb", performed_on=date(_LAST_YEAR, 9, 1))
        await _work(sessionmaker, device, kind="maintenance", performed_on=date(_LAST_YEAR, 5, 1))
        await _work(sessionmaker, device, kind="epb", performed_on=date(_TODAY.year, 1, 15))

        fields = await _fields(async_client, headers)
        assert fields["opo_epb_done"]["suggested"] == "2"
        assert str(_LAST_YEAR) in fields["opo_epb_done"]["suggested_source"]

    async def test_мероприятия_плана_за_прошлый_год_без_отменённых(
        self, async_client, make_auth_headers, sessionmaker
    ) -> None:
        headers = await make_auth_headers()
        await _plan(
            sessionmaker, year=_LAST_YEAR, statuses=("done", "done", "planned", "cancelled")
        )
        await _plan(sessionmaker, year=_TODAY.year, statuses=("planned",))

        fields = await _fields(async_client, headers)
        assert fields["opo_pc_measures_planned"]["suggested"] == "3"
        assert fields["opo_pc_measures_done"]["suggested"] == "2"
        assert str(_LAST_YEAR) in fields["opo_pc_measures_planned"]["suggested_source"]

    async def test_без_плана_на_год_мероприятия_не_подсказываются(
        self, async_client, make_auth_headers, sessionmaker
    ) -> None:
        """«0 запланировано» читалось бы как «план был, но пустой», а плана
        не было — это другой факт."""

        headers = await make_auth_headers()
        await _plan(sessionmaker, year=_TODAY.year, statuses=("planned",))

        fields = await _fields(async_client, headers)
        assert fields["opo_pc_measures_planned"]["suggested"] is None
        assert fields["opo_pc_measures_done"]["suggested"] is None


class TestГраница:
    async def test_аварии_год_и_составитель_не_подсказываются(
        self, async_client, make_auth_headers, sessionmaker
    ) -> None:
        """Реестр происшествий не размечен дисциплиной: выдать все
        происшествия за инциденты на ОПО значило бы вписать в отчёт чужие
        числа. Год и составитель — решение специалиста."""

        headers = await make_auth_headers()
        await _facility(sessionmaker, number="А01-00001")
        fields = await _fields(async_client, headers)
        for key in (
            "opo_incidents_count",
            "opo_report_year",
            "opo_report_author",
            "opo_pc_responsible",
            "opo_violations_note",
        ):
            assert fields[key]["suggested"] is None, key

    async def test_у_каждой_подсказки_есть_источник(
        self, async_client, make_auth_headers, sessionmaker
    ) -> None:
        headers = await make_auth_headers()
        await _facility(sessionmaker, number="А01-00001")
        await _plan(sessionmaker, year=_LAST_YEAR, statuses=("done",))
        fields = await _fields(async_client, headers)
        body = fields.pop("__body__")
        assert body["suggestions_note"]
        assert any(f["suggested"] is not None for f in fields.values())
        for name, field in fields.items():
            if field["suggested"] is None:
                continue
            assert field["suggested_source"], name

    def test_платформа_не_обещает_сроки_и_автосбор(self) -> None:
        """Подсказка — не ответ; поля «срок подачи» или «автозаполнение»
        были бы обещанием несуществующего."""

        keys = {f.name for f in questions_for(PACK_CODE_OPO_REPORTS)}
        assert not any("deadline" in key or "autofill" in key or "auto_" in key for key in keys)
