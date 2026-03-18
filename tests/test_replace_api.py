from __future__ import annotations

from io import BytesIO
from uuid import uuid4

import pytest
from docx import Document
from httpx import ASGITransport, AsyncClient
from sqlalchemy import select
import json

from app.models.document import DocumentVersion
from app.services.file_storage import FileStorageService


async def _seed_document_version(sessionmaker, data_factory):
    async with sessionmaker() as session:
        tenant = await data_factory.ensure_tenant(session=session)
        user = await data_factory.create_user(tenant=tenant, email=f"{uuid4()}@example.com", session=session)
        company = await data_factory.create_company(tenant=tenant, name=f"Company {uuid4()}"[:36], session=session)
        person = await data_factory.create_person(tenant=tenant, company=company, first_name="Jane", last_name="Doe", session=session)
        template = await data_factory.create_template(tenant=tenant, name=f"Template {uuid4()}"[:36], session=session)
        document, version = await data_factory.create_document(
            tenant=tenant,
            company=company,
            person=person,
            template=template,
            creator=user,
            version_payload={"v": 1},
            version_file_key=f"tenants/{tenant.slug}/documents/{uuid4()}.docx",
            storage_key=f"tenants/{tenant.slug}/documents/{uuid4()}.docx",
            session=session,
        )
        await session.commit()
        return tenant.slug, document.id, version.id, version.file_key


@pytest.mark.anyio
async def test_replace_dry_run_apply_rollback(app_fixture, make_auth_headers, sessionmaker, data_factory) -> None:
    headers = await make_auth_headers()
    transport = ASGITransport(app=app_fixture)

    tenant_slug, _, version_id, version_key = await _seed_document_version(sessionmaker, data_factory)

    doc = Document()
    doc.add_paragraph("Hello {{company_name}}")
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "{{company_name}} in table"
    doc.sections[0].header.add_paragraph("{{company_name}}")
    doc.sections[0].footer.add_paragraph("{{company_name}}")
    buffer = BytesIO()
    doc.save(buffer)
    payload = buffer.getvalue()
    FileStorageService.default().put(version_key, payload, content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        create_map = await client.post(
            "/api/v1/replace-maps",
            headers={**headers, "X-Tenant": tenant_slug},
            data={"payload": json.dumps({"code": "demo-map", "name": "Demo map", "rules": [{"from": "{{company_name}}", "to": "OOO Demo"}]})},
        )
        assert create_map.status_code == 201
        map_id = create_map.json()["id"]

        dry = await client.post(
            f"/api/v1/documents/{version_id}/replace:dry-run",
            headers={**headers, "X-Tenant": tenant_slug, "Idempotency-Key": "replace-dry-1"},
            json={"replace_map_id": map_id},
        )
        assert dry.status_code == 200
        dry_payload = dry.json()
        assert dry_payload["replace_run_id"]
        assert dry_payload["hits_count"] and dry_payload["hits_count"] > 0

        apply = await client.post(
            f"/api/v1/documents/{version_id}/replace:apply",
            headers={**headers, "X-Tenant": tenant_slug, "Idempotency-Key": "replace-apply-1"},
            json={"replace_map_id": map_id},
        )
        assert apply.status_code == 200
        apply_payload = apply.json()
        assert apply_payload["new_document_version_id"]

        rollback = await client.post(
            f"/api/v1/replace-runs/{apply_payload['replace_run_id']}/rollback",
            headers={**headers, "X-Tenant": tenant_slug},
            json={"rollback_to_version_number": 1},
        )
        assert rollback.status_code == 200
        rollback_payload = rollback.json()
        assert rollback_payload["restored_from_version_id"] == version_id


@pytest.mark.anyio
async def test_replace_idempotency_conflict_same_key_different_payload(app_fixture, make_auth_headers, sessionmaker, data_factory) -> None:
    headers = await make_auth_headers()
    transport = ASGITransport(app=app_fixture)

    tenant_slug, _, version_id, version_key = await _seed_document_version(sessionmaker, data_factory)

    doc = Document()
    doc.add_paragraph("Hello {{company_name}}")
    buffer = BytesIO()
    doc.save(buffer)
    FileStorageService.default().put(version_key, buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        first_map = await client.post(
            "/api/v1/replace-maps",
            headers={**headers, "X-Tenant": tenant_slug},
            data={"payload": json.dumps({"code": "map-1", "name": "Map 1", "rules": [{"from": "{{company_name}}", "to": "OOO Demo"}]})},
        )
        second_map = await client.post(
            "/api/v1/replace-maps",
            headers={**headers, "X-Tenant": tenant_slug},
            data={"payload": json.dumps({"code": "map-2", "name": "Map 2", "rules": [{"from": "{{company_name}}", "to": "AO Demo"}]})},
        )
        assert first_map.status_code == 201
        assert second_map.status_code == 201

        first = await client.post(
            f"/api/v1/documents/{version_id}/replace:dry-run",
            headers={**headers, "X-Tenant": tenant_slug, "Idempotency-Key": "replace-conflict-1"},
            json={"replace_map_id": first_map.json()["id"]},
        )
        assert first.status_code == 200

        second = await client.post(
            f"/api/v1/documents/{version_id}/replace:dry-run",
            headers={**headers, "X-Tenant": tenant_slug, "Idempotency-Key": "replace-conflict-1"},
            json={"replace_map_id": second_map.json()["id"]},
        )
        assert second.status_code == 409
        assert second.json()["code"] == "IDEMPOTENCY_MISMATCH"

    async with sessionmaker() as session:
        versions = (await session.execute(select(DocumentVersion).where(DocumentVersion.document_id.is_not(None)))).scalars().all()
        assert len(versions) >= 1


def test_replace_engine_persists_patch_and_rollback(tmp_path):
    from app.domains.replace.engine import ReplaceEngine

    storage_path = tmp_path / "patches.json"
    engine = ReplaceEngine(storage_path=str(storage_path))
    patch = engine.commit({"company": "Old", "city": "Kazan"}, {"company": "New"})

    reloaded = ReplaceEngine(storage_path=str(storage_path))
    persisted = reloaded.get_patch(patch.patch_id)
    assert persisted is not None
    assert persisted.status == "applied"
    assert persisted.diff == [{"key": "company", "before": "Old", "after": "New", "replacement": "New"}]

    restored = reloaded.rollback(patch.patch_id)
    assert restored == {"company": "Old", "city": "Kazan"}
    assert reloaded.get_patch(patch.patch_id).status == "rolled_back"



def test_replace_engine_supports_nested_paths_and_idempotent_rollback(tmp_path):
    from app.domains.replace.engine import ReplaceEngine

    storage_path = tmp_path / "patches.json"
    engine = ReplaceEngine(storage_path=str(storage_path))
    context = {"company": {"name": "Old"}, "items": [{"title": "A"}, {"title": "B"}]}
    patch = engine.commit(context, {"company.name": "New", "items[1].title": "B2"})

    assert patch.changed_keys == ["company.name", "items[1].title"]
    assert patch.audit["changed_count"] == 2
    assert patch.after["company"]["name"] == "New"
    assert patch.after["items"][1]["title"] == "B2"

    restored_once = engine.rollback(patch.patch_id)
    restored_twice = engine.rollback(patch.patch_id)
    assert restored_once == context
    assert restored_twice == context
