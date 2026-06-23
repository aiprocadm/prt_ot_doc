"""API печати форм 29н: DOCX 200 + заголовки; невалидный формат 422; tenant-изоляция."""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from fastapi import status

from app.models.models import RoleEnum

REG = "/api/v1/medical/contingent/register/print"
NAMED = "/api/v1/medical/named-list/print"
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _docx_text(b: bytes) -> str:
    doc = Document(BytesIO(b))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


async def _seed_factor_chain(sessionmaker, data_factory, *, tenant=None):
    from app.models.models import MedicalFactor, Position, PositionHazardLink
    from app.models.risk import RiskHazard

    async with sessionmaker() as session:
        t = tenant or await data_factory.ensure_tenant(session=session)
        company = await data_factory.create_company(tenant=t, session=session)
        hazard = RiskHazard(tenant_id=t.id, code="noise", title="Шум", medical_factor_code="4.4")
        session.add(hazard)
        await session.flush()
        pos = Position(tenant_id=t.id, company_id=company.id, name="Сварщик")
        session.add(pos)
        await session.flush()
        session.add(PositionHazardLink(tenant_id=t.id, position_id=pos.id, hazard_id=hazard.id))
        await data_factory.create_person(
            tenant=t, company=company, session=session, position_id=pos.id,
            first_name="Иван", last_name="Петров",
        )
        session.add(
            MedicalFactor(
                tenant_id=t.id, code="4.4", name="Шум",
                exam_kinds=["periodic"], periodicity_months=12,
            )
        )
        await session.commit()


@pytest.mark.asyncio
async def test_register_print_docx_returns_file(async_client, sessionmaker, data_factory, make_auth_headers):
    await _seed_factor_chain(sessionmaker, data_factory)
    headers = await make_auth_headers(RoleEnum.ADMIN)
    resp = await async_client.get(f"{REG}?format=docx", headers=headers)
    assert resp.status_code == status.HTTP_200_OK, resp.text
    assert resp.headers["content-type"].startswith(_DOCX)
    assert "attachment" in resp.headers.get("content-disposition", "")
    assert resp.content[:2] == b"PK"
    assert "Сварщик" in _docx_text(resp.content)


@pytest.mark.asyncio
async def test_named_list_print_docx_returns_file(async_client, sessionmaker, data_factory, make_auth_headers):
    await _seed_factor_chain(sessionmaker, data_factory)
    headers = await make_auth_headers(RoleEnum.ADMIN)
    resp = await async_client.get(f"{NAMED}?format=docx", headers=headers)
    assert resp.status_code == status.HTTP_200_OK, resp.text
    assert resp.content[:2] == b"PK"
    assert "Петров Иван" in _docx_text(resp.content)


@pytest.mark.asyncio
async def test_print_invalid_format_422(async_client, make_auth_headers):
    headers = await make_auth_headers(RoleEnum.ADMIN)
    resp = await async_client.get(f"{REG}?format=xml", headers=headers)
    assert resp.status_code == status.HTTP_422_UNPROCESSABLE_ENTITY, resp.text


@pytest.mark.asyncio
async def test_print_tenant_isolation(async_client, sessionmaker, data_factory, make_auth_headers):
    # Тенант A с данными; тенант B пустой — печать B не содержит данных A.
    ta = await data_factory.ensure_tenant(slug="med-print-ta")
    await data_factory.ensure_tenant(slug="med-print-tb")
    await _seed_factor_chain(sessionmaker, data_factory, tenant=ta)
    headers_b = await make_auth_headers(
        RoleEnum.ADMIN, tenant="med-print-tb", email="admin-med-print-tb@example.com"
    )
    resp = await async_client.get(f"{REG}?format=docx", headers=headers_b)
    assert resp.status_code == status.HTTP_200_OK, resp.text
    assert "Сварщик" not in _docx_text(resp.content)
