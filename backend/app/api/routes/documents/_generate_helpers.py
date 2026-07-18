"""Documents API — generate-side helpers (ARCH-4 slice 5 split).

CSV/XLSX parsing, template scope resolution, company/person fetch, run resolution.
"""

import csv
import hashlib
import json
import logging
from io import StringIO
from typing import Any

from fastapi import (
    UploadFile,
)
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.routes.documents._common import (
    DocGenerateRequest,
    TemplateResolveCandidateRead,
    TemplateResolveRequest,
    _documents_bad_request,
    _documents_conflict,
    _documents_forbidden,
    _documents_not_found,
    _documents_payload_too_large,
)
from app.db.tenant_row_guard import assert_tenant_row_matches_session
from app.models.models import (
    Company,
    Person,
    PipelineRun,
    PipelineRunStatus,
    Template,
    TemplateVersion,
    TemplateVersionStatus,
    Tenant,
    User,
)

logger = logging.getLogger(__name__)


def _serialize_payload(payload: dict[str, Any]) -> str:
    try:
        return json.dumps(payload, sort_keys=True, separators=(",", ":"))
    except (TypeError, ValueError) as exc:
        raise _documents_bad_request("data must be JSON serializable") from exc


def _hash_payload(payload: dict[str, Any]) -> str:
    serialized = _serialize_payload(payload)
    digest = hashlib.sha256(serialized.encode("utf-8"))
    return digest.hexdigest()


def _ensure_payload_size(payload: dict[str, Any], *, limit: int, field: str) -> None:
    serialized = _serialize_payload(payload)
    size = len(serialized.encode("utf-8"))
    if size > limit:
        raise _documents_payload_too_large(f"{field} payload cannot exceed {limit} bytes")


def _parse_csv_payload(file: UploadFile) -> list[dict[str, Any]]:
    raw = file.file.read()
    content = raw.decode("utf-8")
    reader = csv.DictReader(StringIO(content))
    return [dict(row) for row in reader if any(row.values())]


def _parse_xlsx_payload(file: UploadFile) -> list[dict[str, Any]]:
    from io import BytesIO

    from openpyxl import load_workbook

    raw = file.file.read()
    workbook = load_workbook(BytesIO(raw), read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    data_rows = []
    for row in rows[1:]:
        entry = {header: value for header, value in zip(headers, row) if header}
        if any(value is not None and value != "" for value in entry.values()):
            data_rows.append(entry)
    return data_rows


def _apply_naming_pattern(pattern: str | None, row: dict[str, Any], row_index: int) -> str | None:
    if not pattern:
        return None

    class _SafeDict(dict):
        def __missing__(self, key: str) -> str:
            return ""

    return pattern.format_map(_SafeDict(row_index=row_index, **row))


def _extract_branding_generation_metadata(payload: DocGenerateRequest) -> dict[str, Any] | None:
    branding_preview = (
        payload.data.get("branding_preview") if isinstance(payload.data, dict) else None
    )
    if not isinstance(branding_preview, dict):
        return None
    data = branding_preview.get("data") if isinstance(branding_preview.get("data"), dict) else {}
    doc = data.get("doc") if isinstance(data.get("doc"), dict) else {}
    reproducibility = payload.data.get("reproducibility")
    if not isinstance(reproducibility, dict):
        reproducibility = (
            data.get("reproducibility") if isinstance(data.get("reproducibility"), dict) else {}
        )
    return {
        "site_id": payload.data.get("siteId") or data.get("branch", {}).get("id"),
        "preset_code": branding_preview.get("preset_code") or payload.data.get("headerPreset"),
        "document_title": doc.get("title") or payload.template_code,
        "document_number": doc.get("number"),
        "reproducibility": reproducibility,
        "apply_headers_payload": branding_preview,
    }


def _extract_document_version_id(run: PipelineRun) -> str | None:
    metadata = run.result_metadata or {}
    outputs = run.outputs or {}
    return metadata.get("document_version_id") or outputs.get("document_version_id")


def _normalize_scope_level(template: Template) -> str:
    # RC-013: read the indexed relational column first; fall back to the legacy
    # metadata_json["scope"] mirror only for rows written/backfilled before
    # normalization. Alias mapping is unchanged (organization/legal_entity →
    # company, branch → site, global → system) so matching/scoring is identical.
    raw_level = getattr(template, "scope_level", None)
    metadata = template.metadata_json if isinstance(template.metadata_json, dict) else {}
    scope = metadata.get("scope") if isinstance(metadata.get("scope"), dict) else {}
    mirror_level = scope.get("level") or scope.get("type")
    # ``scope_level`` is NOT NULL with a server default of "tenant", so a persisted
    # un-backfilled row (or one created bypassing the catalog API) surfaces
    # scope_level="tenant" even though its real scope still lives in the JSON
    # mirror — making the fallback unreachable if it only triggered on an empty
    # column. Treat the default "tenant" as "unset" and defer to a more specific
    # mirror level; an explicit non-default column value still wins outright.
    if (not raw_level or str(raw_level).strip().lower() == "tenant") and mirror_level:
        raw_level = mirror_level
    raw = str(raw_level or "tenant").strip().lower()
    aliases = {
        "organization": "company",
        "legal_entity": "company",
        "branch": "site",
        "global": "system",
    }
    return aliases.get(raw, raw)


def _scope_target_ids(template: Template) -> tuple[str | None, str | None]:
    # RC-013: prefer the indexed relational columns; fall back to legacy JSON
    # only when both columns are unset (un-backfilled row).
    company_id = getattr(template, "scope_company_id", None)
    site_id = getattr(template, "scope_site_id", None)
    if company_id is None and site_id is None:
        metadata = template.metadata_json if isinstance(template.metadata_json, dict) else {}
        scope = metadata.get("scope") if isinstance(metadata.get("scope"), dict) else {}
        company_id = scope.get("company_id")
        site_id = scope.get("site_id")
    return (str(company_id) if company_id else None, str(site_id) if site_id else None)


def _string_set(values: list[Any] | None) -> set[str]:
    if not values:
        return set()
    return {str(value).strip().lower() for value in values if str(value).strip()}


def _template_matches_request(
    *,
    template: Template,
    payload: TemplateResolveRequest,
) -> bool:
    metadata = template.metadata_json if isinstance(template.metadata_json, dict) else {}
    candidates = _string_set(
        [
            template.code,
            template.name,
            template.domain,
            metadata.get("category"),
            metadata.get("template_type"),
            *((metadata.get("tags") or []) if isinstance(metadata.get("tags"), list) else []),
            *(
                (metadata.get("case_types") or [])
                if isinstance(metadata.get("case_types"), list)
                else []
            ),
        ]
    )
    checks = [payload.case_type, payload.document_type, payload.category]
    required = [str(item).strip().lower() for item in checks if item and str(item).strip()]
    if not required:
        return True
    return all(item in candidates for item in required)


def _scope_score(level: str) -> int:
    return {
        "site": 400,
        "company": 300,
        "tenant": 200,
        "system": 100,
    }.get(level, 0)


def _scope_match_status(
    *,
    level: str,
    payload: TemplateResolveRequest,
    target_company_id: str | None,
    target_site_id: str | None,
) -> str:
    if level == "site":
        if payload.site_id and target_site_id and payload.site_id == target_site_id:
            return "exact"
        return "skip"
    if level == "company":
        if payload.company_id and target_company_id and payload.company_id == target_company_id:
            return "exact"
        return "skip"
    if level in {"tenant", "system"}:
        return "fallback"
    return "skip"


async def _fetch_template(
    session: AsyncSession,
    tenant: Tenant,
    *,
    template_code: str | None = None,
    template_id: str | None = None,
    template_version: int | None = None,
) -> tuple[Template, TemplateVersion]:
    tenant_scope = (str(tenant.id), tenant.slug)
    if not template_code:
        raise _documents_bad_request("template_code is required for template selection")
    if template_version is None:
        raise _documents_bad_request("template_version is required for template selection")
    filters: list[Any] = [
        Template.tenant_id.in_(tenant_scope),
        TemplateVersion.tenant_id.in_(tenant_scope),
        TemplateVersion.status == TemplateVersionStatus.ACTIVE,
        or_(Template.code == template_code, Template.name == template_code),
        TemplateVersion.version == template_version,
    ]

    stmt = (
        select(Template, TemplateVersion)
        .join(TemplateVersion, TemplateVersion.template_id == Template.id)
        .where(*filters)
        .limit(1)
    )
    row = (await session.execute(stmt)).first()
    if row is None:
        raise _documents_not_found(code="DOCUMENT_TEMPLATE_NOT_FOUND", message="Template not found")
    template, version = row
    if template_id and template.id != template_id:
        raise _documents_conflict(
            "template_id does not match template_code selection",
            code="DOCUMENT_TEMPLATE_ID_MISMATCH",
        )
    return template, version


async def _ensure_company(session: AsyncSession, tenant: Tenant, company_id: str) -> Company:
    company = await session.get(Company, company_id)
    if company is None:
        raise _documents_not_found(code="DOCUMENT_COMPANY_NOT_FOUND", message="Company not found")
    try:
        assert_tenant_row_matches_session(
            session,
            company,
            mismatch_event="api.documents.ensure_company.tenant_scope_mismatch",
            not_found_message="tenant_mismatch",
            expected_tenant_id=str(tenant.id),
        )
    except ValueError:
        raise _documents_forbidden(
            code="DOCUMENT_COMPANY_TENANT_MISMATCH",
            message="Tenant mismatch for company resource",
        ) from None
    return company


async def _resolve_template_candidates(
    *,
    session: AsyncSession,
    tenant: Tenant,
    payload: TemplateResolveRequest,
) -> list[TemplateResolveCandidateRead]:
    tenant_scope = (str(tenant.id), tenant.slug)
    stmt = (
        select(Template, TemplateVersion)
        .join(TemplateVersion, TemplateVersion.template_id == Template.id)
        .where(
            Template.tenant_id.in_(tenant_scope),
            TemplateVersion.tenant_id.in_(tenant_scope),
            Template.deleted_at.is_(None),
            TemplateVersion.deleted_at.is_(None),
            TemplateVersion.status == TemplateVersionStatus.ACTIVE,
        )
    )
    rows = (await session.execute(stmt)).all()
    candidates: list[TemplateResolveCandidateRead] = []
    for template, version in rows:
        if not _template_matches_request(template=template, payload=payload):
            continue
        level = _normalize_scope_level(template)
        target_company_id, target_site_id = _scope_target_ids(template)
        scope_match = _scope_match_status(
            level=level,
            payload=payload,
            target_company_id=target_company_id,
            target_site_id=target_site_id,
        )
        if scope_match == "skip":
            continue
        score = _scope_score(level) + int(version.version)
        if template.current_version_id and template.current_version_id == version.id:
            score += 50
        rationale = [f"scope={level}", f"scope_match={scope_match}", f"version={version.version}"]
        if payload.case_type:
            rationale.append(f"case_type={payload.case_type}")
        if payload.document_type:
            rationale.append(f"document_type={payload.document_type}")
        if payload.category:
            rationale.append(f"category={payload.category}")
        candidates.append(
            TemplateResolveCandidateRead(
                template_id=template.id,
                template_code=template.code or template.name,
                template_name=template.name,
                template_version=int(version.version),
                scope_level=level,
                scope_match=scope_match,
                score=score,
                rationale=rationale,
            )
        )
    return sorted(
        candidates,
        key=lambda item: (item.score, item.template_version, item.template_code),
        reverse=True,
    )


async def _ensure_person(
    session: AsyncSession, person_id: str | None, company: Company
) -> Person | None:
    if person_id is None:
        return None
    person = await session.get(Person, person_id)
    if person is None:
        raise _documents_not_found(code="DOCUMENT_PERSON_NOT_FOUND", message="Person not found")
    try:
        assert_tenant_row_matches_session(
            session,
            person,
            mismatch_event="api.documents.ensure_person.tenant_scope_mismatch",
            not_found_message="tenant_mismatch",
            expected_tenant_id=str(company.tenant_id),
        )
    except ValueError:
        raise _documents_forbidden(
            code="DOCUMENT_PERSON_TENANT_MISMATCH",
            message="Tenant mismatch for person resource",
        ) from None
    if person.company_id != company.id:
        raise _documents_bad_request("person_id does not belong to the provided company")
    return person


async def _resolve_run(
    session: AsyncSession,
    *,
    idempotency_key: str,
    tenant: Tenant,
    template: Template,
    template_version: TemplateVersion,
    company: Company,
    person: Person | None,
    payload_hash: str,
    current_user: User,
    context: dict[str, Any],
    correlation_id: str,
    npa_binding_id: str | None,
    visible_passport: bool,
) -> tuple[PipelineRun, bool]:
    stmt = select(PipelineRun).where(
        PipelineRun.tenant_id == tenant.id,
        PipelineRun.idempotency_key == idempotency_key,
    )
    existing = (await session.execute(stmt)).scalar_one_or_none()
    if existing is not None:
        metadata = existing.result_metadata or {}
        if existing.template_id != template.id:
            raise _documents_conflict(
                "Idempotency key already used", code="DOCUMENT_IDEMPOTENCY_KEY_CONFLICT"
            )
        if existing.template_version_id != template_version.id:
            raise _documents_conflict(
                "Idempotency key already used", code="DOCUMENT_IDEMPOTENCY_KEY_CONFLICT"
            )
        if metadata.get("company_id") != company.id:
            raise _documents_conflict(
                "Idempotency key already used", code="DOCUMENT_IDEMPOTENCY_KEY_CONFLICT"
            )
        expected_person = person.id if person else None
        if metadata.get("person_id") != expected_person:
            raise _documents_conflict(
                "Idempotency key already used", code="DOCUMENT_IDEMPOTENCY_KEY_CONFLICT"
            )
        if metadata.get("payload_hash") != payload_hash:
            raise _documents_conflict(
                "Idempotency key already used", code="DOCUMENT_IDEMPOTENCY_KEY_CONFLICT"
            )
        if existing.context != context:
            raise _documents_conflict(
                "Idempotency key already used", code="DOCUMENT_IDEMPOTENCY_KEY_CONFLICT"
            )
        if existing.status == PipelineRunStatus.ERROR:
            raise _documents_conflict(
                "Idempotency key refers to a failed generation task",
                code="DOCUMENT_IDEMPOTENCY_KEY_FAILED_RUN",
            )
        return existing, False

    run = PipelineRun(
        tenant_id=tenant.id,
        template_id=template.id,
        template_version_id=template_version.id,
        status=PipelineRunStatus.QUEUED,
        context=dict(context),
        idempotency_key=idempotency_key,
        result_metadata={
            "company_id": company.id,
            "person_id": person.id if person else None,
            "initiated_by": current_user.id,
            "payload_hash": payload_hash,
            "correlation_id": correlation_id,
            "npa_binding_id": npa_binding_id,
            "visible_passport": visible_passport,
        },
    )
    session.add(run)
    await session.flush()
    return run, True
