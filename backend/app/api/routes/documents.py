"""Document generation API endpoints."""

import csv
import hashlib
import json
import logging
from datetime import datetime, timezone
from io import StringIO
from typing import Annotated, Any
from uuid import UUID

from fastapi import (
    APIRouter,
    Depends,
    Form,
    Header,
    HTTPException,
    Query,
    Request,
    Response,
    UploadFile,
    status,
)
from pydantic import BaseModel, ConfigDict, Field, model_validator
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.api.dependencies import get_file_storage_service, get_session, get_tenant_record
from app.db.tenant_row_guard import assert_tenant_row_matches_session
from app.core.audit_decorator import audit_operation
from app.core.config import get_settings
from app.core.errors import api_problem_detail
from app.core.idempotency import compute_request_hash
from app.core.payload_constraints import PayloadConstraintError, enforce_mapping_constraints
from app.core.rate_limit import generate_per_tenant, ip_tenant_key, limiter
from app.core.security import AccessContext, abac, rbac
from app.core.tracing import get_trace_id
from app.models.document import (
    Document,
    DocumentBatchItem,
    DocumentBatchItemStatus,
    DocumentBatchRun,
    DocumentBatchStatus,
    DocumentJobStatus,
    DocumentStatus,
    DocumentVersion,
    DocumentVersionStatus,
)
from app.models.file import File
from app.models.models import (
    Company,
    Person,
    PipelineRun,
    PipelineRunStatus,
    Template,
    TemplateVersion,
    TemplateVersionStatus,
    Tenant,
    User,
)
from app.schemas.document import (
    DocumentBatchRunRead,
    DocumentCompanySummaryRead,
    DocumentDependencyMapRead,
    DocumentDependencyNpaBindingRead,
    DocumentFileLinkRead,
    DocumentHistoryEntryRead,
    DocumentPaginationRead,
    DocumentPipelineStageRead,
    DocumentRead,
    DocumentReadinessRead,
    DocumentStatusUpdate,
    DocumentUiListResponse,
    DocumentUiRead,
    DocumentVersionCompareRead,
    DocumentVersionDataDiffRead,
)
from app.schemas.task import TaskAcceptedResponse, TaskStatusResponse
from app.services.audit import AuditService
from app.services.billing import BillingService
from app.services.document_insights import (
    build_document_dependency_map,
    diff_version_data_json,
    load_document_versions_for_compare,
)
from app.services.document_readiness import compute_document_readiness
from app.services.documents import (
    DocumentNotFoundError,
    DocumentWorkflowService,
    InvalidStatusTransitionError,
)
from app.services.file_storage import FileStorageService
from app.services.idempotency import IdempotencyService, normalize_idempotency_key
from app.services.pipelines_orchestrator import DocumentPipelineOrchestrator
from app.tasks import generate_document_batch_item_task, generate_document_task

router = APIRouter()
logger = logging.getLogger(__name__)


def _generate_internal_error_problem() -> dict[str, Any]:
    """Клиентский ответ и тело для idempotency store без утечки внутренних исключений."""

    return api_problem_detail(
        code="INTERNAL_ERROR",
        message="Произошла внутренняя ошибка при обработке запроса.",
        error_type="server",
    )


def _documents_bad_request(message: str) -> HTTPException:
    return HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail=api_problem_detail(
            code="DOCUMENTS_BAD_REQUEST",
            message=message,
            error_type="documents",
        ),
    )


def _dispatch_celery_task(task, *, args: list[str], kwargs: dict[str, str], task_id: str | None = None, headers: dict[str, str] | None = None) -> None:
    task.apply_async(args=args, kwargs=kwargs, task_id=task_id, headers=headers)


SessionDep = Depends(get_session)
TenantDep = Depends(get_tenant_record)


def _tenant_resource_id(tenant: Tenant = Depends(get_tenant_record)) -> UUID | None:
    return getattr(tenant, "id", None)


_DOCUMENT_RUN_ROLES = ["admin", "employee", "client_admin"]
_DOCUMENT_READ_ROLES = [
    "owner",
    "admin",
    "employee",
    "line_manager",
    "hr",
    "ot_specialist",
    "ot_head",
    "client_admin",
    "client_user",
    "clerk",
]
_DOCUMENT_STATUS_ROLES = ["admin"]

AccessDep = Depends(
    abac(
        _tenant_resource_id,
        required_roles=_DOCUMENT_RUN_ROLES,
        action="manage documents",
    )
)
ReadAccessDep = Depends(
    abac(
        _tenant_resource_id,
        required_roles=_DOCUMENT_READ_ROLES,
        action="read documents",
    )
)
StatusAccessDep = Depends(rbac(_DOCUMENT_STATUS_ROLES))




class GenerateTemplateRef(BaseModel):
    code: str
    version: int


class GenerateDataPayload(BaseModel):
    type: str
    payload: dict[str, Any] | None = None
    file_id: str | None = None


class GenerateStepOptions(BaseModel):
    apply_headers: dict[str, Any] | None = None
    replace: dict[str, Any] | None = None
    pdf: dict[str, Any] | None = None
    zip: dict[str, Any] | None = None


class GeneratePipelinePayload(BaseModel):
    profile_code: str | None = None
    steps: GenerateStepOptions | None = None


class DocGeneratePipelineRequest(BaseModel):
    template: GenerateTemplateRef
    data: GenerateDataPayload
    pipeline: GeneratePipelinePayload | None = None
    npa_binding_id: str | None = None


class DocGenerateRequest(BaseModel):
    """Incoming payload for document generation requests."""

    model_config = ConfigDict(extra="forbid")

    template_code: str | None = Field(default=None, min_length=1, max_length=255)
    template_id: str | None = Field(default=None, min_length=1)
    template_version: int | None = Field(default=None, ge=1)
    company_id: str | None = Field(default=None, min_length=1)
    person_id: str | None = Field(default=None)
    data: dict[str, Any] = Field(default_factory=dict)
    pipeline_profile_id: str | None = Field(default=None)
    input_source_id: str | None = Field(default=None)
    inline_data: dict[str, Any] | None = Field(default=None)
    options: dict[str, Any] = Field(default_factory=dict)
    visible_passport: bool = Field(default=True)
    npa_binding_id: str | None = Field(default=None)

    @model_validator(mode="after")
    def _ensure_identifier(self) -> "DocGenerateRequest":
        if not self.template_code:
            raise ValueError("template_code is required to select a template")
        if self.template_version is None:
            raise ValueError("template_version is required to select a template")
        if self.inline_data is None and self.input_source_id is None and self.company_id is None:
            raise ValueError("company_id required for legacy mode")
        return self


def _map_document_status(document: Document) -> str:
    if document.job and document.job.status in {DocumentJobStatus.QUEUED, DocumentJobStatus.PROCESSING}:
        return "generating"
    if document.job and document.job.status == DocumentJobStatus.FAILED:
        return "error"
    if document.status in {DocumentStatus.GENERATED, DocumentStatus.APPROVED, DocumentStatus.SIGNED, DocumentStatus.ARCHIVED}:
        return "ready"
    if document.status == DocumentStatus.REVOKED:
        return "error"
    return "draft"


def _document_matches_frontend_status(document: Document, status_value: str | None) -> bool:
    if not status_value:
        return True
    return _map_document_status(document) == status_value


def _build_file_link(
    *,
    file_record: File | None,
    document_id: str,
    download_path: str,
    storage_key: str | None = None,
    created_at: datetime | None = None,
) -> DocumentFileLinkRead | None:
    if file_record is None and not storage_key:
        return None
    name = None
    if file_record is not None:
        name = file_record.original_name
    elif storage_key:
        name = storage_key.rsplit("/", 1)[-1]
    return DocumentFileLinkRead(
        id=file_record.id if file_record is not None else f"document:{document_id}",
        url=download_path,
        name=name or f"document-{document_id}",
        mime_type=file_record.mime if file_record is not None else "application/octet-stream",
        size=file_record.size if file_record is not None else 0,
        created_at=file_record.created_at if file_record is not None else (created_at or datetime.now(timezone.utc)),
    )


def _resolve_latest_version(document: Document) -> DocumentVersion | None:
    if not document.versions:
        return None
    return max(document.versions, key=lambda version: (version.version_number, version.created_at))


def _build_document_history(document: Document) -> list[DocumentHistoryEntryRead]:
    history: list[DocumentHistoryEntryRead] = []
    for version in sorted(document.versions, key=lambda item: (item.version_number, item.created_at), reverse=True):
        history.append(
            DocumentHistoryEntryRead(
                id=version.id,
                created_at=version.created_at,
                updated_at=version.updated_at,
                document_id=version.document_id,
                status=version.status.value if isinstance(version.status, DocumentVersionStatus) else str(version.status),
                storage=_build_file_link(
                    file_record=version.file,
                    document_id=document.id,
                    download_path=f"/api/v1/documents/{document.id}/download",
                    storage_key=version.file_key,
                    created_at=version.created_at,
                ),
            )
        )
    return history


def _build_document_ui_read(document: Document) -> DocumentUiRead:
    latest_version = _resolve_latest_version(document)
    current_file = document.file or (latest_version.file if latest_version else None)
    company = document.company
    if company is None:
        raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR, "Document company is missing")
    template_name = (document.template.name if document.template else None) or f"Document {document.id[:8]}"
    template_type = (document.template.domain if document.template else None) or (document.template.code if document.template else None) or "document"
    return DocumentUiRead(
        id=document.id,
        created_at=document.created_at,
        updated_at=document.updated_at,
        name=template_name,
        type=template_type,
        company=DocumentCompanySummaryRead(
            id=company.id,
            created_at=company.created_at,
            updated_at=company.updated_at,
            name=company.name,
            inn=company.inn or "",
            status="active",
        ),
        status=_map_document_status(document),
        version=str(latest_version.version_number if latest_version else 1),
        current_version_id=latest_version.id if latest_version else None,
        template_id=document.template_id,
        storage=_build_file_link(
            file_record=current_file,
            document_id=document.id,
            download_path=f"/api/v1/documents/{document.id}/download",
            storage_key=document.storage_key or (latest_version.file_key if latest_version else None),
            created_at=document.updated_at,
        ),
        history=_build_document_history(document),
    )


def _document_read_query(tenant_id: str):
    return (
        select(Document)
        .where(Document.tenant_id == tenant_id)
        .options(
            selectinload(Document.company),
            selectinload(Document.template),
            selectinload(Document.template_version),
            selectinload(Document.file),
            selectinload(Document.job),
            selectinload(Document.versions).selectinload(DocumentVersion.file),
        )
    )


@router.get("", response_model=DocumentUiListResponse)
async def list_documents(
    tenant: Tenant = TenantDep,
    session: AsyncSession = SessionDep,
    access: AccessContext = ReadAccessDep,
    search: str | None = Query(default=None),
    status_value: str | None = Query(default=None, alias="status"),
    company_id: str | None = Query(default=None),
    type_value: str | None = Query(default=None, alias="type"),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=10, ge=1, le=200),
) -> DocumentUiListResponse:
    access.ensure_tenant_access(tenant.id, action="read documents")

    stmt = _document_read_query(str(tenant.id))
    if company_id:
        stmt = stmt.where(Document.company_id == company_id)

    scoped_company_ids = access.claims.get("company_ids")
    if isinstance(scoped_company_ids, list) and scoped_company_ids:
        stmt = stmt.where(Document.company_id.in_([str(value) for value in scoped_company_ids]))
    elif access.company_id:
        stmt = stmt.where(Document.company_id == str(access.company_id))

    if search:
        pattern = f"%{search.strip()}%"
        stmt = stmt.where(
            or_(
                Document.id.ilike(pattern),
                Document.company.has(Company.name.ilike(pattern)),
                Document.template.has(or_(Template.name.ilike(pattern), Template.code.ilike(pattern))),
            )
        )
    if type_value:
        stmt = stmt.where(
            Document.template.has(
                or_(Template.domain == type_value, Template.code == type_value)
            )
        )
    if status_value == "generating":
        stmt = stmt.where(
            or_(
                Document.job.has(status=DocumentJobStatus.QUEUED),
                Document.job.has(status=DocumentJobStatus.PROCESSING),
            )
        )
    elif status_value == "ready":
        stmt = stmt.where(
            Document.status.in_(
                [
                    DocumentStatus.GENERATED,
                    DocumentStatus.APPROVED,
                    DocumentStatus.SIGNED,
                    DocumentStatus.ARCHIVED,
                ]
            )
        )
    elif status_value == "draft":
        stmt = stmt.where(Document.status.in_([DocumentStatus.DRAFT, DocumentStatus.REVIEW]))
    elif status_value == "error":
        stmt = stmt.where(
            or_(
                Document.status == DocumentStatus.REVOKED,
                Document.job.has(status=DocumentJobStatus.FAILED),
            )
        )

    total = int(await session.scalar(select(func.count()).select_from(stmt.subquery())) or 0)
    rows = list(
        (
            await session.execute(
                stmt.order_by(Document.updated_at.desc()).offset((page - 1) * page_size).limit(page_size)
            )
        ).scalars().all()
    )

    items: list[DocumentUiRead] = []
    for document in rows:
        access.ensure_abac(
            action="read document",
            company_id=document.company_id,
            document_id=document.id,
            document_status=document.status.value if hasattr(document.status, "value") else str(document.status),
            document_owner_id=document.created_by,
            site_id=document.site_id,
            site_company_id=document.company_id,
        )
        items.append(_build_document_ui_read(document))

    return DocumentUiListResponse(
        items=items,
        pagination=DocumentPaginationRead(page=page, page_size=page_size, total=total),
    )


@router.get("/{document_id}", response_model=DocumentUiRead)
async def get_document(
    document_id: str,
    tenant: Tenant = TenantDep,
    session: AsyncSession = SessionDep,
    access: AccessContext = ReadAccessDep,
) -> DocumentUiRead:
    access.ensure_tenant_access(tenant.id, action="read document")
    document = (
        await session.execute(_document_read_query(str(tenant.id)).where(Document.id == document_id))
    ).scalar_one_or_none()
    if document is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Document not found")
    access.ensure_abac(
        action="read document",
        company_id=document.company_id,
        document_id=document.id,
        document_status=document.status.value if hasattr(document.status, "value") else str(document.status),
        document_owner_id=document.created_by,
        site_id=document.site_id,
        site_company_id=document.company_id,
    )
    return _build_document_ui_read(document)


@router.get("/{document_id}/readiness", response_model=DocumentReadinessRead)
async def get_document_readiness_endpoint(
    document_id: str,
    tenant: Tenant = TenantDep,
    session: AsyncSession = SessionDep,
    access: AccessContext = ReadAccessDep,
) -> DocumentReadinessRead:
    access.ensure_tenant_access(tenant.id, action="read document")
    document = (
        await session.execute(_document_read_query(str(tenant.id)).where(Document.id == document_id))
    ).scalar_one_or_none()
    if document is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Document not found")
    access.ensure_abac(
        action="read document",
        company_id=document.company_id,
        document_id=document.id,
        document_status=document.status.value if hasattr(document.status, "value") else str(document.status),
        document_owner_id=document.created_by,
        site_id=document.site_id,
        site_company_id=document.company_id,
    )
    snap = compute_document_readiness(document)
    return DocumentReadinessRead(
        score=snap.score,
        blockers=snap.blockers,
        recommended_actions=snap.recommended_actions,
        pipeline_stages=[
            DocumentPipelineStageRead(
                stage_id=s.stage_id,
                label=s.label,
                complete=s.complete,
                detail=s.detail,
            )
            for s in snap.pipeline_stages
        ],
    )


@router.get("/{document_id}/versions/compare", response_model=DocumentVersionCompareRead)
async def compare_document_versions_endpoint(
    document_id: str,
    left_version_id: str = Query(...),
    right_version_id: str = Query(...),
    tenant: Tenant = TenantDep,
    session: AsyncSession = SessionDep,
    access: AccessContext = ReadAccessDep,
) -> DocumentVersionCompareRead:
    access.ensure_tenant_access(tenant.id, action="read document")
    document = (
        await session.execute(_document_read_query(str(tenant.id)).where(Document.id == document_id))
    ).scalar_one_or_none()
    if document is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Document not found")
    access.ensure_abac(
        action="read document",
        company_id=document.company_id,
        document_id=document.id,
        document_status=document.status.value if hasattr(document.status, "value") else str(document.status),
        document_owner_id=document.created_by,
        site_id=document.site_id,
        site_company_id=document.company_id,
    )
    try:
        left_v, right_v = await load_document_versions_for_compare(
            session,
            tenant_id=str(tenant.id),
            document_id=document_id,
            left_version_id=left_version_id,
            right_version_id=right_version_id,
        )
    except ValueError:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Version not found") from None

    left_data = dict(left_v.data_json or {})
    right_data = dict(right_v.data_json or {})
    raw_diffs = diff_version_data_json(left_data, right_data)
    return DocumentVersionCompareRead(
        document_id=document_id,
        left_version_id=left_v.id,
        right_version_id=right_v.id,
        diffs=[
            DocumentVersionDataDiffRead(
                field=d["field"],
                before=d.get("before"),
                after=d.get("after"),
                change=str(d["change"]),
            )
            for d in raw_diffs
        ],
        template_version_changed=(
            (left_v.template_version_id or "") != (right_v.template_version_id or "")
            or (left_v.template_version or "") != (right_v.template_version or "")
        ),
    )


@router.get("/{document_id}/dependency-map", response_model=DocumentDependencyMapRead)
async def get_document_dependency_map_endpoint(
    document_id: str,
    tenant: Tenant = TenantDep,
    session: AsyncSession = SessionDep,
    access: AccessContext = ReadAccessDep,
) -> DocumentDependencyMapRead:
    access.ensure_tenant_access(tenant.id, action="read document")
    document = (
        await session.execute(_document_read_query(str(tenant.id)).where(Document.id == document_id))
    ).scalar_one_or_none()
    if document is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Document not found")
    access.ensure_abac(
        action="read document",
        company_id=document.company_id,
        document_id=document.id,
        document_status=document.status.value if hasattr(document.status, "value") else str(document.status),
        document_owner_id=document.created_by,
        site_id=document.site_id,
        site_company_id=document.company_id,
    )
    raw = await build_document_dependency_map(session, tenant_id=str(tenant.id), document=document)
    npa = [
        DocumentDependencyNpaBindingRead(
            binding_id=b["binding_id"],
            npa_id=b["npa_id"],
            npa_code=b["npa_code"],
            npa_title=b["npa_title"],
            ref=b.get("ref"),
            entity_type=b["entity_type"],
        )
        for b in raw.get("npa_bindings") or []
    ]
    return DocumentDependencyMapRead(
        template=raw.get("template"),
        template_version=raw.get("template_version"),
        npa_bindings=npa,
        pipeline_profile_hint=raw.get("pipeline_profile_hint"),
    )


@router.get("/{document_id}/status", response_model=DocumentUiRead)
async def get_document_status(
    document_id: str,
    tenant: Tenant = TenantDep,
    session: AsyncSession = SessionDep,
    access: AccessContext = ReadAccessDep,
) -> DocumentUiRead:
    return await get_document(document_id=document_id, tenant=tenant, session=session, access=access)


@router.get("/{document_id}/download")
async def download_document(
    document_id: str,
    tenant: Tenant = TenantDep,
    session: AsyncSession = SessionDep,
    access: AccessContext = ReadAccessDep,
    storage: FileStorageService = Depends(get_file_storage_service),
) -> Response:
    access.ensure_tenant_access(tenant.id, action="download document")
    document = (
        await session.execute(_document_read_query(str(tenant.id)).where(Document.id == document_id))
    ).scalar_one_or_none()
    if document is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Document not found")
    access.ensure_abac(
        action="download document",
        company_id=document.company_id,
        document_id=document.id,
        document_status=document.status.value if hasattr(document.status, "value") else str(document.status),
        document_owner_id=document.created_by,
        site_id=document.site_id,
        site_company_id=document.company_id,
    )

    latest_version = _resolve_latest_version(document)
    current_file = document.file or (latest_version.file if latest_version else None)
    storage_key = document.storage_key or (current_file.storage_key if current_file else None) or (latest_version.file_key if latest_version else None)
    if not storage_key:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Document file not found")

    try:
        payload = storage.download(storage_key)
    except KeyError as exc:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Document file not found") from exc

    filename = (current_file.original_name if current_file and current_file.original_name else f"document-{document.id}.bin").replace('"', "")
    media_type = current_file.mime if current_file else "application/octet-stream"
    return Response(
        content=payload,
        media_type=media_type,
        headers={"Content-Disposition": f'inline; filename="{filename}"'},
    )

def _serialize_payload(payload: dict[str, Any]) -> str:
    try:
        return json.dumps(payload, sort_keys=True, separators=(",", ":"))
    except (TypeError, ValueError) as exc:
        raise _documents_bad_request("data must be JSON serializable") from exc


def _hash_payload(payload: dict[str, Any]) -> str:
    serialized = _serialize_payload(payload)
    digest = hashlib.sha256(serialized.encode("utf-8"))
    return digest.hexdigest()


def _ensure_payload_size(payload: dict[str, Any], *, limit: int, field: str) -> None:
    serialized = _serialize_payload(payload)
    size = len(serialized.encode("utf-8"))
    if size > limit:
        raise HTTPException(
            status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            f"{field} payload cannot exceed {limit} bytes",
        )


def _parse_csv_payload(file: UploadFile) -> list[dict[str, Any]]:
    raw = file.file.read()
    content = raw.decode("utf-8")
    reader = csv.DictReader(StringIO(content))
    return [dict(row) for row in reader if any(row.values())]


def _parse_xlsx_payload(file: UploadFile) -> list[dict[str, Any]]:
    from io import BytesIO

    from openpyxl import load_workbook

    raw = file.file.read()
    workbook = load_workbook(BytesIO(raw), read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    data_rows = []
    for row in rows[1:]:
        entry = {
            header: value for header, value in zip(headers, row) if header
        }
        if any(value is not None and value != "" for value in entry.values()):
            data_rows.append(entry)
    return data_rows


def _apply_naming_pattern(pattern: str | None, row: dict[str, Any], row_index: int) -> str | None:
    if not pattern:
        return None

    class _SafeDict(dict):
        def __missing__(self, key: str) -> str:
            return ""

    return pattern.format_map(_SafeDict(row_index=row_index, **row))

def _extract_branding_generation_metadata(payload: DocGenerateRequest) -> dict[str, Any] | None:
    branding_preview = payload.data.get("branding_preview") if isinstance(payload.data, dict) else None
    if not isinstance(branding_preview, dict):
        return None
    data = branding_preview.get("data") if isinstance(branding_preview.get("data"), dict) else {}
    doc = data.get("doc") if isinstance(data.get("doc"), dict) else {}
    reproducibility = payload.data.get("reproducibility")
    if not isinstance(reproducibility, dict):
        reproducibility = data.get("reproducibility") if isinstance(data.get("reproducibility"), dict) else {}
    return {
        "site_id": payload.data.get("siteId") or data.get("branch", {}).get("id"),
        "preset_code": branding_preview.get("preset_code") or payload.data.get("headerPreset"),
        "document_title": doc.get("title") or payload.template_code,
        "document_number": doc.get("number"),
        "reproducibility": reproducibility,
        "apply_headers_payload": branding_preview,
    }


def _extract_document_version_id(run: PipelineRun) -> str | None:
    metadata = run.result_metadata or {}
    outputs = run.outputs or {}
    return metadata.get("document_version_id") or outputs.get("document_version_id")


async def _fetch_template(
    session: AsyncSession,
    tenant: Tenant,
    *,
    template_code: str | None = None,
    template_id: str | None = None,
    template_version: int | None = None,
) -> tuple[Template, TemplateVersion]:
    tenant_scope = (str(tenant.id), tenant.slug)
    if not template_code:
        raise _documents_bad_request("template_code is required for template selection")
    if template_version is None:
        raise _documents_bad_request("template_version is required for template selection")
    filters: list[Any] = [
        Template.tenant_id.in_(tenant_scope),
        TemplateVersion.tenant_id.in_(tenant_scope),
        TemplateVersion.status == TemplateVersionStatus.ACTIVE,
        or_(Template.code == template_code, Template.name == template_code),
        TemplateVersion.version == template_version,
    ]

    stmt = (
        select(Template, TemplateVersion)
        .join(TemplateVersion, TemplateVersion.template_id == Template.id)
        .where(*filters)
        .limit(1)
    )
    row = (await session.execute(stmt)).first()
    if row is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Template not found")
    template, version = row
    if template_id and template.id != template_id:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            "template_id does not match template_code selection",
        )
    return template, version


async def _ensure_company(session: AsyncSession, tenant: Tenant, company_id: str) -> Company:
    company = await session.get(Company, company_id)
    if company is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Company not found")
    try:
        assert_tenant_row_matches_session(
            session,
            company,
            mismatch_event="api.documents.ensure_company.tenant_scope_mismatch",
            not_found_message="tenant_mismatch",
            expected_tenant_id=str(tenant.id),
        )
    except ValueError:
        raise HTTPException(
            status.HTTP_403_FORBIDDEN,
            "Tenant mismatch for company resource",
        ) from None
    return company


async def _ensure_person(
    session: AsyncSession, person_id: str | None, company: Company
) -> Person | None:
    if person_id is None:
        return None
    person = await session.get(Person, person_id)
    if person is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Person not found")
    try:
        assert_tenant_row_matches_session(
            session,
            person,
            mismatch_event="api.documents.ensure_person.tenant_scope_mismatch",
            not_found_message="tenant_mismatch",
            expected_tenant_id=str(company.tenant_id),
        )
    except ValueError:
        raise HTTPException(
            status.HTTP_403_FORBIDDEN,
            "Tenant mismatch for person resource",
        ) from None
    if person.company_id != company.id:
        raise _documents_bad_request("person_id does not belong to the provided company")
    return person


async def _resolve_run(
    session: AsyncSession,
    *,
    idempotency_key: str,
    tenant: Tenant,
    template: Template,
    template_version: TemplateVersion,
    company: Company,
    person: Person | None,
    payload_hash: str,
    current_user: User,
    context: dict[str, Any],
    correlation_id: str,
    npa_binding_id: str | None,
    visible_passport: bool,
) -> tuple[PipelineRun, bool]:
    stmt = select(PipelineRun).where(
        PipelineRun.tenant_id == tenant.id,
        PipelineRun.idempotency_key == idempotency_key,
    )
    existing = (await session.execute(stmt)).scalar_one_or_none()
    if existing is not None:
        metadata = existing.result_metadata or {}
        if existing.template_id != template.id:
            raise HTTPException(status.HTTP_409_CONFLICT, "Idempotency key already used")
        if existing.template_version_id != template_version.id:
            raise HTTPException(status.HTTP_409_CONFLICT, "Idempotency key already used")
        if metadata.get("company_id") != company.id:
            raise HTTPException(status.HTTP_409_CONFLICT, "Idempotency key already used")
        expected_person = person.id if person else None
        if metadata.get("person_id") != expected_person:
            raise HTTPException(status.HTTP_409_CONFLICT, "Idempotency key already used")
        if metadata.get("payload_hash") != payload_hash:
            raise HTTPException(status.HTTP_409_CONFLICT, "Idempotency key already used")
        if existing.context != context:
            raise HTTPException(status.HTTP_409_CONFLICT, "Idempotency key already used")
        if existing.status == PipelineRunStatus.ERROR:
            raise HTTPException(
                status.HTTP_409_CONFLICT,
                "Idempotency key refers to a failed generation task",
            )
        return existing, False

    run = PipelineRun(
        tenant_id=tenant.id,
        template_id=template.id,
        template_version_id=template_version.id,
        status=PipelineRunStatus.QUEUED,
        context=dict(context),
        idempotency_key=idempotency_key,
        result_metadata={
            "company_id": company.id,
            "person_id": person.id if person else None,
            "initiated_by": current_user.id,
            "payload_hash": payload_hash,
            "correlation_id": correlation_id,
            "npa_binding_id": npa_binding_id,
            "visible_passport": visible_passport,
        },
    )
    session.add(run)
    await session.flush()
    return run, True


@router.post(
    "/generate",
    response_model=TaskAcceptedResponse,
    status_code=status.HTTP_202_ACCEPTED,
)
@router.post(
    "/documents:generate",
    response_model=TaskAcceptedResponse,
    status_code=status.HTTP_202_ACCEPTED,
)
@limiter.limit(lambda: generate_per_tenant(), key_func=ip_tenant_key)
@audit_operation("generate", "document_task", id_attr="task_id")
async def generate_document(
    payload: DocGenerateRequest | DocGeneratePipelineRequest,
    request: Request,
    response: Response,
    idempotency_key: Annotated[str | None, Header(alias="Idempotency-Key")] = None,
    session: AsyncSession = SessionDep,
    tenant: Tenant = TenantDep,
    access: AccessContext = AccessDep,
) -> TaskAcceptedResponse:
    settings = get_settings()
    correlation_id = get_trace_id()
    await BillingService(session).assert_allowed(tenant, "documents.generate")
    engine_payload: dict[str, Any] | None = None
    if isinstance(payload, DocGeneratePipelineRequest):
        template_code = payload.template.code
        template_version = payload.template.version
        options: dict[str, Any] = {}
        if payload.pipeline and payload.pipeline.steps and payload.pipeline.steps.zip:
            options["zip"] = bool(payload.pipeline.steps.zip.get("enabled"))
        engine_payload = {
            "template_code": template_code,
            "template_version": template_version,
            "pipeline_profile_id": (payload.pipeline.profile_code if payload.pipeline else None),
            "input_source_id": payload.data.file_id if payload.data.type == "file" else None,
            "inline_data": payload.data.payload or {},
            "options": options,
        }
    try:
        payload_for_checks = payload.data if isinstance(payload, DocGenerateRequest) else (payload.data.payload or {})
        enforce_mapping_constraints(payload_for_checks, field="data")
        _ensure_payload_size(payload_for_checks, limit=settings.document_payload_max_bytes, field="data")
    except PayloadConstraintError as exc:
        raise HTTPException(exc.status_code, str(exc)) from exc

    normalized_key = normalize_idempotency_key(idempotency_key)
    idem_state = getattr(request.state, "idempotency", {})
    request_hash = idem_state.get("fingerprint")
    if request_hash is None:
        request_hash = compute_request_hash(payload)
        idem_state = {"key": normalized_key, "fingerprint": request_hash}
        request.state.idempotency = idem_state
    idempotency = IdempotencyService(
        session=session,
        tenant_id=str(tenant.id),
        endpoint="documents.generate",
    )
    record, created_record = await idempotency.acquire(
        key=normalized_key,
        request_hash=request_hash,
        method=request.method.upper(),
        path=request.url.path,
    )

    created_run = False
    try:
        if engine_payload is not None or (isinstance(payload, DocGenerateRequest) and (payload.inline_data is not None or payload.input_source_id is not None)):
            orchestrator = DocumentPipelineOrchestrator(session)
            if engine_payload is None:
                legacy_payload = payload
                engine_payload = {
                    "template_code": legacy_payload.template_code,
                    "template_version": legacy_payload.template_version,
                    "pipeline_profile_id": legacy_payload.pipeline_profile_id,
                    "input_source_id": legacy_payload.input_source_id,
                    "inline_data": legacy_payload.inline_data or {},
                    "options": legacy_payload.options or {},
                }
            request_hash = compute_request_hash({
                "tenant_id": str(tenant.id),
                "endpoint": "documents.generate",
                **engine_payload,
            })
            if record.request_hash and record.request_hash != request_hash:
                raise HTTPException(
                    status.HTTP_409_CONFLICT,
                    {
                        "code": "IDEMPOTENCY_MISMATCH",
                        "type": "idempotency",
                        "message": "Idempotency key cannot be reused with a different request payload",
                    },
                )
            record.request_hash = request_hash
            if created_record:
                job = await orchestrator.start_document_job(
                    tenant_id=str(tenant.id),
                    created_by=getattr(access.user, "id", None),
                    payload=engine_payload,
                    idempotency_key=normalized_key,
                    request_hash=request_hash,
                    correlation_id=correlation_id,
                )
                await orchestrator.run_job(job_id=job.id)
                body = {
                    "task_id": job.id,
                    "job_id": job.id,
                    "correlation_id": job.correlation_id or correlation_id,
                    "status_url": f"/api/v1/jobs/{job.id}",
                    "document_version_id": None,
                }
                await BillingService(session).add_usage(tenant_id=str(tenant.id), docs_generated=1)
                await idempotency.store_success(record, status_code=status.HTTP_202_ACCEPTED, body=body)
                await session.commit()
                response.status_code = status.HTTP_202_ACCEPTED
                return TaskAcceptedResponse(**body)
            await session.commit()
            return await idempotency.respond_from_store(record, model=TaskAcceptedResponse, response=response)

        if not isinstance(payload, DocGenerateRequest):
            raise _documents_bad_request("legacy mode requires old payload shape")

        payload_hash = _hash_payload(payload.data)

        template, template_version = await _fetch_template(
            session,
            tenant,
            template_code=payload.template_code,
            template_id=payload.template_id,
            template_version=payload.template_version,
        )
        company = await _ensure_company(session, tenant, payload.company_id)
        person = await _ensure_person(session, payload.person_id, company)
        current_user: User = access.user
        if access.role in {"client_admin", "client_user"}:
            access.ensure_company_access(company.id, action="generate documents")

        branding_generation = _extract_branding_generation_metadata(payload)

        run, created_run = await _resolve_run(
            session,
            idempotency_key=normalized_key,
            tenant=tenant,
            template=template,
            template_version=template_version,
            company=company,
            person=person,
            payload_hash=payload_hash,
            current_user=current_user,
            context=payload.data,
            correlation_id=correlation_id,
            npa_binding_id=payload.npa_binding_id,
            visible_passport=payload.visible_passport,
        )

        if branding_generation:
            metadata = dict(run.result_metadata or {})
            metadata["branding"] = branding_generation
            run.result_metadata = metadata

        if created_run:
            audit_service = AuditService(session)
            ip = request.client.host if request.client else "unknown"
            user_agent = request.headers.get("user-agent")
            await audit_service.log_event(
                tenant_id=str(tenant.id),
                action="render_start",
                object_type="pipeline_run",
                object_id=run.id,
                user_id=current_user.id,
                ip=ip,
                user_agent=user_agent,
                details={
                    "template_id": template.id,
                    "company_id": company.id,
                    "person_id": person.id if person else None,
                },
            )

        if not created_record:
            document_version_id = _extract_document_version_id(run)
            if document_version_id:
                await idempotency.update_document_version_id(
                    key=normalized_key,
                    document_version_id=document_version_id,
                )
            await session.commit()
            return await idempotency.respond_from_store(
                record, model=TaskAcceptedResponse, response=response
            )

        status_url = f"/api/v1/documents/tasks/{run.id}"
        metadata = run.result_metadata or {}
        result = TaskAcceptedResponse(
            task_id=run.id,
            correlation_id=correlation_id,
            status_url=status_url,
            document_version_id=metadata.get("document_version_id"),
        )
        await BillingService(session).add_usage(tenant_id=str(tenant.id), docs_generated=1)
        await idempotency.store_success(
            record,
            status_code=status.HTTP_202_ACCEPTED,
            body=result.model_dump(mode="json"),
        )
        await session.commit()
    except HTTPException as exc:
        await idempotency.store_failure(
            record,
            status_code=exc.status_code,
            detail={"detail": exc.detail},
        )
        await session.commit()
        raise
    except Exception as exc:
        logger.exception(
            "documents.generate.unexpected_failure",
            extra={"correlation_id": correlation_id},
        )
        await idempotency.store_failure(
            record,
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail={"detail": _generate_internal_error_problem()},
        )
        await session.commit()
        raise

    if created_run:
        try:
            trace_id = get_trace_id()
            _dispatch_celery_task(
                generate_document_task,
                args=[run.id],
                kwargs={"tenant_slug": tenant.slug},
                task_id=run.id,
                headers={"trace_id": trace_id},
            )
        except Exception as exc:  # pragma: no cover - defensive logging
            logger.exception("documents.generate.dispatch_failed", exc_info=exc)
            stored = await idempotency.get(key=normalized_key)
            if stored is not None:
                await idempotency.store_failure(
                    stored,
                    status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                    detail={"detail": _generate_internal_error_problem()},
                )
                await session.commit()
            raise HTTPException(
                status.HTTP_500_INTERNAL_SERVER_ERROR,
                _generate_internal_error_problem(),
            ) from exc

    response.status_code = status.HTTP_202_ACCEPTED
    return result


@router.post(
    "/batch",
    response_model=DocumentBatchRunRead,
    status_code=status.HTTP_202_ACCEPTED,
)
@limiter.limit(lambda: generate_per_tenant(), key_func=ip_tenant_key)
@audit_operation("batch_generate", "document_batch")
async def generate_document_batch(
    file: UploadFile,
    request: Request,
    response: Response,
    template_code: str | None = Form(default=None),
    template_id: str | None = Form(default=None),
    template_version: int | None = Form(default=None),
    company_id: str | None = Form(default=None),
    naming_pattern: str | None = Form(default=None),
    session: AsyncSession = SessionDep,
    tenant: Tenant = TenantDep,
    access: AccessContext = AccessDep,
) -> DocumentBatchRunRead:
    settings = get_settings()
    await BillingService(session).assert_allowed(tenant, "documents.generate")
    if template_version is None or not company_id or not template_code:
        raise _documents_bad_request("template_code, template_version and company_id required")

    filename = (file.filename or "").lower()
    if filename.endswith(".csv"):
        rows = _parse_csv_payload(file)
    elif filename.endswith(".xlsx"):
        rows = _parse_xlsx_payload(file)
    else:
        raise _documents_bad_request("Unsupported batch file type")

    if not rows:
        raise _documents_bad_request("Batch file is empty")
    if len(rows) > settings.document_batch_max_rows:
        raise HTTPException(
            status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            f"Batch cannot exceed {settings.document_batch_max_rows} documents",
        )

    template, template_version_row = await _fetch_template(
        session,
        tenant,
        template_code=template_code,
        template_id=template_id,
        template_version=template_version,
    )
    company = await _ensure_company(session, tenant, company_id)
    if access.role in {"client_admin", "client_user"}:
        access.ensure_company_access(company.id, action="generate documents")

    current_user: User = access.user
    batch = DocumentBatchRun(
        tenant_id=tenant.id,
        template_id=template.id,
        template_version_id=template_version_row.id,
        company_id=company.id,
        naming_pattern=naming_pattern,
        status=DocumentBatchStatus.RUNNING,
        total=len(rows),
        processed=0,
        succeeded=0,
        failed=0,
        created_by=current_user.id,
        started_at=datetime.now(tz=timezone.utc),
    )
    session.add(batch)
    await session.flush()

    items: list[DocumentBatchItem] = []
    for index, row in enumerate(rows, start=1):
        row_payload = {k: v for k, v in row.items() if k not in {"person_id"}}
        try:
            enforce_mapping_constraints(row_payload, field="data")
            _ensure_payload_size(
                row_payload,
                limit=settings.document_payload_max_bytes,
                field="data",
            )
        except PayloadConstraintError as exc:
            raise HTTPException(exc.status_code, str(exc)) from exc
        person_id = row.get("person_id") if isinstance(row, dict) else None
        person = await _ensure_person(session, person_id, company)
        output_name = _apply_naming_pattern(naming_pattern, row_payload, index)
        payload_hash = _hash_payload(row_payload)
        run = PipelineRun(
            tenant_id=tenant.id,
            template_id=template.id,
            template_version_id=template_version_row.id,
            status=PipelineRunStatus.QUEUED,
            context=dict(row_payload),
            idempotency_key=f"batch-{batch.id}-{index}",
            result_metadata={
                "company_id": company.id,
                "person_id": person.id if person else None,
                "initiated_by": current_user.id,
                "payload_hash": payload_hash,
                "batch_id": batch.id,
                "row_index": index,
                "output_name": output_name,
            },
        )
        session.add(run)
        await session.flush()
        item = DocumentBatchItem(
            tenant_id=tenant.id,
            batch_id=batch.id,
            row_index=index,
            payload=row_payload,
            person_id=person.id if person else None,
            output_name=output_name,
            pipeline_run_id=run.id,
            status=DocumentBatchItemStatus.PENDING,
        )
        session.add(item)
        items.append(item)

    await session.commit()

    trace_id = get_trace_id()
    for item in items:
        _dispatch_celery_task(
            generate_document_batch_item_task,
            args=[batch.id, item.id],
            kwargs={"tenant_slug": tenant.slug},
            headers={"trace_id": trace_id},
        )

    return DocumentBatchRunRead(
        id=batch.id,
        status=batch.status.value,
        total=batch.total,
        processed=batch.processed,
        succeeded=batch.succeeded,
        failed=batch.failed,
        items=[
            {
                "id": item.id,
                "row_index": item.row_index,
                "status": item.status.value,
                "output_name": item.output_name,
            }
            for item in items
        ],
    )


@router.get(
    "/batch/{batch_id}",
    response_model=DocumentBatchRunRead,
)
async def get_document_batch(
    batch_id: str,
    session: AsyncSession = SessionDep,
    tenant: Tenant = TenantDep,
    access: AccessContext = AccessDep,
) -> DocumentBatchRunRead:
    access.ensure_tenant_access(tenant.id, action="read document batch")
    batch = (
        await session.execute(
            select(DocumentBatchRun)
            .where(
                DocumentBatchRun.id == batch_id,
                DocumentBatchRun.tenant_id == str(tenant.id),
            )
            .options(selectinload(DocumentBatchRun.items))
        )
    ).scalar_one_or_none()
    if batch is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Batch not found")
    items = batch.items
    return DocumentBatchRunRead(
        id=batch.id,
        status=batch.status.value,
        total=batch.total,
        processed=batch.processed,
        succeeded=batch.succeeded,
        failed=batch.failed,
        items=[
            {
                "id": item.id,
                "row_index": item.row_index,
                "status": item.status.value,
                "error": item.error,
                "document_id": item.document_id,
                "document_version_id": item.document_version_id,
                "output_name": item.output_name,
            }
            for item in items
        ],
    )


@router.get("/tasks/{task_id}", response_model=TaskStatusResponse)
async def get_generation_task_status(
    task_id: str,
    session: AsyncSession = SessionDep,
    tenant: Tenant = TenantDep,
    access: AccessContext = AccessDep,
) -> TaskStatusResponse:
    stmt = select(PipelineRun).where(
        PipelineRun.id == task_id,
        PipelineRun.tenant_id == tenant.id,
    )
    run = (await session.execute(stmt)).scalar_one_or_none()
    if run is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Task not found")

    metadata = dict(run.result_metadata or {})
    outputs = dict(run.outputs or {})
    if outputs:
        metadata.setdefault("outputs", outputs)

    pipeline_status = (
        run.status.value if hasattr(run.status, "value") else str(run.status)
    )
    metadata.setdefault("pipeline_status", pipeline_status)

    document_id = metadata.get("document_id") or outputs.get("document_id")
    document_version_id = metadata.get("document_version_id") or outputs.get(
        "document_version_id"
    )

    return TaskStatusResponse(
        task_id=run.id,
        status=pipeline_status,
        document_id=document_id,
        document_version_id=document_version_id,
        error=run.error,
        metadata=metadata or None,
    )


@router.patch("/{document_id}/status", response_model=DocumentRead)
@audit_operation("change_status", "document")
async def update_document_status(
    document_id: str,
    payload: DocumentStatusUpdate,
    request: Request,
    session: AsyncSession = SessionDep,
    tenant: Tenant = TenantDep,
    access: AccessContext = StatusAccessDep,
) -> DocumentRead:
    access.ensure_tenant_access(tenant.id, action="manage document status")
    service = DocumentWorkflowService(session=session)
    document = await service.get_document(document_id=document_id, tenant_id=str(tenant.id))
    access.ensure_abac(
        action="manage document status",
        document_id=document.id,
        document_status=document.status.value if hasattr(document.status, "value") else str(document.status),
        document_owner_id=document.created_by,
        site_id=document.site_id,
        site_company_id=document.company_id,
    )
    actor_id = getattr(access.user, "id", None)
    ip = request.client.host if request.client else "unknown"
    user_agent = request.headers.get("user-agent")
    try:
        document = await service.change_status(
            document_id=document_id,
            tenant_id=str(tenant.id),
            new_status=payload.to,
            actor_id=actor_id,
            ip=ip,
            user_agent=user_agent,
        )
    except DocumentNotFoundError as exc:
        raise HTTPException(status.HTTP_404_NOT_FOUND, str(exc)) from exc
    except InvalidStatusTransitionError as exc:
        await session.commit()
        raise HTTPException(status.HTTP_409_CONFLICT, str(exc)) from exc

    await session.commit()
    await session.refresh(document)
    return DocumentRead.model_validate(document)
