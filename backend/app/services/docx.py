from __future__ import annotations

import base64
import re
from collections.abc import Iterable, Mapping, Sequence
from io import BytesIO
from typing import Any

from docx import Document
from docx.shared import Mm
from docx.text.paragraph import Paragraph
from docxtpl import DocxTemplate, InlineImage

__all__ = ["DocxService"]


def _iter_paragraphs(container: object) -> Iterable[Paragraph]:
    if hasattr(container, "paragraphs"):
        for paragraph in getattr(container, "paragraphs"):
            yield paragraph
    if hasattr(container, "tables"):
        for table in getattr(container, "tables"):
            for row in table.rows:
                for cell in row.cells:
                    yield from _iter_paragraphs(cell)


def _iter_section_paragraphs(section: object) -> Iterable[Paragraph]:
    yield from _iter_paragraphs(getattr(section, "header", None))
    yield from _iter_paragraphs(getattr(section, "footer", None))


def _iter_document_paragraphs(document: Any) -> Iterable[Paragraph]:
    yield from _iter_paragraphs(document)
    if hasattr(document, "sections"):
        for section in document.sections:
            yield from _iter_section_paragraphs(section)


def _get_paragraph_text(paragraph: Paragraph) -> str:
    if paragraph.runs:
        return "".join(run.text for run in paragraph.runs)
    return paragraph.text


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in list(paragraph.runs)[1:]:
            paragraph._element.remove(run._element)
    else:
        paragraph.text = text


class DocxService:
    """Utility helpers built on top of python-docx and docxtpl libraries."""

    @staticmethod
    def mass_replace(docx_bytes: bytes, replacements: Mapping[str, str]) -> bytes:
        doc = Document(BytesIO(docx_bytes))
        for paragraph in _iter_document_paragraphs(doc):
            original_text = _get_paragraph_text(paragraph)
            updated_text = original_text
            for key, value in replacements.items():
                if key in updated_text:
                    updated_text = updated_text.replace(key, value)
            if updated_text == original_text:
                continue
            _set_paragraph_text(paragraph, updated_text)
        out = BytesIO()
        doc.save(out)
        return out.getvalue()

    @staticmethod
    def set_headers_footers(
        docx_bytes: bytes,
        header_texts: str | Sequence[str | None] | None,
        footer_texts: str | Sequence[str | None] | None,
    ) -> bytes:
        doc = Document(BytesIO(docx_bytes))
        for idx, section in enumerate(doc.sections):
            header_text = DocxService._select_section_text(header_texts, idx)
            footer_text = DocxService._select_section_text(footer_texts, idx)
            if header_text is not None:
                header = section.header
                DocxService._prepare_section_container(header)
                header.add_paragraph(header_text)
            if footer_text is not None:
                footer = section.footer
                DocxService._prepare_section_container(footer)
                footer.add_paragraph(footer_text)
        out = BytesIO()
        doc.save(out)
        return out.getvalue()

    @staticmethod
    def _select_section_text(
        value: str | Sequence[str | None] | None, index: int
    ) -> str | None:
        if value is None:
            return None
        if isinstance(value, Sequence) and not isinstance(value, (str, bytes)):
            if not value:
                return None
            if index < len(value):
                return value[index]
            return value[-1]
        return value

    @staticmethod
    def _prepare_section_container(container: object) -> None:
        if container is None:
            return
        if hasattr(container, "is_linked_to_previous"):
            container.is_linked_to_previous = False
        if hasattr(container, "paragraphs") and hasattr(container, "_element"):
            for paragraph in list(container.paragraphs):
                container._element.remove(paragraph._element)
        if hasattr(container, "tables") and hasattr(container, "_element"):
            for table in list(container.tables):
                container._element.remove(table._element)

    @staticmethod
    def extract_placeholders(docx_bytes: bytes) -> set[str]:
        doc = Document(BytesIO(docx_bytes))
        text = "\n".join(_get_paragraph_text(p) for p in _iter_document_paragraphs(doc))
        return set(re.findall(r"\{\{\s*([a-zA-Z0-9_\.]+)\s*\}\}", text))

    @staticmethod
    def _prepare_context(tpl: DocxTemplate, context: dict[str, Any]) -> dict[str, Any]:
        def transform(value: Any) -> Any:
            if isinstance(value, dict):
                if value.get("_type") == "inline_image":
                    data = value.get("data")
                    width = value.get("width_mm")
                    height = value.get("height_mm")
                    if isinstance(data, str):
                        payload = base64.b64decode(data)
                    elif isinstance(data, (bytes, bytearray)):
                        payload = bytes(data)
                    else:
                        raise ValueError("inline image descriptor missing data")
                    stream = BytesIO(payload)
                    kwargs: dict[str, Any] = {}
                    if width is not None:
                        kwargs["width"] = Mm(float(width))
                    if height is not None:
                        kwargs["height"] = Mm(float(height))
                    return InlineImage(tpl, stream, **kwargs)
                return {key: transform(val) for key, val in value.items()}
            if isinstance(value, list):
                return [transform(item) for item in value]
            return value

        return {key: transform(val) for key, val in context.items()}

    @staticmethod
    def render_template(docx_bytes: bytes, context: dict) -> bytes:
        tpl = DocxTemplate(BytesIO(docx_bytes))
        prepared = DocxService._prepare_context(tpl, context)
        tpl.render(prepared)
        out = BytesIO()
        tpl.save(out)
        return out.getvalue()
