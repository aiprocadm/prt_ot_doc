"""Чистый домен импорта отчёта СОУТ (без sqlalchemy / без I/O БД).

Три формат-парсера (csv/xlsx/ФГИС xml) сходятся к одному нормализованному
промежуточному формату ParsedWorkplace; валидация и diff — формат-независимы.
RU-метки классов берём из соседнего print_form.py (один контур)."""
from __future__ import annotations

import csv
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from io import BytesIO, StringIO

from app.domains.sout.print_form import HARMFUL_CLASSES


class UnsupportedImportFormat(Exception):
    """Расширение файла не имеет парсера."""


# Прямая карта распознавания класса: нормализованный токен -> значение SoutClass.
_CLASS_ALIASES = {
    "1": "optimal", "optimal": "optimal", "оптимальный": "optimal",
    "2": "acceptable", "acceptable": "acceptable", "допустимый": "acceptable",
    "3.1": "harmful_3_1", "harmful_3_1": "harmful_3_1",
    "3.2": "harmful_3_2", "harmful_3_2": "harmful_3_2",
    "3.3": "harmful_3_3", "harmful_3_3": "harmful_3_3",
    "3.4": "harmful_3_4", "harmful_3_4": "harmful_3_4",
    "4": "dangerous", "dangerous": "dangerous", "опасный": "dangerous",
}


def parse_class_label(raw) -> tuple[str | None, str | None]:
    """(value, unparsed). value = значение SoutClass, либо unparsed = исходная строка."""
    if raw is None:
        return None, None
    token = str(raw).strip().lower()
    if not token:
        return None, None
    for prefix in ("подкласс", "класс", "subclass", "class"):
        if token.startswith(prefix):
            token = token[len(prefix):].strip()
            break
    token = token.replace(",", ".")
    if token in _CLASS_ALIASES:
        return _CLASS_ALIASES[token], None
    parts = token.split()
    if parts and parts[-1] in _CLASS_ALIASES:
        return _CLASS_ALIASES[parts[-1]], None
    return None, str(raw).strip()


@dataclass
class ParsedFactor:
    code: str | None
    name: str
    measured_class: str | None
    class_unparsed: str | None


@dataclass
class ParsedWorkplace:
    workplace_code: str
    position_name: str
    assessed_class: str | None
    class_unparsed: str | None
    factors: list[ParsedFactor] = field(default_factory=list)
    conflict: bool = False  # один код встретился с разными классом/должностью


@dataclass
class RowIssues:
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass
class DiffRow:
    workplace_code: str
    change: str  # new | changed | unchanged | removed
    parsed_class: str | None
    current_class: str | None


# --- column aliases (RU + EN), нижний регистр ---
_WP_CODE_KEYS = ("workplace_code", "код рм", "код", "номер рм", "рабочее место")
_WP_POS_KEYS = ("position_name", "должность", "профессия", "наименование должности")
_WP_CLASS_KEYS = ("assessed_class", "класс", "итоговый класс", "класс условий труда")
_F_CODE_KEYS = ("factor_code", "код фактора")
_F_NAME_KEYS = ("factor_name", "фактор", "наименование фактора", "вредный фактор")
_F_CLASS_KEYS = ("factor_class", "класс фактора")


def _pick(norm_row: dict, keys) -> object | None:
    for k in keys:
        if k in norm_row and norm_row[k] not in (None, ""):
            return norm_row[k]
    return None


def _rows_to_workplaces(dict_rows) -> list[ParsedWorkplace]:
    order: list[str] = []
    by_code: dict[str, ParsedWorkplace] = {}
    for row in dict_rows:
        norm = {(str(k).strip().lower() if k is not None else ""): v for k, v in row.items()}
        code = str(_pick(norm, _WP_CODE_KEYS) or "").strip()
        pos = str(_pick(norm, _WP_POS_KEYS) or "").strip()
        ac_val, ac_unparsed = parse_class_label(_pick(norm, _WP_CLASS_KEYS))
        if code not in by_code:
            by_code[code] = ParsedWorkplace(
                workplace_code=code, position_name=pos,
                assessed_class=ac_val, class_unparsed=ac_unparsed, factors=[],
            )
            order.append(code)
        else:
            wp = by_code[code]
            # повторная строка с тем же кодом: конфликт, если непустой класс/должность отличаются
            if pos and pos != wp.position_name:
                wp.conflict = True
            if ac_val is not None and wp.assessed_class is not None and ac_val != wp.assessed_class:
                wp.conflict = True
        wp = by_code[code]
        fname = _pick(norm, _F_NAME_KEYS)
        if fname:
            fc_val, fc_unparsed = parse_class_label(_pick(norm, _F_CLASS_KEYS))
            fcode = _pick(norm, _F_CODE_KEYS)
            wp.factors.append(ParsedFactor(
                code=(str(fcode).strip() if fcode else None),
                name=str(fname).strip(), measured_class=fc_val, class_unparsed=fc_unparsed,
            ))
    return [by_code[c] for c in order]


def parse_csv(content: bytes) -> list[ParsedWorkplace]:
    reader = csv.DictReader(StringIO(content.decode("utf-8-sig")))
    rows = [dict(r) for r in reader if any(v not in (None, "") for v in r.values())]
    return _rows_to_workplaces(rows)


def parse_xlsx(content: bytes) -> list[ParsedWorkplace]:
    from openpyxl import load_workbook  # ленивый импорт (тяжёлый), как в documents.py

    wb = load_workbook(BytesIO(content), read_only=True)
    raw = list(wb.active.iter_rows(values_only=True))
    if not raw:
        return []
    headers = [str(c).strip() if c is not None else "" for c in raw[0]]
    rows = []
    for r in raw[1:]:
        entry = {h: v for h, v in zip(headers, r) if h}
        if any(v not in (None, "") for v in entry.values()):
            rows.append(entry)
    return _rows_to_workplaces(rows)


def parse_fgis_xml(content: bytes) -> list[ParsedWorkplace]:
    """Прагматичный субсет ФГИС СОУТ. Схема-допущение:

    <sout><workplace code="РМ-01" position="Слесарь">
        <assessed_class>acceptable</assessed_class>
        <factors><factor code="4.50" name="Шум" class="harmful_3_1"/></factors>
    </workplace></sout>

    TODO: сверить с реальной выгрузкой ФГИС СОУТ (образца в репо нет)."""
    root = ET.fromstring(content)
    out: list[ParsedWorkplace] = []
    for wp_el in root.iter("workplace"):
        code = (wp_el.get("code") or wp_el.findtext("code") or "").strip()
        pos = (wp_el.get("position") or wp_el.findtext("position") or "").strip()
        ac_val, ac_unparsed = parse_class_label(wp_el.findtext("assessed_class"))
        factors = []
        for f_el in wp_el.iter("factor"):
            fname = (f_el.get("name") or f_el.findtext("name") or "").strip()
            if not fname:
                continue
            fc_val, fc_unparsed = parse_class_label(f_el.get("class") or f_el.findtext("class"))
            factors.append(ParsedFactor(
                code=(f_el.get("code") or None), name=fname,
                measured_class=fc_val, class_unparsed=fc_unparsed,
            ))
        out.append(ParsedWorkplace(
            workplace_code=code, position_name=pos,
            assessed_class=ac_val, class_unparsed=ac_unparsed, factors=factors,
        ))
    return out


def parse_report(content: bytes, filename: str) -> list[ParsedWorkplace]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return parse_csv(content)
    if name.endswith(".xlsx"):
        return parse_xlsx(content)
    if name.endswith(".xml"):
        return parse_fgis_xml(content)
    raise UnsupportedImportFormat(f"неподдерживаемый формат файла: {filename}")


def validate_parsed(workplaces) -> dict[int, RowIssues]:
    issues: dict[int, RowIssues] = {}
    for i, wp in enumerate(workplaces):
        ri = RowIssues()
        if not wp.workplace_code:
            ri.errors.append("пустой код рабочего места")
        if not wp.position_name:
            ri.errors.append("пустая должность")
        if wp.class_unparsed:
            ri.errors.append(f"не распознан класс условий труда: {wp.class_unparsed!r}")
        if wp.conflict:
            ri.errors.append("конфликт атрибутов РМ (один код, разные класс/должность)")
        if wp.assessed_class in ("optimal", "acceptable") and any(
            f.measured_class in HARMFUL_CLASSES for f in wp.factors
        ):
            ri.warnings.append("класс 1-2, но указан вредный фактор (3.1+)")
        for f in wp.factors:
            if f.class_unparsed:
                ri.warnings.append(f"не распознан класс фактора {f.name!r}: {f.class_unparsed!r}")
        issues[i] = ri
    return issues


def diff_campaign(parsed, existing) -> list[DiffRow]:
    """existing: list[(workplace_code, assessed_class_value)]."""
    existing_map = {code: cls for code, cls in existing}
    seen: set[str] = set()
    out: list[DiffRow] = []
    for wp in parsed:
        seen.add(wp.workplace_code)
        if wp.workplace_code not in existing_map:
            change = "new"
        elif existing_map[wp.workplace_code] != wp.assessed_class:
            change = "changed"
        else:
            change = "unchanged"
        out.append(DiffRow(wp.workplace_code, change, wp.assessed_class, existing_map.get(wp.workplace_code)))
    for code, cls in existing:
        if code not in seen:
            out.append(DiffRow(code, "removed", None, cls))
    return out
