"""Чистый сборщик печатных форм СОУТ: Карта СОУТ + Сводная ведомость.

Без sqlalchemy / без I/O — как domains/medical/print_form.py. Принимает уже
собранный снимок (*PrintData), возвращает байты DOCX. RU-метки классов/гарантий
и канонический порядок классов — локальные константы (модуль не зависит ни от
ORM, ни от фронта)."""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO

from docx import Document

# --- RU vocab + канонический порядок классов (локально для чистого модуля) ---
SOUT_CLASS_LABELS = {
    "optimal": "1 (оптимальный)",
    "acceptable": "2 (допустимый)",
    "harmful_3_1": "3.1",
    "harmful_3_2": "3.2",
    "harmful_3_3": "3.3",
    "harmful_3_4": "3.4",
    "dangerous": "4 (опасный)",
}
# Канонический порядок (метки заголовков + порядок статистики берут отсюда).
SOUT_CLASS_ORDER = (
    "optimal", "acceptable", "harmful_3_1", "harmful_3_2",
    "harmful_3_3", "harmful_3_4", "dangerous",
)
# Вредные/опасные классы (3.1 и выше) — для подсчёта «вредных факторов».
HARMFUL_CLASSES = frozenset(
    {"harmful_3_1", "harmful_3_2", "harmful_3_3", "harmful_3_4", "dangerous"}
)
GUARANTEE_KIND_LABELS = {
    "additional_leave": "Дополнительный отпуск",
    "extra_pay": "Повышенная оплата труда",
    "reduced_hours": "Сокращённая рабочая неделя",
    "milk": "Молоко / лечебно-профилактическое питание",
    "early_pension": "Досрочная пенсия",
    "medical_exam": "Обязательный медосмотр",
}


def class_label(code: str | None) -> str:
    if code is None:
        return "—"
    return SOUT_CLASS_LABELS.get(code, code)


def guarantee_label(code: str) -> str:
    return GUARANTEE_KIND_LABELS.get(code, code)


@dataclass
class SoutCardFactorRow:
    code: str | None
    name: str
    measured_class: str | None  # raw enum value


@dataclass
class SoutCardGuaranteeRow:
    kind: str  # raw enum value
    detail: str | None


@dataclass
class SoutCardPrintData:
    org_header: str
    generated_at: str | None
    expert_org_name: str | None
    report_number: str | None
    report_date: str | None
    workplace_code: str
    position_name: str
    assessed_class: str | None
    assessment_date: str | None
    next_assessment_date: str | None
    factors: list[SoutCardFactorRow] = field(default_factory=list)
    guarantees: list[SoutCardGuaranteeRow] = field(default_factory=list)


@dataclass
class SoutSummaryRow:
    workplace_code: str
    position_name: str
    assessed_class: str | None
    harmful_factor_count: int
    next_assessment_date: str | None


@dataclass
class SoutSummaryPrintData:
    org_header: str
    generated_at: str | None
    campaign_name: str
    expert_org_name: str | None
    report_number: str | None
    report_date: str | None
    rows: list[SoutSummaryRow] = field(default_factory=list)
    class_counts: list[tuple[str, int]] = field(default_factory=list)  # (label, count)


def harmful_factor_count(factors: list[SoutCardFactorRow]) -> int:
    """Число факторов с классом 3.1+ (по measured_class фактора, не итоговому классу РМ)."""
    return sum(1 for f in factors if f.measured_class in HARMFUL_CLASSES)


def class_counts(assessed_classes: list[str | None]) -> list[tuple[str, int]]:
    """Статистика РМ по классам в каноническом порядке (только классы с count > 0)."""
    tally: dict[str, int] = {}
    for code in assessed_classes:
        if code is None:
            continue
        tally[code] = tally.get(code, 0) + 1
    return [
        (SOUT_CLASS_LABELS.get(code, code), tally[code])
        for code in SOUT_CLASS_ORDER
        if code in tally
    ]


def _kv(doc, label: str, value: str | None) -> None:
    doc.add_paragraph(f"{label}: {value if value is not None else '—'}")


def build_sout_card_docx(data: SoutCardPrintData) -> bytes:
    doc = Document()
    doc.add_heading("Карта специальной оценки условий труда", level=0)
    _kv(doc, "Организация", data.org_header)
    _kv(doc, "Экспертная организация", data.expert_org_name)
    _kv(doc, "Отчёт №", data.report_number)
    _kv(doc, "Дата отчёта", data.report_date)
    _kv(doc, "Код рабочего места", data.workplace_code)
    _kv(doc, "Должность", data.position_name)
    _kv(doc, "Итоговый класс условий труда", class_label(data.assessed_class))
    _kv(doc, "Дата оценки", data.assessment_date)
    _kv(doc, "Дата следующей оценки", data.next_assessment_date)

    doc.add_heading("Вредные и (или) опасные производственные факторы", level=1)
    if data.factors:
        tbl = doc.add_table(rows=1, cols=4)
        tbl.style = "Table Grid"
        for i, title in enumerate(["№", "Код (29н)", "Наименование фактора", "Класс"]):
            tbl.rows[0].cells[i].text = title
        for idx, f in enumerate(data.factors, start=1):
            cells = tbl.add_row().cells
            cells[0].text = str(idx)
            cells[1].text = f.code or "—"
            cells[2].text = f.name
            cells[3].text = class_label(f.measured_class)
    else:
        doc.add_paragraph("—")

    doc.add_heading("Гарантии и компенсации", level=1)
    if data.guarantees:
        for g in data.guarantees:
            detail = f": {g.detail}" if g.detail else ""
            doc.add_paragraph(f"• {guarantee_label(g.kind)}{detail}")
    else:
        doc.add_paragraph("—")

    doc.add_heading("Комиссия по проведению СОУТ", level=1)
    doc.add_paragraph("Председатель комиссии: ____________________ / ____________________")
    doc.add_paragraph("Члены комиссии: ____________________ / ____________________")
    _kv(doc, "Дата формирования", data.generated_at)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_summary_sheet_docx(data: SoutSummaryPrintData) -> bytes:
    doc = Document()
    doc.add_heading("Сводная ведомость результатов СОУТ", level=0)
    _kv(doc, "Организация", data.org_header)
    _kv(doc, "Кампания", data.campaign_name)
    _kv(doc, "Экспертная организация", data.expert_org_name)
    _kv(doc, "Отчёт №", data.report_number)
    _kv(doc, "Дата отчёта", data.report_date)

    if data.rows:
        tbl = doc.add_table(rows=1, cols=6)
        tbl.style = "Table Grid"
        for i, title in enumerate(
            ["№", "Код РМ", "Должность", "Итоговый класс", "Вредных факторов", "След. оценка"]
        ):
            tbl.rows[0].cells[i].text = title
        for idx, r in enumerate(data.rows, start=1):
            cells = tbl.add_row().cells
            cells[0].text = str(idx)
            cells[1].text = r.workplace_code
            cells[2].text = r.position_name
            cells[3].text = class_label(r.assessed_class)
            cells[4].text = str(r.harmful_factor_count)
            cells[5].text = r.next_assessment_date or "—"
    else:
        doc.add_paragraph("—")

    doc.add_heading("Итоги по классам условий труда", level=1)
    if data.class_counts:
        for label, count in data.class_counts:
            doc.add_paragraph(f"Класс {label}: {count}")
    else:
        doc.add_paragraph("—")
    _kv(doc, "Дата формирования", data.generated_at)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
