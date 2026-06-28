"""Чистый домен декларации соответствия условий труда (ст. 11 ФЗ-426, форма 406н).

Без sqlalchemy / без I/O — как print_form.py. Определяет, какие РМ подлежат
декларированию (класс 1-2 И без вредных факторов), и собирает печатную форму.
RU-метки/подсчёт вредных факторов переиспользуются из print_form (один контур)."""
from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO

from docx import Document

from app.domains.sout.print_form import (
    SoutCardFactorRow,
    class_label,
    harmful_factor_count,
)

#: Классы условий труда, при которых РМ подлежит декларированию (1 и 2).
DECLARABLE_CLASSES = frozenset({"optimal", "acceptable"})


def evaluate_eligibility(
    assessed_class: str | None, factors: list[SoutCardFactorRow]
) -> tuple[bool, str | None]:
    """Подлежит ли РМ декларированию. Возвращает (eligible, ineligible_reason).

    Консервативно-юридический критерий: любой вредный фактор (класс 3.1+)
    дисквалифицирует РМ, даже если итоговый класс проставлен 1-2.
    """
    if assessed_class is None:
        return False, "класс условий труда не проставлен"
    if assessed_class not in DECLARABLE_CLASSES:
        return False, f"класс {class_label(assessed_class)} — вредные/опасные условия"
    if harmful_factor_count(factors) > 0:
        return False, "выявлен вредный фактор (класс 3.1+)"
    return True, None


@dataclass
class DeclarationRow:
    workplace_code: str
    position_name: str
    assessed_class: str | None
    headcount: str          # "1" если есть person_id, иначе "—"
    report_ref: str | None
    eligible: bool
    ineligible_reason: str | None


@dataclass
class DeclarationPrintData:
    org_header: str
    generated_at: str | None
    employer_inn: str | None = None
    employer_ogrn: str | None = None
    employer_address: str | None = None
    campaign_name: str = ""
    report_number: str | None = None
    report_date: str | None = None
    rows: list[DeclarationRow] = field(default_factory=list)  # только eligible идут в печать


def _kv(doc, label: str, value: str | None) -> None:
    doc.add_paragraph(f"{label}: {value if value else '____________________'}")


def build_declaration_docx(data: DeclarationPrintData) -> bytes:
    doc = Document()
    doc.add_heading(
        "Декларация соответствия условий труда государственным нормативным "
        "требованиям охраны труда",
        level=0,
    )
    doc.add_paragraph(f"Наименование работодателя: {data.org_header}")
    _kv(doc, "ИНН", data.employer_inn)
    _kv(doc, "ОГРН", data.employer_ogrn)
    _kv(doc, "Адрес", data.employer_address)

    basis = None
    if data.report_number or data.report_date:
        basis = " ".join(
            p for p in [
                f"№ {data.report_number}" if data.report_number else None,
                f"от {data.report_date}" if data.report_date else None,
            ] if p
        )
    doc.add_paragraph(f"Основание — отчёт о проведении СОУТ: {basis or '____________________'}")

    doc.add_heading("Рабочие места, подлежащие декларированию", level=1)
    if data.rows:
        tbl = doc.add_table(rows=1, cols=5)
        tbl.style = "Table Grid"
        for i, title in enumerate(
            ["№", "Код РМ", "Должность", "Численность", "Реквизиты отчёта СОУТ"]
        ):
            tbl.rows[0].cells[i].text = title
        for idx, r in enumerate(data.rows, start=1):
            cells = tbl.add_row().cells
            cells[0].text = str(idx)
            cells[1].text = r.workplace_code
            cells[2].text = r.position_name
            cells[3].text = r.headcount
            cells[4].text = r.report_ref or "—"
    else:
        doc.add_paragraph("Нет рабочих мест, подлежащих декларированию.")

    # Срок действия декларации регулируется ст. 11 ФЗ-426.
    # TODO (юр.): точную формулировку срока/условий прекращения сверить с юристом.
    doc.add_paragraph(
        "Декларация подаётся в территориальный орган федерального органа "
        "исполнительной власти (Роструд). Срок действия — согласно ст. 11 ФЗ-426."
    )
    doc.add_paragraph("Руководитель: ____________________ / ____________________")
    _kv(doc, "Дата формирования", data.generated_at)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
