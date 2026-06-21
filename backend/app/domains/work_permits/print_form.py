"""Ф4: чистый сборщик печатного бланка наряда-допуска 782н (python-docx).

Без sqlalchemy / без I/O — как domains/work_permits/lifecycle.py. Принимает уже
собранный снимок (WorkPermitPrintData), возвращает байты DOCX. RU-метки —
локальные словари (модуль не зависит от фронта)."""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from docx import Document

# --- RU vocab (зеркало lib/workPermitVocab.ts, но локально для чистого модуля) ---
WORK_TYPE_LABELS = {
    "hot_work": "Огневые работы", "gas_hazardous": "Газоопасные работы",
    "height": "Работа на высоте", "confined_space": "Замкнутые пространства",
    "excavation": "Земляные работы", "electrical": "Электроустановки",
}
MEMBER_ROLE_LABELS = {
    "issuer": "Выдающий наряд", "supervisor": "Ответственный руководитель",
    "admitter": "Допускающий", "foreman": "Производитель работ",
    "observer": "Наблюдающий", "member": "Член бригады",
}
SAFETY_SYSTEM_LABELS = {
    "restraint": "Удерживающие системы", "positioning": "Системы позиционирования",
    "fall_arrest": "Страховочные системы", "rescue_evacuation": "Системы для эвакуации и спасения",
    "access": "Системы для подъёма и спуска",
}
STATUS_LABELS = {
    "draft": "Черновик", "issued": "Выдан", "suspended": "Приостановлен",
    "closed": "Закрыт", "cancelled": "Отменён",
}
SIGN_GROUP_LABELS = {"permit": "Наряд", "briefing": "Целевой инструктаж", "closing": "Закрытие"}
SIGN_MODE_LABELS = {"attested": "При оформителе", "code": "По коду"}


def work_type_label(code: str) -> str:
    return WORK_TYPE_LABELS.get(code, code)


def member_role_label(code: str) -> str:
    return MEMBER_ROLE_LABELS.get(code, code)


def safety_system_label(code: str) -> str:
    return SAFETY_SYSTEM_LABELS.get(code, code)


def status_label(code: str) -> str:
    return STATUS_LABELS.get(code, code)


@dataclass
class SignatureLine:
    group: str          # "permit" | "briefing" | "closing"
    role_label: str
    fio: str
    status_label: str
    signed_at: str | None
    mode: str           # "attested" | "code" | "—"
    hash_short: str


@dataclass
class WorkPermitPrintData:
    number: str
    work_type_label: str
    status_label: str
    org_header: str
    subdivision: str | None
    planned_start: str | None
    planned_end: str | None
    zone_text: str
    content_text: str | None
    conditions_text: str | None
    equipment_text: str | None
    hazards_text: str | None
    safety_systems_labels: list[str]
    measures_before: str | None
    measures_during: str | None
    special_conditions: str | None
    ppe_text: str | None
    members: list[tuple[str, str]]          # (role_label, fio)
    briefing: dict | None                    # {conducted_by_fio, conducted_at, topics}
    daily_admissions: list[dict]             # [{date, start, end, admitted_by_fio}]
    extensions: list[dict]                   # [{old_end, new_end, at}]
    completion: dict | None                  # {text, recorded_at}
    closed_at: str | None
    signatures: list[SignatureLine]


def _kv(doc, label: str, value: str | None) -> None:
    doc.add_paragraph(f"{label}: {value if value is not None else '—'}")


def build_work_permit_docx(data: WorkPermitPrintData) -> bytes:
    doc = Document()
    doc.add_heading(
        f"НАРЯД-ДОПУСК на производство работ повышенной опасности "
        f"({data.work_type_label}, Приказ Минтруда № 782н)", level=0,
    )

    # 1. Шапка
    _kv(doc, "Организация", data.org_header)
    _kv(doc, "Подразделение", data.subdivision)
    _kv(doc, "Номер наряда", data.number)
    _kv(doc, "Вид работ", data.work_type_label)
    _kv(doc, "Статус", data.status_label)
    _kv(doc, "Плановое начало", data.planned_start)
    _kv(doc, "Плановое окончание", data.planned_end)

    # 2. Ответственные и бригада
    doc.add_heading("Ответственные лица и состав бригады", level=1)
    if data.members:
        tbl = doc.add_table(rows=1, cols=2)
        tbl.style = "Table Grid"
        hdr = tbl.rows[0].cells
        hdr[0].text, hdr[1].text = "Роль", "Ф.И.О."
        for role_label, fio in data.members:
            cells = tbl.add_row().cells
            cells[0].text, cells[1].text = role_label, fio
    else:
        doc.add_paragraph("—")

    # 3. Описание работ
    doc.add_heading("Описание работ", level=1)
    _kv(doc, "Место (зона)", data.zone_text)
    _kv(doc, "Содержание работ", data.content_text)
    _kv(doc, "Условия проведения", data.conditions_text)
    _kv(doc, "Оборудование", data.equipment_text)
    _kv(doc, "Опасные факторы", data.hazards_text)
    _kv(doc, "Системы обеспечения безопасности",
        ", ".join(data.safety_systems_labels) if data.safety_systems_labels else None)
    _kv(doc, "Мероприятия до начала работ", data.measures_before)
    _kv(doc, "Мероприятия в процессе работ", data.measures_during)
    _kv(doc, "Особые условия", data.special_conditions)
    _kv(doc, "Средства индивидуальной защиты", data.ppe_text)

    # 4. Целевой инструктаж
    doc.add_heading("Целевой инструктаж", level=1)
    if data.briefing:
        _kv(doc, "Провёл", data.briefing.get("conducted_by_fio"))
        _kv(doc, "Дата", data.briefing.get("conducted_at"))
        _kv(doc, "Темы", data.briefing.get("topics"))
    else:
        doc.add_paragraph("—")

    # 5. Ежедневный допуск
    doc.add_heading("Ежедневный допуск к работе", level=1)
    if data.daily_admissions:
        tbl = doc.add_table(rows=1, cols=4)
        tbl.style = "Table Grid"
        hdr = tbl.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Дата", "Начало", "Окончание", "Допустил"
        for a in data.daily_admissions:
            cells = tbl.add_row().cells
            cells[0].text = str(a.get("date") or "")
            cells[1].text = str(a.get("start") or "")
            cells[2].text = str(a.get("end") or "")
            cells[3].text = str(a.get("admitted_by_fio") or "")
    else:
        doc.add_paragraph("—")

    # 6. Продления
    doc.add_heading("Продления наряда", level=1)
    if data.extensions:
        for e in data.extensions:
            doc.add_paragraph(
                f"Продлён с {e.get('old_end') or '—'} до {e.get('new_end') or '—'} "
                f"({e.get('at') or '—'})"
            )
    else:
        doc.add_paragraph("—")

    # 7. Окончание и закрытие
    doc.add_heading("Окончание работ и закрытие наряда", level=1)
    if data.completion:
        _kv(doc, "Акт окончания работ", data.completion.get("text"))
        _kv(doc, "Дата оформления", data.completion.get("recorded_at"))
    else:
        doc.add_paragraph("Акт окончания не оформлен")
    _kv(doc, "Наряд закрыт", data.closed_at)

    # 8. Подписи ПЭП
    doc.add_heading("Подписи (простая электронная подпись)", level=1)
    if data.signatures:
        tbl = doc.add_table(rows=1, cols=6)
        tbl.style = "Table Grid"
        hdr = tbl.rows[0].cells
        for i, title in enumerate(["Раздел", "Ф.И.О.", "Статус", "Дата", "Режим", "Хэш ПЭП"]):
            hdr[i].text = title
        for s in data.signatures:
            cells = tbl.add_row().cells
            cells[0].text = f"{SIGN_GROUP_LABELS.get(s.group, s.group)} · {s.role_label}"
            cells[1].text = s.fio
            cells[2].text = s.status_label
            cells[3].text = s.signed_at or "—"
            cells[4].text = SIGN_MODE_LABELS.get(s.mode, s.mode)
            cells[5].text = s.hash_short
    else:
        doc.add_paragraph("Подписи отсутствуют")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
