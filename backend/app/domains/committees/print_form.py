"""Срез-4: чистый сборщик печатной формы протокола заседания (python-docx).

Без sqlalchemy / без I/O — как domains/work_permits/print_form.py. Принимает
уже собранный снимок (CommitteeProtocolPrintData), возвращает байты DOCX.
RU-метки — локальные словари (модуль не зависит от фронта).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO

from docx import Document

KIND_LABELS = {
    "osms": "Комитет по охране труда",
    "pb": "Комитет по промышленной безопасности",
    "commission_training": "Комиссия по обучению",
    "commission_investigation": "Комиссия по расследованию",
    "other": "Комиссия",
}
OUTCOME_LABELS = {"carried": "Принято", "rejected": "Отклонено"}
TASK_STATUS_LABELS = {"open": "Открыта", "in_progress": "В работе", "done": "Выполнена"}


def kind_label(code: str) -> str:
    return KIND_LABELS.get(code, code)


def outcome_label(code: str | None) -> str:
    if code is None:
        return "Без голосования"
    return OUTCOME_LABELS.get(code, code)


def task_status_label(code: str) -> str:
    return TASK_STATUS_LABELS.get(code, code)


def quorum_label(
    quorum_met: bool | None,
    present_count: int | None,
    members_total: int | None,
    threshold_pct: int | None,
) -> str:
    if quorum_met is None or present_count is None or members_total is None:
        return "Кворум: —"
    verdict = "есть" if quorum_met else "нет"
    rule = "простое большинство" if threshold_pct is None else f"порог {threshold_pct}%"
    return f"Кворум: {verdict} ({present_count} из {members_total}, {rule})"


@dataclass
class DecisionBlock:
    seq: int
    agenda_title: str | None
    text: str
    votes_for: int
    votes_against: int
    votes_abstain: int
    outcome_label: str
    #: (исполнитель ФИО, срок, статус-метка)
    tasks: list[tuple[str, str | None, str]] = field(default_factory=list)


@dataclass
class CommitteeProtocolPrintData:
    committee_name: str
    kind_label: str
    protocol_no: str
    held_at: str | None
    location: str | None
    quorum_text: str
    #: (ФИО, присутствовал)
    attendance: list[tuple[str, bool]]
    invited: list[str]
    agenda: list[str]
    decisions: list[DecisionBlock]


def _kv(doc, label: str, value: str | None) -> None:
    doc.add_paragraph(f"{label}: {value if value else '—'}")


def build_protocol_docx(data: CommitteeProtocolPrintData) -> bytes:
    doc = Document()
    doc.add_heading(f"ПРОТОКОЛ № {data.protocol_no}", level=0)
    doc.add_paragraph(f"{data.committee_name} ({data.kind_label})")

    _kv(doc, "Дата и время проведения", data.held_at)
    _kv(doc, "Место проведения", data.location)
    doc.add_paragraph(data.quorum_text)

    doc.add_heading("Участники", level=1)
    if data.attendance:
        tbl = doc.add_table(rows=1, cols=2)
        tbl.style = "Table Grid"
        hdr = tbl.rows[0].cells
        hdr[0].text, hdr[1].text = "Ф.И.О.", "Присутствие"
        for fio, present in data.attendance:
            cells = tbl.add_row().cells
            cells[0].text = fio
            cells[1].text = "присутствовал(а)" if present else "отсутствовал(а)"
    else:
        doc.add_paragraph("—")

    doc.add_heading("Приглашённые", level=1)
    if data.invited:
        for fio in data.invited:
            doc.add_paragraph(fio, style="List Bullet")
    else:
        doc.add_paragraph("—")

    doc.add_heading("Повестка", level=1)
    if data.agenda:
        for i, title in enumerate(data.agenda, start=1):
            doc.add_paragraph(f"{i}. {title}")
    else:
        doc.add_paragraph("—")

    doc.add_heading("Решения", level=1)
    if data.decisions:
        for block in data.decisions:
            doc.add_heading(f"Решение {block.seq}", level=2)
            if block.agenda_title:
                _kv(doc, "Вопрос повестки", block.agenda_title)
            doc.add_paragraph(block.text)
            doc.add_paragraph(
                f"Итоги голосования — За: {block.votes_for}, "
                f"Против: {block.votes_against}, Воздержались: {block.votes_abstain}. "
                f"Итог: {block.outcome_label}."
            )
            if block.tasks:
                tbl = doc.add_table(rows=1, cols=3)
                tbl.style = "Table Grid"
                hdr = tbl.rows[0].cells
                hdr[0].text, hdr[1].text, hdr[2].text = "Исполнитель", "Срок", "Статус"
                for fio, due, status_text in block.tasks:
                    cells = tbl.add_row().cells
                    cells[0].text = fio
                    cells[1].text = due or "—"
                    cells[2].text = status_text
    else:
        doc.add_paragraph("—")

    doc.add_paragraph("")
    doc.add_paragraph("Председатель: ______________________")
    doc.add_paragraph("Секретарь: ______________________")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
