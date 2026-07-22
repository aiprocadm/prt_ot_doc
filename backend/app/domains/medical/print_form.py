"""Чистый сборщик печатных форм приказа 29н: «контингент» + «поименный список».

Без sqlalchemy / без I/O — как domains/work_permits/print_form.py. Принимает уже
собранный снимок (*PrintData), возвращает байты DOCX. RU-метки — локальные словари
(модуль не зависит ни от ORM, ни от фронта)."""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO

from docx import Document

# --- RU vocab (локально для чистого модуля) ---
EXAM_KIND_LABELS = {
    "periodic": "Периодический",
    "preliminary": "Предварительный",
    "psychiatric": "Психиатрическое освидетельствование",
    "fluorography": "Флюорография",
    "health_book": "Личная медицинская книжка",
}
CONTINGENT_STATUS_LABELS = {
    "ok": "В норме",
    "due_soon": "Истекает",
    "overdue": "Просрочен",
    "missing": "Отсутствует",
}


def exam_kind_label(code: str) -> str:
    return EXAM_KIND_LABELS.get(code, code)


def contingent_status_label(code: str) -> str:
    return CONTINGENT_STATUS_LABELS.get(code, code)


def _factors_text(factors: list[tuple[str, str]]) -> str:
    return "; ".join(f"{code} {name}".strip() for code, name in factors) or "—"


def _kinds_text(kinds: list[str]) -> str:
    return ", ".join(exam_kind_label(k) for k in kinds) or "—"


@dataclass
class RegisterPrintRow:
    position_name: str
    factors: list[tuple[str, str]]  # (code, name)
    headcount: int
    exam_kinds: list[str]  # raw kind codes; метки ставит сборщик
    periodicity_months: int | None


@dataclass
class RegisterPrintData:
    org_header: str
    generated_at: str | None
    rows: list[RegisterPrintRow] = field(default_factory=list)


@dataclass
class NamedListPrintRow:
    full_name: str
    position_name: str | None
    department: str | None
    factors: list[tuple[str, str]]
    required_kinds: list[str]  # raw kind codes
    last_exam_date: str | None
    next_due_date: str | None
    status: str  # raw status code; метку ставит сборщик


@dataclass
class NamedListPrintData:
    org_header: str
    generated_at: str | None
    rows: list[NamedListPrintRow] = field(default_factory=list)


def _kv(doc, label: str, value: str | None) -> None:
    doc.add_paragraph(f"{label}: {value if value is not None else '—'}")


def build_contingent_register_docx(data: RegisterPrintData) -> bytes:
    doc = Document()
    doc.add_heading(
        "Контингент работников, подлежащих медицинским осмотрам (приказ № 29н)",
        level=0,
    )
    _kv(doc, "Организация", data.org_header)
    _kv(doc, "Дата формирования", data.generated_at)

    if data.rows:
        tbl = doc.add_table(rows=1, cols=5)
        tbl.style = "Table Grid"
        for i, title in enumerate(
            [
                "Должность",
                "Вредные факторы (29н)",
                "Численность",
                "Виды осмотров",
                "Периодичность, мес.",
            ]
        ):
            tbl.rows[0].cells[i].text = title
        for r in data.rows:
            cells = tbl.add_row().cells
            cells[0].text = r.position_name
            cells[1].text = _factors_text(r.factors)
            cells[2].text = str(r.headcount)
            cells[3].text = _kinds_text(r.exam_kinds)
            cells[4].text = str(r.periodicity_months) if r.periodicity_months is not None else "—"
    else:
        doc.add_paragraph("—")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


@dataclass
class ReferralPrintData:
    org_header: str
    generated_at: str | None
    full_name: str
    birth_date: str | None
    position_name: str | None
    department: str | None
    exam_kind: str  # raw kind code; метку ставит сборщик
    medical_org_name: str | None
    due_date: str | None
    snils: str | None
    factors: list[tuple[str, str]] = field(default_factory=list)  # (code, name)


def build_referral_docx(data: ReferralPrintData) -> bytes:
    doc = Document()
    doc.add_heading("Направление на медицинский осмотр (приказ № 29н)", level=0)
    _kv(doc, "Организация", data.org_header)
    _kv(doc, "Дата формирования", data.generated_at)
    doc.add_paragraph("")
    _kv(doc, "Ф.И.О.", data.full_name)
    _kv(doc, "Дата рождения", data.birth_date)
    _kv(doc, "СНИЛС", data.snils)
    _kv(doc, "Должность", data.position_name)
    _kv(doc, "Подразделение", data.department)
    _kv(doc, "Вид осмотра", exam_kind_label(data.exam_kind))
    _kv(doc, "Медицинская организация", data.medical_org_name)
    _kv(doc, "Срок прохождения", data.due_date)
    _kv(doc, "Вредные факторы (29н)", _factors_text(data.factors))

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_named_list_docx(data: NamedListPrintData) -> bytes:
    doc = Document()
    doc.add_heading(
        "Поименный список работников, подлежащих медицинским осмотрам (приказ № 29н)",
        level=0,
    )
    _kv(doc, "Организация", data.org_header)
    _kv(doc, "Дата формирования", data.generated_at)

    if data.rows:
        tbl = doc.add_table(rows=1, cols=7)
        tbl.style = "Table Grid"
        for i, title in enumerate(
            [
                "Ф.И.О.",
                "Должность",
                "Подразделение",
                "Вредные факторы (29н)",
                "Виды осмотров",
                "Последний / следующий",
                "Статус",
            ]
        ):
            tbl.rows[0].cells[i].text = title
        for r in data.rows:
            cells = tbl.add_row().cells
            cells[0].text = r.full_name
            cells[1].text = r.position_name or "—"
            cells[2].text = r.department or "—"
            cells[3].text = _factors_text(r.factors)
            cells[4].text = _kinds_text(r.required_kinds)
            cells[5].text = f"{r.last_exam_date or '—'} / {r.next_due_date or '—'}"
            cells[6].text = contingent_status_label(r.status)
    else:
        doc.add_paragraph("—")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
