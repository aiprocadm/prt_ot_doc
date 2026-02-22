from __future__ import annotations

import re
from pathlib import Path


def normalize_text(text: str) -> str:
    cleaned = re.sub(r"\s+", " ", text or "").strip()
    return cleaned


def extract_text_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except Exception:
        return ""
    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return normalize_text("\n".join(parts))


def extract_text_docx(path: Path) -> str:
    try:
        from docx import Document
    except Exception:
        return ""
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return normalize_text("\n".join(parts))
