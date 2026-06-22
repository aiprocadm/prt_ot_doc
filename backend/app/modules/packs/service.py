from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import re
from datetime import datetime, timezone
from typing import Any

from fastapi import HTTPException, status
from openpyxl import load_workbook
from sqlalchemy import Select, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import api_problem_detail
from app.core.idempotency import compute_request_hash
from app.models.file import File
from app.models.models import (
    PackageEntityStatus,
    PackagePresetConfig,
    PackagePresetItem,
    PackageProfileConfig,
    PackRun,
    PackRunItem,
    PackRunItemStatus,
    PackRunLifecycleStatus,
    TemplateUsage,
    TemplateVersion,
    TemplateVersionStatus,
)
from app.modules.packs.schemas import PackagePresetItemCreate, PackRunCreate

_PACKAGE_SOURCE_FILE_NOT_FOUND = api_problem_detail(
    code="PACKAGE_SOURCE_FILE_NOT_FOUND",
    message="Source file not found",
    error_type="packs",
)


class NamingRuleEngine:
    TOKEN_RE = re.compile(r"<([a-zA-Z0-9_]+)>")
    BAD_CHARS_RE = re.compile(r"[\\/:*?\"<>|]+")
    TOKEN_ALIASES: dict[str, str] = {
        "yyyymmdd": "date",
    }
    ALLOWED_TOKENS: set[str] = {
        "org",
        "unit",
        "project",
        "client",
        "doc",
        "topic",
        "version",
        "date",
        "flags",
    }

    def _resolve_token_value(self, key: str, payload: dict[str, Any]) -> str:
        normalized = key.strip().lower()
        canonical = self.TOKEN_ALIASES.get(normalized, normalized)
        if canonical == "date" and canonical not in payload:
            return datetime.now(timezone.utc).strftime("%Y%m%d")

        for candidate in {canonical, key, key.lower(), key.upper()}:
            if candidate in payload and payload[candidate] is not None:
                return str(payload[candidate]).strip()
        return ""

    def render(self, rule: str, payload: dict[str, Any], ext: str = "docx") -> str:
        def _replace(match: re.Match[str]) -> str:
            return self._resolve_token_value(match.group(1), payload)

        value = self.TOKEN_RE.sub(_replace, rule)
        value = self.BAD_CHARS_RE.sub("_", value)
        value = re.sub(r"\s+", " ", value).strip(" ._")
        value = value[:180] if len(value) > 180 else value
        return f"{value or 'document'}.{ext}"

    def validate_rule(self, rule: str) -> list[str]:
        warnings: list[str] = []
        for token in self.TOKEN_RE.findall(rule):
            key = token.strip().lower()
            canonical = self.TOKEN_ALIASES.get(key, key)
            if canonical not in self.ALLOWED_TOKENS:
                warnings.append(f"unknown naming token <{token}>")
        return warnings

    def ensure_unique(self, rendered_filename: str, used_filenames: set[str]) -> str:
        if rendered_filename not in used_filenames:
            used_filenames.add(rendered_filename)
            return rendered_filename

        stem, dot, extension = rendered_filename.rpartition(".")
        if not dot:
            stem, extension = rendered_filename, ""

        suffix = 2
        while True:
            candidate_stem = f"{stem}_{suffix}"
            if extension:
                candidate = f"{candidate_stem}.{extension}"
            else:
                candidate = candidate_stem
            if candidate not in used_filenames:
                used_filenames.add(candidate)
                return candidate
            suffix += 1


class SourceImportService:
    def parse(self, source_type: str, content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
        if source_type == "json":
            rows = json.loads(content.decode("utf-8"))
            if not isinstance(rows, list):
                raise HTTPException(status.HTTP_400_BAD_REQUEST, "JSON source must be an array")
            normalized = [self._normalize_row(dict(row)) for row in rows]
            columns = sorted({k for row in normalized for k in row.keys()})
            return columns, normalized
        if source_type == "csv":
            text = content.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            rows = [self._normalize_row(dict(row)) for row in reader]
            return self._normalize_columns(reader.fieldnames or []), rows
        if source_type == "xlsx":
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            raw_header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if raw_header is None:
                return [], []

            columns = self._normalize_columns(list(raw_header))
            rows: list[dict[str, Any]] = []
            for values in sheet.iter_rows(min_row=2, values_only=True):
                row_payload = {columns[index]: values[index] for index in range(len(columns))}
                rows.append(self._normalize_row(row_payload))
            return columns, rows
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST, "source type is not supported in this build"
        )

    def _norm_col(self, column: str | None) -> str:
        return (column or "").strip().lower()

    def _normalize_row(self, row: dict[str, Any]) -> dict[str, Any]:
        normalized: dict[str, Any] = {}
        for key, value in row.items():
            normalized[self._norm_col(key)] = value
        return normalized

    def _normalize_columns(self, columns: list[Any]) -> list[str]:
        result: list[str] = []
        seen: dict[str, int] = {}
        for index, column in enumerate(columns, start=1):
            normalized = self._norm_col(str(column or "")) or f"column_{index}"
            counter = seen.get(normalized, 0)
            if counter > 0:
                normalized = f"{normalized}_{counter + 1}"
            seen[self._norm_col(str(column or "")) or f"column_{index}"] = counter + 1
            result.append(normalized)
        return result


class MappingValidationService:
    def validate(self, mapping: dict[str, Any], columns: list[str]) -> None:
        known = set(columns)
        for target, source in mapping.items():
            if isinstance(source, str) and source.strip().lower() not in known:
                raise HTTPException(
                    status.HTTP_400_BAD_REQUEST, f"missing source column for mapping {target}"
                )
            if isinstance(source, dict):
                if (
                    source.get("type") == "column"
                    and str(source.get("value", "")).strip().lower() not in known
                ):
                    raise HTTPException(
                        status.HTTP_400_BAD_REQUEST, f"missing source column for mapping {target}"
                    )
                if source.get("type") not in {"column", "literal"}:
                    raise HTTPException(
                        status.HTTP_400_BAD_REQUEST, "invalid computed mapping type"
                    )

    def apply(self, mapping: dict[str, Any], row: dict[str, Any]) -> dict[str, Any]:
        result: dict[str, Any] = {}
        for target, source in mapping.items():
            if isinstance(source, str):
                result[target] = row.get(source.strip().lower())
            elif isinstance(source, dict) and source.get("type") == "literal":
                result[target] = source.get("value")
            elif isinstance(source, dict) and source.get("type") == "column":
                result[target] = row.get(str(source.get("value", "")).strip().lower())
        return result


class PackageService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session

    async def create_profile(self, tenant_id: str, payload: dict[str, Any]) -> PackageProfileConfig:
        profile = PackageProfileConfig(tenant_id=tenant_id, **payload)
        self.session.add(profile)
        await self.session.flush()
        return profile

    async def list_profiles(self, tenant_id: str) -> list[PackageProfileConfig]:
        stmt: Select[tuple[PackageProfileConfig]] = select(PackageProfileConfig).where(
            PackageProfileConfig.tenant_id == tenant_id,
            PackageProfileConfig.deleted_at.is_(None),
        )
        return list((await self.session.execute(stmt)).scalars().all())

    async def get_profile(self, tenant_id: str, profile_id: str) -> PackageProfileConfig | None:
        stmt = select(PackageProfileConfig).where(
            PackageProfileConfig.tenant_id == tenant_id,
            PackageProfileConfig.id == profile_id,
            PackageProfileConfig.deleted_at.is_(None),
        )
        return (await self.session.execute(stmt)).scalar_one_or_none()

    async def create_preset(self, tenant_id: str, payload: dict[str, Any]) -> PackagePresetConfig:
        preset = PackagePresetConfig(tenant_id=tenant_id, **payload)
        self.session.add(preset)
        await self.session.flush()
        return preset

    async def get_preset(self, tenant_id: str, preset_id: str) -> PackagePresetConfig | None:
        stmt = select(PackagePresetConfig).where(
            PackagePresetConfig.tenant_id == tenant_id,
            PackagePresetConfig.id == preset_id,
            PackagePresetConfig.deleted_at.is_(None),
        )
        return (await self.session.execute(stmt)).scalar_one_or_none()

    async def list_presets(self, tenant_id: str) -> list[PackagePresetConfig]:
        stmt = select(PackagePresetConfig).where(
            PackagePresetConfig.tenant_id == tenant_id,
            PackagePresetConfig.deleted_at.is_(None),
        )
        return list((await self.session.execute(stmt)).scalars().all())

    async def add_item(
        self, tenant_id: str, preset: PackagePresetConfig, payload: PackagePresetItemCreate
    ) -> PackagePresetItem:
        tv = await self.session.get(TemplateVersion, payload.template_version_id)
        if tv is None or str(tv.tenant_id) != str(tenant_id):
            raise HTTPException(status.HTTP_404_NOT_FOUND, "template version not found")
        if tv.deleted_at is not None or tv.status in {
            TemplateVersionStatus.ARCHIVED,
            TemplateVersionStatus.DEPRECATED,
        }:
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST, "template version is invalid for package"
            )
        replace_mode = payload.replace_mode
        if replace_mode == "dry-run":
            replace_mode = "preview"

        item = PackagePresetItem(
            tenant_id=tenant_id,
            package_preset_id=preset.id,
            template_id=payload.template_id or tv.template_id,
            template_version_id=payload.template_version_id,
            order_no=payload.order_no,
            header_preset_json=payload.header_preset_json,
            replace_mode=replace_mode,
            replace_map_json=payload.replace_map_json,
            output_format=payload.output_format,
            is_required=payload.is_required,
            conditions_json=payload.conditions_json,
        )
        self.session.add(item)
        self.session.add(
            TemplateUsage(
                tenant_id=tenant_id,
                template_version_id=payload.template_version_id,
                used_by_type="package_preset",
                used_by_id=preset.id,
            )
        )
        await self.session.flush()
        return item

    async def get_item(
        self, tenant_id: str, preset_id: str, item_id: str
    ) -> PackagePresetItem | None:
        stmt = select(PackagePresetItem).where(
            PackagePresetItem.tenant_id == tenant_id,
            PackagePresetItem.package_preset_id == preset_id,
            PackagePresetItem.id == item_id,
            PackagePresetItem.deleted_at.is_(None),
        )
        return (await self.session.execute(stmt)).scalar_one_or_none()


class PackRunService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        self.naming = NamingRuleEngine()
        self.mapping = MappingValidationService()

    async def create_run(
        self,
        *,
        tenant_id: str,
        payload: PackRunCreate,
        idempotency_key: str,
        rows: list[dict[str, Any]],
    ) -> PackRun:
        preset = await self.session.get(PackagePresetConfig, payload.package_preset_id)
        if (
            preset is None
            or str(preset.tenant_id) != str(tenant_id)
            or preset.deleted_at is not None
        ):
            raise HTTPException(status.HTTP_404_NOT_FOUND, "package preset not found")
        if preset.status != PackageEntityStatus.ACTIVE and preset.status != "active":
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "package preset must be active")

        selected_rows = payload.selected_rows or list(range(1, len(rows) + 1))
        req_hash = compute_request_hash(
            {
                "package_preset_id": payload.package_preset_id,
                "selected_rows": selected_rows,
                "override_profile_id": payload.override_profile_id,
                "override_options": payload.override_options,
                "rows": rows,
            }
        )
        existing = (
            await self.session.execute(
                select(PackRun).where(
                    PackRun.tenant_id == tenant_id,
                    PackRun.idempotency_key == idempotency_key,
                    PackRun.request_hash == req_hash,
                )
            )
        ).scalar_one_or_none()
        if existing:
            return existing

        self.mapping.validate(preset.mapping_json or {}, list(rows[0].keys()) if rows else [])
        run = PackRun(
            tenant_id=tenant_id,
            package_preset_id=preset.id,
            package_profile_id=payload.override_profile_id or preset.package_profile_id,
            source_file_id=payload.source_file_id,
            source_type=preset.source_type,
            source_rows_count=len(rows),
            selected_rows_count=len(selected_rows),
            status=PackRunLifecycleStatus.SUCCESS,
            stats_json={"queued": len(selected_rows), "success": len(selected_rows), "failed": 0},
            started_at=datetime.now(timezone.utc),
            ended_at=datetime.now(timezone.utc),
            idempotency_key=idempotency_key,
            request_hash=req_hash,
        )
        self.session.add(run)
        await self.session.flush()
        used_filenames: set[str] = set()
        for row_no in selected_rows:
            if row_no < 1 or row_no > len(rows):
                continue
            mapped = self.mapping.apply(preset.mapping_json or {}, rows[row_no - 1])
            filename = self.naming.render(preset.naming_rule, mapped, ext="docx")
            filename = self.naming.ensure_unique(filename, used_filenames)
            self.session.add(
                PackRunItem(
                    tenant_id=tenant_id,
                    pack_run_id=run.id,
                    row_no=row_no,
                    source_record_hash=hashlib.sha1(  # nosec B324 - content-addressing hash, not security; FIPS-safe via usedforsecurity=False
                        json.dumps(rows[row_no - 1], sort_keys=True).encode("utf-8"),
                        usedforsecurity=False,
                    ).hexdigest(),
                    status=PackRunItemStatus.SUCCESS,
                    filename=filename,
                )
            )
        return run


async def ensure_template_version_deletable(
    session: AsyncSession, tenant_id: str, template_version_id: str
) -> None:
    usage = (
        await session.execute(
            select(TemplateUsage).where(
                TemplateUsage.tenant_id == tenant_id,
                TemplateUsage.template_version_id == template_version_id,
                TemplateUsage.used_by_type == "package_preset",
            )
        )
    ).scalar_one_or_none()
    if usage:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            detail=api_problem_detail(
                code="TEMPLATE_VERSION_USED_BY_PACKAGE_PRESET",
                message="Template version is used by a package preset",
                error_type="packs",
            ),
        )


async def load_source_rows(
    session: AsyncSession,
    source_file_id: str | None,
    inline_rows: list[dict[str, Any]],
    *,
    tenant_id: str | None = None,
) -> tuple[str, list[str], list[dict[str, Any]]]:
    if inline_rows:
        cols = sorted({k.strip().lower() for row in inline_rows for k in row.keys()})
        normalized = [{k.strip().lower(): v for k, v in row.items()} for row in inline_rows]
        return "json", cols, normalized
    if not source_file_id:
        return "json", [], []
    file = await session.get(File, source_file_id)
    if file is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail=_PACKAGE_SOURCE_FILE_NOT_FOUND)
    if tenant_id is not None and str(file.tenant_id) != str(tenant_id):
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail=_PACKAGE_SOURCE_FILE_NOT_FOUND)
    meta = file.meta_json or {}
    raw = meta.get("inline_content")
    raw_b64 = meta.get("inline_content_b64")
    if raw_b64:
        content = base64.b64decode(raw_b64)
    elif raw:
        content = str(raw).encode("utf-8")
    else:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail=api_problem_detail(
                code="PACKAGE_SOURCE_FILE_NO_INLINE_CONTENT",
                message="Source file has no inline content for preview",
                error_type="packs",
            ),
        )
    source_type = meta.get("source_type", "csv")
    columns, rows = SourceImportService().parse(source_type=source_type, content=content)
    return source_type, columns, rows
