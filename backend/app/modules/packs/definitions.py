"""Default document pack catalogue shipped with the service."""

from __future__ import annotations

from dataclasses import dataclass
from enum import Enum
from io import BytesIO
from typing import Callable, Sequence

from docx import Document

from app.core.disciplines import Discipline
from app.models.models import DocumentPackModule, DocumentPackScenario
from app.modules.packs.assets import (
    DEFAULT_LOGO_BYTES,
    DEFAULT_STAMP_BYTES,
    build_inline_image_descriptor,
)

__all__ = [
    "PackTemplateSpec",
    "PackDefinition",
    "PackScenario",
    "DEFAULT_PACKS",
    "PACK_DEFINITIONS_BY_CODE",
    "PACK_CODE_SITE_ACCESS",
    "PACK_CODE_NEW_COMPANY",
    "PACK_CODE_INCIDENT",
    "PACK_CODE_INSPECTION_PREP",
    "PACK_CODE_OPO",
    "PACK_CODE_WASTE",
    "PACK_CODE_CEO_SHIELD",
    "PACK_CODE_NEW_EMPLOYEE",
    "PACK_CODE_CONTRACTOR",
    "PACK_CODE_FIRE_INSPECTION",
    "PACK_CODE_CIVIL_DEFENCE",
    "PACK_CODE_BDD_REPORTS",
    "PACK_CODE_GOCHS_REPORTS",
    "PACK_CODE_ECO_REPORTS",
    "PACK_CODE_ROAD_SAFETY",
    "PACKS_WITHOUT_DISCIPLINE",
    "discipline_pack_coverage",
    "packs_for",
]


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

PACK_CODE_SITE_ACCESS = "OT_ENTER_SITE"
PACK_CODE_NEW_COMPANY = "NEW_COMPANY"
PACK_CODE_INCIDENT = "INCIDENT_RESPONSE"
PACK_CODE_INSPECTION_PREP = "OT_INSPECTION_PREP"
PACK_CODE_OPO = "OPO_COMPLIANCE"
PACK_CODE_WASTE = "ECO_WASTE"
PACK_CODE_CEO_SHIELD = "CEO_SHIELD"
# Срез-2: четыре сценария из «минимального набора для запуска» разд. 50.1,
# которых в каталоге не было. Коды новые, существующие НЕ переименованы:
# код пакета — естественный ключ, по нему уже заведены строки у арендаторов.
PACK_CODE_NEW_EMPLOYEE = "OT_NEW_EMPLOYEE"
PACK_CODE_CONTRACTOR = "OT_NEW_CONTRACTOR"
PACK_CODE_FIRE_INSPECTION = "PB_SITE_INSPECTION"
PACK_CODE_CIVIL_DEFENCE = "GOCHS_BASE"
# Срез-5: БДД была ЕДИНСТВЕННОЙ дисциплиной словаря без сценарного комплекта —
# в таблице разд. 50.1 её нет, а приёмка §58.3 требует комплект для КАЖДОЙ.
PACK_CODE_ROAD_SAFETY = "BDD_BASE"
# Срез-8 экологии (разд. 55.3): отчётность как формы фабрики. До него каталог
# держал только сценарные комплекты — ни одной отчётной формы.
PACK_CODE_ECO_REPORTS = "ECO_REPORTS"
# Срез-9 БДД (разд. 56.2, последний пункт раздела): «отчётность по БДД и
# документы (приказы, положения, планы мероприятий)». Приказ, инструкция и план
# мероприятий печатались базовым комплектом с первого среза; не хватало
# ОТЧЁТНОСТИ и ПОЛОЖЕНИЯ — а положение о системе управления БДД это как раз
# документ, с которого проверка начинается.
PACK_CODE_BDD_REPORTS = "BDD_REPORTS"
# Срез-6 ГО и ЧС (разд. 56.1, последний пункт раздела): «документы и
# отчётность: приказы, положения, инструкции, отчётность в органы управления
# ГОЧС». Приказ, план, состав комиссии, программа учений и журнал печатались
# базовым комплектом; не хватало ПОЛОЖЕНИЯ, ИНСТРУКЦИИ и самой ОТЧЁТНОСТИ.
PACK_CODE_GOCHS_REPORTS = "GOCHS_REPORTS"
# Срез-42 ПромБез (разд. 54.2 «документы и отчётность: ПЛА, положения,
# приказы, отчёты для Ростехнадзора»): ПЛА печатался базовым комплектом ОПО с
# первого среза; не хватало ПОЛОЖЕНИЯ о производственном контроле, ПРИКАЗА о
# назначении ответственного и самих ОТЧЁТОВ в Ростехнадзор.
PACK_CODE_OPO_REPORTS = "OPO_REPORTS"


class PackScenario(str, Enum):
    ENTER_SITE = "enter_site"
    INCIDENT = "incident"
    INSPECTION_PREP = "inspection_preparation"
    OPO = "hazardous_production_facility"
    WASTE = "waste_and_ecology"
    ECOLOGY_REPORTING = "ecology_reporting"
    CEO_SHIELD = "ceo_shield"
    NEW_COMPANY = "new_company"
    NEW_EMPLOYEE = "new_employee"
    CONTRACTOR = "contractor_onboarding"
    FIRE_INSPECTION = "fire_inspection"
    CIVIL_DEFENCE = "civil_defence"
    ROAD_SAFETY = "road_safety"
    ROAD_SAFETY_REPORTING = "road_safety_reporting"
    CIVIL_DEFENCE_REPORTING = "civil_defence_reporting"
    OPO_REPORTING = "opo_reporting"


@dataclass(slots=True, frozen=True)
class PackTemplateSpec:
    code: str
    name: str
    description: str
    category: str
    builder: Callable[[], bytes]


@dataclass(slots=True, frozen=True)
class PackDefinition:
    code: str
    name: str
    description: str
    scenario: PackScenario
    module: DocumentPackModule
    scenario_type: DocumentPackScenario
    templates: Sequence[PackTemplateSpec]
    item_order: Sequence[str]
    metadata: dict[str, str]
    #: Дисциплины комплекта КОДАМИ из общего словаря (BIZ-54-57 срез-5).
    #:
    #: Раньше дисциплина была только свободной строкой в ``metadata``
    #: («Охрана труда и промышленная безопасность», «Любая — по типу надзора»),
    #: и требование приёмки §58.3 «для каждой дисциплины есть комплект» нельзя
    #: было ни проверить, ни заметить пропажу. Строка осталась подписью для
    #: человека, а разметка кодами — для проверки.
    #:
    #: Пустой набор ДОПУСТИМ и означает «общая охрана труда»: отдельной
    #: дисциплины для неё в словаре нет (она разложена на медосмотры, СИЗ и
    #: обучение), и приписать комплект к любой из трёх ради заполненной клетки
    #: значило бы подменить дисциплину — тот же довод, что у нарядов-допусков
    #: в срезе-3. Такие комплекты названы в ``PACKS_WITHOUT_DISCIPLINE``.
    disciplines: tuple[Discipline, ...] = ()


def _doc(*paragraphs: str, header: str | None = None, footer: str | None = None) -> bytes:
    doc = Document()
    section = doc.sections[0]
    if header:
        header_para = section.header.paragraphs[0]
        header_para.text = header
    if footer:
        footer_para = section.footer.paragraphs[0]
        footer_para.text = footer
    for text in paragraphs:
        doc.add_paragraph(text)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _site_access_order() -> bytes:
    return _doc(
        "Приказ на выход на объект",
        "Компания: {{ company.name }} (ИНН {{ company.inn }})",
        "Адрес объекта: {{ site.address }}",
        "Ответственный по охране труда: {{ ot_responsible }}",
        "Ответственный по пожарной безопасности: {{ pb_responsible }}",
        "Виды работ: {{ work_types }}",
        "Опасные факторы: {{ hazards }}",
        header="{{ logo }}",
        footer="{{ stamp }}",
    )


def _site_access_briefing() -> bytes:
    return _doc(
        "Инструктаж перед выходом на объект",
        "Дата: {{ data.session_date }}",
        "Состав бригады: {{ data.team_members }}",
        "Ответственный по охране труда: {{ ot_responsible }}",
        "Ответственный по пожарной безопасности: {{ pb_responsible }}",
        "Перечень опасностей: {{ hazards }}",
        footer="Контроль: {{ stamp }}",
    )


def _site_access_checklist() -> bytes:
    return _doc(
        "Чек-лист допуска",
        "1. Обучение пройдено: {{ data.training_passed }}",
        "2. Медосмотр действителен: {{ data.medical_clearance }}",
        "3. СИЗ выданы: {{ data.ppe_ready }}",
        "4. Дополнительные отметки: {{ data.notes }}",
        footer="Подпись ответственного: {{ data.supervisor }}",
    )


def _site_access_ppe_card() -> bytes:
    return _doc(
        "Личная карточка учета СИЗ",
        "Работник: {{ person.full_name }}",
        "Должность: {{ position.name }}",
        "Выданные СИЗ: {{ data.issues }}",
        "Нормы по должности: {{ data.norms }}",
        footer="Ответственный: {{ data.ppe_officer }}",
    )


def _site_access_journal_entry() -> bytes:
    return _doc(
        "Лист журнала инструктажей",
        "Журнал: {{ journal.title }}",
        "Тип: {{ journal.journal_type }}",
        "Дата: {{ data.entry_date }}",
        "Инструктируемый: {{ data.person }}",
        "Ответственный: {{ data.instructor }}",
        footer="Подпись ответственного: {{ data.instructor }}",
    )


def _new_company_order() -> bytes:
    return _doc(
        "Приказ о создании службы охраны труда",
        "Организация: {{ company.name }}",
        "ИНН: {{ company.inn }}",
        "Адрес: {{ company.address }}",
        "Дата: {{ data.issued_at }}",
        "Назначить ответственным: {{ data.ot_lead }}",
        footer="Подпись директора: {{ stamp }}",
    )


def _new_company_instruction() -> bytes:
    return _doc(
        "Инструкция по охране труда",
        "Раздел 1. Общие положения",
        "{{ data.policy_intro }}",
        "Раздел 2. Обязанности работодателя",
        "{{ data.employer_duties }}",
        "Раздел 3. Обязанности работников",
        "{{ data.employee_duties }}",
        header="{{ logo }}",
    )


def _new_company_journal() -> bytes:
    return _doc(
        "Журнал регистрации инструктажей",
        "Организация: {{ company.name }}",
        "Ответственный: {{ data.ot_lead }}",
        "Начат: {{ data.journal_start }}",
        "Завершён: {{ data.journal_end }}",
        "Примечания: {{ data.journal_notes }}",
    )


def _new_company_policy() -> bytes:
    return _doc(
        "Политика в области охраны труда и пожарной безопасности",
        "Организация: {{ company.name }}",
        "Политика утверждена: {{ data.issued_at }}",
        "Ответственный за исполнение: {{ data.pb_lead }}",
        footer="{{ stamp }}",
    )


def _incident_report() -> bytes:
    return _doc(
        "Акт о несчастном случае",
        "Дата происшествия: {{ incident.date }}",
        "Место: {{ incident.location }}",
        "Пострадавший: {{ victim.name }} ({{ victim.position }})",
        "Описание события: {{ incident.description }}",
        "Свидетели: {{ witnesses }}",
        "Причины: {{ causes }}",
        "Мероприятия: {{ actions }}",
        footer="Комиссия: {{ commission }}",
    )


def _incident_commission_order() -> bytes:
    return _doc(
        "Приказ о создании комиссии",
        "Организация: {{ company.name }}",
        "Дата: {{ incident.date }}",
        "Состав комиссии: {{ commission }}",
        "Срок расследования: {{ incident.deadline }}",
        footer="Подпись руководителя: {{ stamp }}",
    )


def _incident_plan() -> bytes:
    return _doc(
        "План корректирующих мероприятий",
        "Организация: {{ company.name }}",
        "Инцидент: {{ incident.description }}",
        "Причины: {{ causes }}",
        "Мероприятия: {{ actions }}",
        "Контроль исполнения: {{ data.follow_up }}",
        header="{{ logo }}",
    )


def _inspection_plan() -> bytes:
    return _doc(
        "План подготовки к проверке",
        "Надзорный орган: {{ inspection.agency }}",
        "Дата визита: {{ inspection.date }}",
        "Ответственный: {{ inspection.lead }}",
        "Перечень документов: {{ inspection.documents }}",
        footer="{{ stamp }}",
    )


def _inspection_checklist() -> bytes:
    return _doc(
        "Чек-лист готовности к проверке",
        "Область проверки: {{ inspection.scope }}",
        "Ключевые риски: {{ inspection.risks }}",
        "Контакты на месте: {{ inspection.contacts }}",
        footer="{{ logo }}",
    )


def _inspection_briefing() -> bytes:
    return _doc(
        "Инструктаж команды перед проверкой",
        "Дата инструктажа: {{ data.briefing_date }}",
        "Участники: {{ data.team }}",
        "Роли и обязанности: {{ data.roles }}",
        "Согласованные действия: {{ data.actions }}",
        footer="{{ stamp }}",
    )


def _opo_passport() -> bytes:
    return _doc(
        "Паспорт ОПО",
        "Категория опасности: {{ opo.category }}",
        "Технологический процесс: {{ opo.process }}",
        "Опасные факторы: {{ hazards }}",
        "Ответственные лица: {{ opo.responsibles }}",
        header="{{ logo }}",
    )


def _opo_emergency_plan() -> bytes:
    return _doc(
        "План ликвидации аварий",
        "Контакты аварийных служб: {{ opo.contacts }}",
        "Сценарии аварий: {{ opo.incidents }}",
        "Ресурсы и средства: {{ opo.resources }}",
        "Порядок оповещения: {{ opo.alerts }}",
        footer="{{ stamp }}",
    )


def _opo_training_matrix() -> bytes:
    return _doc(
        "Матрица обучения ОПО",
        "Персонал: {{ opo.personnel }}",
        "Требуемые курсы: {{ opo.courses }}",
        "Допуски и удостоверения: {{ opo.permits }}",
        footer="{{ logo }}",
    )


def _waste_register() -> bytes:
    return _doc(
        "Реестр отходов",
        "Типы отходов: {{ waste.types }}",
        "Классы опасности: {{ waste.classes }}",
        "Места накопления: {{ waste.storage }}",
        "Учёт и вывоз: {{ waste.removal }}",
        header="{{ logo }}",
    )


def _waste_contract() -> bytes:
    return _doc(
        "Договор на обращение с отходами",
        "Подрядчик: {{ waste.contractor }}",
        "Номер договора: {{ waste.contract_number }}",
        "Срок действия: {{ waste.contract_validity }}",
        "Ответственные: {{ waste.responsibles }}",
        footer="{{ stamp }}",
    )


def _waste_briefing() -> bytes:
    return _doc(
        "Памятка по обращению с отходами",
        "Персонал: {{ data.personnel }}",
        "Инструкции: {{ waste.instructions }}",
        "Контроль: {{ waste.control }}",
        footer="{{ logo }}",
    )


def _eco_fee_declaration() -> bytes:
    """Декларация о плате за НВОС — заготовка, а не копия госформы.

    Суммы — ответы специалиста по справочнику ставок и строкам расчёта
    (разд. 55.3, экран «Экология → Плата за НВОС»), не вычисления платформы.
    """

    return _doc(
        "Декларация о плате за негативное воздействие на окружающую среду",
        "Организация: {{ company.name }} (ИНН {{ company.inn }})",
        "Отчётный год: {{ data.eco_report_year }}",
        "Плата по выбросам загрязняющих веществ: {{ data.eco_fee_emissions }}",
        "Плата по сбросам загрязняющих веществ: {{ data.eco_fee_discharges }}",
        "Плата за размещение отходов: {{ data.eco_fee_waste }}",
        "Авансовые платежи по кварталам: {{ data.eco_fee_advances }}",
        header="{{ logo }}",
        footer="Составил: {{ data.eco_responsible }}",
    )


def _eco_2tp_waste() -> bytes:
    return _doc(
        "Форма 2-ТП (отходы): сведения об образовании и обращении с отходами",
        "Организация: {{ company.name }} (ИНН {{ company.inn }})",
        "Отчётный год: {{ data.eco_report_year }}",
        "Образовано отходов за год: {{ data.eco_waste_generated }}",
        "Передано операторам по договорам: {{ data.eco_waste_transferred }}",
        "Размещено на объектах размещения: {{ data.eco_waste_disposed }}",
        footer="Составил: {{ data.eco_responsible }}",
    )


def _eco_2tp_air() -> bytes:
    return _doc(
        "Форма 2-ТП (воздух): сведения об охране атмосферного воздуха",
        "Организация: {{ company.name }} (ИНН {{ company.inn }})",
        "Отчётный год: {{ data.eco_report_year }}",
        "Выброшено загрязняющих веществ за год: {{ data.eco_air_emitted }}",
        "Стационарных источников выбросов: {{ data.eco_air_sources }}",
        "Установки очистки газа: {{ data.eco_air_treatment }}",
        footer="Составил: {{ data.eco_responsible }}",
    )


def _eco_2tp_water() -> bytes:
    return _doc(
        "Форма 2-ТП (водхоз): сведения об использовании воды",
        "Организация: {{ company.name }} (ИНН {{ company.inn }})",
        "Отчётный год: {{ data.eco_report_year }}",
        "Забрано воды за год: {{ data.eco_water_intake }}",
        "Отведено сточных вод: {{ data.eco_water_discharge }}",
        "Приборы учёта: {{ data.eco_water_meters }}",
        footer="Составил: {{ data.eco_responsible }}",
    )


def _eco_pek_report() -> bytes:
    return _doc(
        "Отчёт об организации и результатах производственного экологического контроля",
        "Организация: {{ company.name }} (ИНН {{ company.inn }})",
        "Отчётный год: {{ data.eco_report_year }}",
        "Программа ПЭК: {{ data.eco_pek_program }}",
        "Выполненные замеры: {{ data.eco_pek_measurements }}",
        "Превышения по замерам: {{ data.eco_pek_exceedances }}",
        "Аккредитованная лаборатория: {{ data.eco_pek_laboratory }}",
        footer="Составил: {{ data.eco_responsible }}",
    )


def _eco_nvos_report() -> bytes:
    return _doc(
        "Отчётность по объекту негативного воздействия на окружающую среду",
        "Организация: {{ company.name }} (ИНН {{ company.inn }})",
        "Отчётный год: {{ data.eco_report_year }}",
        "Код объекта в государственном реестре: {{ data.eco_nvos_number }}",
        "Категория объекта: {{ data.eco_nvos_category }}",
        "Актуализация сведений об объекте: {{ data.eco_nvos_actualization }}",
        footer="Составил: {{ data.eco_responsible }}",
        header="{{ logo }}",
    )


def _ceo_order() -> bytes:
    return _doc(
        "Приказ генерального директора",
        "Организация: {{ company.name }}",
        "Назначить ответственным по ОТ: {{ shield.ot_lead }}",
        "Назначить ответственным по ПБ: {{ shield.pb_lead }}",
        "Назначить ответственным по экологии: {{ shield.eco_lead }}",
        footer="{{ stamp }}",
    )


def _ceo_delegation() -> bytes:
    return _doc(
        "Делегирование полномочий",
        "Директор: {{ shield.ceo }}",
        "Заместитель: {{ shield.deputy }}",
        "Объём полномочий: {{ shield.scope }}",
        "Срок действия: {{ shield.valid_until }}",
        header="{{ logo }}",
    )


def _ceo_control_card() -> bytes:
    return _doc(
        "Карта контроля директора",
        "Ключевые риски: {{ shield.risks }}",
        "Отчётность: {{ shield.reporting }}",
        "Каналы связи: {{ shield.contacts }}",
        footer="{{ stamp }}",
    )


# --- Срез-2: сценарии, которых не хватало каталогу (разд. 50.1) ---------------
#
# Модуль пакета берётся из существующего перечня БД (ot / fire_safety / health /
# custom): добавить в него «экологию» или «ГО и ЧС» значило бы миграцию типа
# ради ярлыка. Дисциплина по ТЗ записана в metadata пакета — так её видно и в
# каталоге, и в отчётах, а тип в базе остаётся прежним.


def _new_employee_intro_briefing() -> bytes:
    return _doc(
        "Вводный инструктаж по охране труда",
        "Компания: {{ company.name }} (ИНН {{ company.inn }})",
        "Работник: {{ data.employee_name }}, должность: {{ data.position }}",
        "Дата приёма: {{ data.hire_date }}",
        "Программа вводного инструктажа: {{ data.program }}",
        "Инструктаж провёл: {{ ot_responsible }}",
        header="{{ logo }}",
        footer="Подпись работника: {{ data.employee_sign }} · {{ stamp }}",
    )


def _new_employee_primary_briefing() -> bytes:
    return _doc(
        "Первичный инструктаж на рабочем месте",
        "Работник: {{ data.employee_name }}",
        "Рабочее место: {{ data.workplace }}",
        "Опасные и вредные факторы: {{ hazards }}",
        "Инструкции по охране труда: {{ data.instructions }}",
        "Стажировка: {{ data.internship_days }} смен",
        footer="Инструктаж провёл: {{ ot_responsible }}",
    )


def _new_employee_medical_referral() -> bytes:
    return _doc(
        "Направление на предварительный медицинский осмотр",
        "Компания: {{ company.name }}",
        "Работник: {{ data.employee_name }}",
        "Должность: {{ data.position }}",
        "Вредные факторы и виды работ: {{ hazards }}",
        "Медицинская организация: {{ data.medical_org }}",
        header="{{ logo }}",
        footer="{{ stamp }}",
    )


def _new_employee_ppe_card() -> bytes:
    return _doc(
        "Личная карточка учёта выдачи СИЗ",
        "Работник: {{ data.employee_name }}",
        "Должность: {{ data.position }}",
        "Нормы выдачи: {{ data.ppe_norms }}",
        "Размеры (одежда/обувь/СИЗОД): {{ data.sizes }}",
        footer="Выдал: {{ ot_responsible }} · Получил: {{ data.employee_sign }}",
    )


def _new_employee_acknowledgement() -> bytes:
    return _doc(
        "Лист ознакомления с локальными нормативными актами",
        "Работник: {{ data.employee_name }}",
        "Перечень документов: {{ data.documents }}",
        "Дата ознакомления: {{ data.hire_date }}",
        footer="Подпись работника: {{ data.employee_sign }}",
    )


def _contractor_questionnaire() -> bytes:
    return _doc(
        "Анкета подрядной организации",
        "Подрядчик: {{ data.contractor_name }} (ИНН {{ data.contractor_inn }})",
        "Виды выполняемых работ: {{ work_types }}",
        "Численность привлекаемого персонала: {{ data.headcount }}",
        "Наличие СРО/допусков: {{ data.permits }}",
        "Ответственный от подрядчика: {{ data.contractor_responsible }}",
        header="{{ logo }}",
        footer="{{ stamp }}",
    )


def _contractor_readiness_check() -> bytes:
    return _doc(
        "Проверка готовности подрядчика",
        "1. Обучение по охране труда: {{ data.training_status }}",
        "2. Медосмотры персонала: {{ data.medical_status }}",
        "3. Обеспеченность СИЗ: {{ data.ppe_status }}",
        "4. Наряды-допуски на опасные работы: {{ data.permits_status }}",
        "5. Страхование ответственности: {{ data.insurance_status }}",
        footer="Проверку провёл: {{ ot_responsible }}",
    )


def _contractor_interaction_order() -> bytes:
    return _doc(
        "Приказ о порядке взаимодействия с подрядной организацией",
        "Заказчик: {{ company.name }}",
        "Подрядчик: {{ data.contractor_name }}",
        "Объект: {{ site.address }}",
        "Ответственный от заказчика: {{ ot_responsible }}",
        "Разграничение зон ответственности: {{ data.responsibility_split }}",
        header="{{ logo }}",
        footer="{{ stamp }}",
    )


def _contractor_briefing_log() -> bytes:
    return _doc(
        "Журнал инструктажей персонала подрядчика",
        "Подрядчик: {{ data.contractor_name }}",
        "Дата: {{ data.session_date }}",
        "Состав: {{ data.team_members }}",
        "Опасности объекта: {{ hazards }}",
        footer="Инструктаж провёл: {{ ot_responsible }}",
    )


def _fire_inspection_checklist() -> bytes:
    return _doc(
        "Чек-лист проверки противопожарного состояния объекта",
        "Объект: {{ site.address }}",
        "1. Первичные средства пожаротушения: {{ data.extinguishers }}",
        "2. Пути эвакуации и выходы: {{ data.escape_routes }}",
        "3. Системы сигнализации и оповещения: {{ data.alarm_systems }}",
        "4. Противопожарные двери и преграды: {{ data.fire_doors }}",
        "5. Замечания: {{ data.findings }}",
        header="{{ logo }}",
        footer="Проверку провёл: {{ pb_responsible }} · {{ stamp }}",
    )


def _fire_inspection_order() -> bytes:
    return _doc(
        "Приказ о противопожарном режиме на объекте",
        "Компания: {{ company.name }}",
        "Объект: {{ site.address }}",
        "Ответственный за пожарную безопасность: {{ pb_responsible }}",
        "Места курения и порядок огневых работ: {{ data.fire_works_rules }}",
        "Порядок обесточивания по окончании работ: {{ data.shutdown_rules }}",
        header="{{ logo }}",
        footer="{{ stamp }}",
    )


def _fire_inspection_instruction() -> bytes:
    return _doc(
        "Инструкция о мерах пожарной безопасности",
        "Объект: {{ site.address }}",
        "Порядок содержания территории и помещений: {{ data.housekeeping }}",
        "Действия при обнаружении пожара: {{ data.fire_actions }}",
        "Порядок эвакуации людей и материальных ценностей: {{ data.evacuation }}",
        footer="Ответственный: {{ pb_responsible }}",
    )


def _fire_inspection_drill_program() -> bytes:
    return _doc(
        "Программа практической тренировки по эвакуации",
        "Объект: {{ site.address }}",
        "Дата тренировки: {{ data.drill_date }}",
        "Участники: {{ data.team_members }}",
        "Вводная обстановка: {{ data.scenario }}",
        "Оценка результатов: {{ data.drill_result }}",
        footer="Организатор: {{ pb_responsible }}",
    )


def _fire_inspection_journal() -> bytes:
    return _doc(
        "Журнал эксплуатации систем противопожарной защиты",
        "Объект: {{ site.address }}",
        "Проверяемая система: {{ data.system_name }}",
        "Результат проверки: {{ data.system_result }}",
        "Дата следующей проверки: {{ data.next_check_date }}",
        footer="Отметку внёс: {{ pb_responsible }}",
    )


def _civil_defence_plan() -> bytes:
    return _doc(
        "План действий по предупреждению и ликвидации чрезвычайных ситуаций",
        "Компания: {{ company.name }}",
        "Объект: {{ site.address }}",
        "Возможные чрезвычайные ситуации: {{ data.threats }}",
        "Порядок оповещения: {{ data.notification_order }}",
        "Силы и средства: {{ data.resources }}",
        header="{{ logo }}",
        footer="{{ stamp }}",
    )


def _civil_defence_order() -> bytes:
    return _doc(
        "Приказ об организации гражданской обороны",
        "Компания: {{ company.name }}",
        "Ответственный за гражданскую оборону: {{ data.gochs_responsible }}",
        "Категория объекта: {{ data.facility_category }}",
        "Порядок хранения средств индивидуальной защиты: {{ data.protection_storage }}",
        header="{{ logo }}",
        footer="{{ stamp }}",
    )


def _civil_defence_commission() -> bytes:
    return _doc(
        "Состав комиссии по предупреждению и ликвидации ЧС",
        "Председатель комиссии: {{ data.commission_head }}",
        "Члены комиссии: {{ data.commission_members }}",
        "Задачи комиссии: {{ data.commission_tasks }}",
        footer="Утверждаю: {{ data.gochs_responsible }}",
    )


def _civil_defence_drill_program() -> bytes:
    return _doc(
        "Программа учений и тренировок по гражданской обороне",
        "Объект: {{ site.address }}",
        "Тема учения: {{ data.scenario }}",
        "Дата проведения: {{ data.drill_date }}",
        "Привлекаемые работники: {{ data.team_members }}",
        footer="Организатор: {{ data.gochs_responsible }}",
    )


def _civil_defence_journal() -> bytes:
    return _doc(
        "Журнал учёта занятий по гражданской обороне",
        "Тема занятия: {{ data.lesson_topic }}",
        "Дата: {{ data.session_date }}",
        "Присутствовали: {{ data.team_members }}",
        footer="Занятие провёл: {{ data.gochs_responsible }}",
    )


def _bdd_management_regulation() -> bytes:
    """Положение о системе управления БДД.

    ТЗ просит «положения», и это ровно тот документ: с него начинается
    проверка организации, эксплуатирующей транспорт. В базовом комплекте его
    не было — печатались приказ, инструкция и план.
    """

    return _doc(
        "Положение о системе управления безопасностью дорожного движения",
        "Организация: {{ company.name }} (ИНН {{ company.inn }})",
        "Ответственный за БДД: {{ data.bdd_responsible }}",
        "Цели в области БДД: {{ data.bdd_goals }}",
        "Порядок контроля и учёта: {{ data.bdd_control_order }}",
        "Порядок разбора происшествий: {{ data.bdd_review_order }}",
        header="{{ logo }}",
        footer="Утверждено: {{ data.bdd_approved_by }} · {{ stamp }}",
    )


def _bdd_accident_report() -> bytes:
    """Отчёт о состоянии аварийности — ЗАГОТОВКА, а не выборка из реестров.

    Числа вносит специалист: платформа умеет показать их в сводке контура, но
    подставлять их в документ автоматически она пока не умеет ни для одной
    дисциплины (то же у отчётных форм экологии). Предзаполнение мастера из
    реестров — отдельная работа, названная в журнале волн.
    """

    return _doc(
        "Отчёт о состоянии аварийности",
        "Организация: {{ company.name }}",
        "Отчётный период: {{ data.bdd_report_period }}",
        "Транспортных средств в эксплуатации: {{ data.bdd_report_vehicles }}",
        "Водителей допущено: {{ data.bdd_report_drivers }}",
        "Дорожно-транспортных происшествий: {{ data.bdd_report_accidents }}",
        "Пострадало / погибло: {{ data.bdd_report_injured }}",
        "Нарушений ПДД: {{ data.bdd_report_violations }}",
        footer="Составил: {{ data.bdd_report_author }} · {{ stamp }}",
    )


def _bdd_measures_report() -> bytes:
    """Отчёт о выполнении мероприятий по предупреждению аварийности."""

    return _doc(
        "Отчёт о выполнении мероприятий по предупреждению ДТП",
        "Организация: {{ company.name }}",
        "Отчётный период: {{ data.bdd_report_period }}",
        "Запланировано мероприятий: {{ data.bdd_measures_planned }}",
        "Выполнено: {{ data.bdd_measures_done }}",
        "Причины невыполнения: {{ data.bdd_measures_failed_reason }}",
        footer="Составил: {{ data.bdd_report_author }} · {{ stamp }}",
    )


def _bdd_vehicle_assignment_order() -> bytes:
    """Приказ о закреплении ТС за водителями.

    ГРАНИЦА: платформа не решает, за кем закреплять машину, — она печатает
    внесённое.
    """

    return _doc(
        "Приказ о закреплении транспортных средств за водителями",
        "Организация: {{ company.name }}",
        "Дата: {{ data.bdd_assignment_date }}",
        "Закрепление: {{ data.bdd_assignment_list }}",
        "Ответственный за БДД: {{ data.bdd_responsible }}",
        footer="Руководитель: {{ data.bdd_approved_by }} · {{ stamp }}",
    )


def _gochs_rsches_regulation() -> bytes:
    """Положение об объектовом звене РСЧС.

    ТЗ просит «положения», и это главное из них: им организация определяет
    своё звено единой системы предупреждения и ликвидации ЧС. В базовом
    комплекте печатались приказ и план, положения не было.
    """

    return _doc(
        "Положение об объектовом звене РСЧС",
        "Организация: {{ company.name }} (ИНН {{ company.inn }})",
        "Категория объекта по ГО: {{ data.facility_category }}",
        "Состав звена: {{ data.gochs_unit_structure }}",
        "Задачи звена: {{ data.gochs_unit_tasks }}",
        "Режимы функционирования: {{ data.gochs_unit_modes }}",
        header="{{ logo }}",
        footer="Утверждено: {{ data.gochs_approved_by }} · {{ stamp }}",
    )


def _gochs_emergency_instruction() -> bytes:
    """Инструкция о действиях при угрозе и возникновении ЧС."""

    return _doc(
        "Инструкция о действиях при угрозе и возникновении ЧС",
        "Организация: {{ company.name }}",
        "Порядок оповещения: {{ data.gochs_alert_order }}",
        "Порядок сбора и эвакуации: {{ data.gochs_evacuation_order }}",
        "Места укрытия: {{ data.gochs_shelters }}",
        "Действия персонала: {{ data.gochs_staff_actions }}",
        footer="Ответственный за ГО и ЧС: {{ data.gochs_responsible }} · {{ stamp }}",
    )


def _gochs_emergency_report() -> bytes:
    """Донесение о чрезвычайной ситуации.

    ГРАНИЦА: платформа НЕ решает, какое донесение и в какой срок подавать —
    это зависит от вида и масштаба ЧС и от требований органа управления ГОЧС.
    Здесь заготовка, которую заполняет специалист.
    """

    return _doc(
        "Донесение о чрезвычайной ситуации",
        "Организация: {{ company.name }}",
        "Дата и время ЧС: {{ data.gochs_event_at }}",
        "Вид и краткая характеристика: {{ data.gochs_event_kind }}",
        "Пострадало / погибло: {{ data.gochs_event_injured }}",
        "Принятые меры: {{ data.gochs_event_measures }}",
        "Привлечённые силы и средства: {{ data.gochs_event_forces }}",
        footer="Донесение подготовил: {{ data.gochs_report_author }} · {{ stamp }}",
    )


def _gochs_forces_report() -> bytes:
    """Сведения о силах и средствах ГО — отчётность в орган управления."""

    return _doc(
        "Сведения о силах и средствах гражданской обороны",
        "Организация: {{ company.name }}",
        "Отчётный период: {{ data.gochs_report_period }}",
        "Нештатных формирований: {{ data.gochs_formations_count }}",
        "Численность личного состава: {{ data.gochs_personnel_count }}",
        "Обеспеченность СИЗ: {{ data.gochs_ppe_coverage }}",
        "Средства оповещения: {{ data.gochs_alert_means }}",
        footer="Составил: {{ data.gochs_report_author }} · {{ stamp }}",
    )


def _opo_pc_regulation() -> bytes:
    """Положение о производственном контроле.

    ТЗ просит «положения», и для ОПО главное из них — это: ФЗ-116 ст. 11
    обязывает эксплуатирующую организацию организовать производственный
    контроль, а Правила организации ПК — иметь положение о нём. В базовом
    комплекте ОПО печатались паспорт, ПЛА и матрица обучения; положения не
    было.
    """

    return _doc(
        "Положение о производственном контроле за соблюдением требований промышленной безопасности",
        "Организация: {{ company.name }} (ИНН {{ company.inn }})",
        "Эксплуатируемых ОПО: {{ data.opo_facilities_count }}",
        "Ответственный за производственный контроль: {{ data.opo_pc_responsible }}",
        "Задачи производственного контроля: {{ data.opo_pc_tasks }}",
        "Порядок проверок и обследований: {{ data.opo_pc_inspection_order }}",
        "Порядок учёта и отчётности: {{ data.opo_pc_reporting_order }}",
        header="{{ logo }}",
        footer="Утверждено: {{ data.opo_approved_by }} · {{ stamp }}",
    )


def _opo_pc_order() -> bytes:
    """Приказ о назначении ответственного за производственный контроль."""

    return _doc(
        "Приказ о назначении ответственного за осуществление производственного контроля",
        "Организация: {{ company.name }}",
        "Дата приказа: {{ data.opo_order_date }}",
        "Ответственный за производственный контроль: {{ data.opo_pc_responsible }}",
        "Основание: аттестация в области промышленной безопасности {{ data.opo_pc_attestation }}",
        "Ответственные за эксплуатацию ОПО: {{ data.opo_facility_responsibles }}",
        header="{{ logo }}",
        footer="Руководитель: {{ data.opo_director }} · {{ stamp }}",
    )


def _opo_pc_report() -> bytes:
    """Сведения об организации производственного контроля — годовой отчёт.

    Правила организации ПК требуют представлять в Ростехнадзор сведения об
    организации производственного контроля ЗА ПРОШЕДШИЙ год. Здесь заготовка,
    которую заполняет и подписывает специалист.

    ГРАНИЦА: платформа НЕ решает, обязана ли организация подавать сведения и
    по каким ОПО — это зависит от того, что и как она эксплуатирует, и от
    требований территориального органа.
    """

    return _doc(
        "Сведения об организации производственного контроля за соблюдением требований промышленной безопасности",
        "Организация: {{ company.name }} (ИНН {{ company.inn }})",
        "Отчётный год: {{ data.opo_report_year }}",
        "Ответственный за производственный контроль: {{ data.opo_pc_responsible }}",
        "Эксплуатируемых ОПО: {{ data.opo_facilities_count }}",
        "ОПО по классам опасности: {{ data.opo_facilities_by_class }}",
        "Технических устройств в эксплуатации: {{ data.opo_devices_count }}",
        "Экспертиз промышленной безопасности проведено: {{ data.opo_epb_done }}",
        "Устройств с просроченным заключением ЭПБ: {{ data.opo_epb_overdue }}",
        "Действующих аттестаций работников: {{ data.opo_attestations_count }}",
        "Просроченных аттестаций: {{ data.opo_attestations_overdue }}",
        "Мероприятий плана ПК запланировано: {{ data.opo_pc_measures_planned }}",
        "Мероприятий плана ПК выполнено: {{ data.opo_pc_measures_done }}",
        "Аварий и инцидентов: {{ data.opo_incidents_count }}",
        "Выявленных нарушений и принятые меры: {{ data.opo_violations_note }}",
        footer="Составил: {{ data.opo_report_author }} · {{ stamp }}",
    )


def _opo_incident_report() -> bytes:
    """Сведения об инцидентах на ОПО — периодический отчёт в Ростехнадзор.

    Учёт инцидентов ведёт организация, а сведения о них подаёт в
    территориальный орган за период. Второй из «отчётов для Ростехнадзора»
    формулировки ТЗ.
    """

    return _doc(
        "Сведения об инцидентах, происшедших на опасных производственных объектах",
        "Организация: {{ company.name }}",
        "Отчётный период: {{ data.opo_incident_period }}",
        "Инцидентов за период: {{ data.opo_incidents_count }}",
        "Причины инцидентов: {{ data.opo_incident_causes }}",
        "Принятые меры: {{ data.opo_incident_measures }}",
        "Продолжительность простоя: {{ data.opo_incident_downtime }}",
        footer="Составил: {{ data.opo_report_author }} · {{ stamp }}",
    )


def _road_safety_order() -> bytes:
    return _doc(
        "Приказ о назначении ответственного за обеспечение безопасности дорожного движения",
        "Компания: {{ company.name }}",
        "Ответственный за БДД: {{ data.bdd_responsible }}",
        "Основание: аттестация по БДД от {{ data.bdd_attestation_date }}",
        "Зона ответственности: {{ data.bdd_fleet_scope }}",
        header="{{ logo }}",
        footer="{{ stamp }}",
    )


def _road_safety_instruction() -> bytes:
    return _doc(
        "Инструкция по обеспечению безопасности дорожного движения",
        "Компания: {{ company.name }}",
        "Порядок выпуска транспорта на линию: {{ data.bdd_dispatch_order }}",
        "Предрейсовый медицинский осмотр: {{ data.bdd_medical_check_order }}",
        "Предрейсовый контроль технического состояния: {{ data.bdd_tech_check_order }}",
        "Режим труда и отдыха водителей: {{ data.bdd_driver_schedule }}",
        header="{{ logo }}",
        footer="{{ stamp }}",
    )


def _road_safety_action_plan() -> bytes:
    return _doc(
        "План мероприятий по предупреждению дорожно-транспортных происшествий",
        "Период: {{ data.bdd_plan_period }}",
        "Мероприятия: {{ data.bdd_measures }}",
        "Ответственные: {{ data.bdd_responsible }}",
        "Отметка о выполнении: {{ data.bdd_completion_note }}",
        footer="Утверждаю: {{ data.bdd_director }}",
    )


def _road_safety_briefing_program() -> bytes:
    return _doc(
        "Программа инструктажа водителей по безопасности дорожного движения",
        "Вид инструктажа: {{ data.bdd_briefing_kind }}",
        "Темы: {{ data.bdd_briefing_topics }}",
        "Продолжительность: {{ data.bdd_briefing_hours }}",
        "Инструктаж провёл: {{ data.bdd_responsible }}",
        footer="Отметка о проведении: {{ data.bdd_briefing_date }}",
    )


DEFAULT_PACKS: Sequence[PackDefinition] = (
    PackDefinition(
        code=PACK_CODE_SITE_ACCESS,
        name="Выход на объект",
        description="Комплект документов для допуска бригады на объект",
        scenario=PackScenario.ENTER_SITE,
        module=DocumentPackModule.OT,
        scenario_type=DocumentPackScenario.DOCUMENT_BATCH,
        templates=(
            PackTemplateSpec(
                code="pack_ot_site_order",
                name="Приказ на выход",
                description="Распоряжение о выходе на объект",
                category="order",
                builder=_site_access_order,
            ),
            PackTemplateSpec(
                code="pack_ot_site_briefing",
                name="Инструктаж перед выходом",
                description="Протокол вводного инструктажа",
                category="instruction",
                builder=_site_access_briefing,
            ),
            PackTemplateSpec(
                code="pack_ot_site_ppe_card",
                name="Карточка СИЗ",
                description="Личная карточка выдачи СИЗ",
                category="ppe_card",
                builder=_site_access_ppe_card,
            ),
            PackTemplateSpec(
                code="pack_ot_site_journal",
                name="Журнал инструктажей",
                description="Лист регистрации инструктажа",
                category="journal",
                builder=_site_access_journal_entry,
            ),
            PackTemplateSpec(
                code="pack_ot_site_checklist",
                name="Чек-лист допуска",
                description="Контроль требований перед выходом",
                category="checklist",
                builder=_site_access_checklist,
            ),
        ),
        item_order=(
            "pack_ot_site_order",
            "pack_ot_site_briefing",
            "pack_ot_site_ppe_card",
            "pack_ot_site_journal",
            "pack_ot_site_checklist",
        ),
        metadata={
            "logo": build_inline_image_descriptor(DEFAULT_LOGO_BYTES)["data"],
            "stamp": build_inline_image_descriptor(DEFAULT_STAMP_BYTES)["data"],
            "discipline": "Охрана труда и промышленная безопасность",
        },
        disciplines=(),
    ),
    PackDefinition(
        code=PACK_CODE_NEW_COMPANY,
        name="Новая компания",
        description="Стартовый комплект документов по охране труда и пожарной безопасности",
        scenario=PackScenario.NEW_COMPANY,
        module=DocumentPackModule.OT,
        scenario_type=DocumentPackScenario.DOCUMENT_BATCH,
        templates=(
            PackTemplateSpec(
                code="pack_new_company_order",
                name="Приказ",
                description="Приказ о создании службы ОТ",
                category="order",
                builder=_new_company_order,
            ),
            PackTemplateSpec(
                code="pack_new_company_instruction",
                name="Инструкция",
                description="Базовая инструкция по охране труда",
                category="instruction",
                builder=_new_company_instruction,
            ),
            PackTemplateSpec(
                code="pack_new_company_journal",
                name="Журнал",
                description="Журнал регистрации инструктажей",
                category="journal",
                builder=_new_company_journal,
            ),
            PackTemplateSpec(
                code="pack_new_company_policy",
                name="Политика",
                description="Политика в области ОТ/ПБ",
                category="policy",
                builder=_new_company_policy,
            ),
        ),
        item_order=(
            "pack_new_company_order",
            "pack_new_company_instruction",
            "pack_new_company_journal",
            "pack_new_company_policy",
        ),
        metadata={
            "logo": build_inline_image_descriptor(DEFAULT_LOGO_BYTES)["data"],
            "stamp": build_inline_image_descriptor(DEFAULT_STAMP_BYTES)["data"],
            "discipline": "Охрана труда и пожарная безопасность",
        },
        disciplines=(Discipline.FIRE_SAFETY,),
    ),
    PackDefinition(
        code=PACK_CODE_INCIDENT,
        name="Несчастный случай",
        description="Документы для расследования несчастного случая",
        scenario=PackScenario.INCIDENT,
        module=DocumentPackModule.OT,
        scenario_type=DocumentPackScenario.WORKFLOW,
        templates=(
            PackTemplateSpec(
                code="pack_incident_report",
                name="Акт",
                description="Акт о несчастном случае",
                category="report",
                builder=_incident_report,
            ),
            PackTemplateSpec(
                code="pack_incident_commission",
                name="Приказ",
                description="Приказ о создании комиссии",
                category="order",
                builder=_incident_commission_order,
            ),
            PackTemplateSpec(
                code="pack_incident_plan",
                name="План мероприятий",
                description="Перечень корректирующих действий",
                category="plan",
                builder=_incident_plan,
            ),
        ),
        item_order=(
            "pack_incident_report",
            "pack_incident_commission",
            "pack_incident_plan",
        ),
        metadata={
            "logo": build_inline_image_descriptor(DEFAULT_LOGO_BYTES)["data"],
            "stamp": build_inline_image_descriptor(DEFAULT_STAMP_BYTES)["data"],
            "discipline": "Охрана труда",
            "entity": "incident",
            "models": "Incident,IncidentLog",
            "context_keys": "incident_id,company_id,site_id,victim_ids",
        },
        disciplines=(),
    ),
    PackDefinition(
        code=PACK_CODE_INSPECTION_PREP,
        name="Подготовка к проверке",
        description="Комплект для подготовки к проверке ГИТ/МЧС/Ростехнадзора",
        scenario=PackScenario.INSPECTION_PREP,
        module=DocumentPackModule.OT,
        scenario_type=DocumentPackScenario.WORKFLOW,
        templates=(
            PackTemplateSpec(
                code="pack_inspection_plan",
                name="План",
                description="План подготовки к визиту",
                category="plan",
                builder=_inspection_plan,
            ),
            PackTemplateSpec(
                code="pack_inspection_checklist",
                name="Чек-лист",
                description="Чек-лист готовности к проверке",
                category="checklist",
                builder=_inspection_checklist,
            ),
            PackTemplateSpec(
                code="pack_inspection_briefing",
                name="Инструктаж",
                description="Инструктаж команды перед проверкой",
                category="instruction",
                builder=_inspection_briefing,
            ),
        ),
        item_order=(
            "pack_inspection_plan",
            "pack_inspection_checklist",
            "pack_inspection_briefing",
        ),
        metadata={
            "logo": build_inline_image_descriptor(DEFAULT_LOGO_BYTES)["data"],
            "stamp": build_inline_image_descriptor(DEFAULT_STAMP_BYTES)["data"],
            "discipline": "Любая — по типу надзора",
            "entity": "inspection",
            "context_keys": "inspection_id,company_id,site_id",
        },
        disciplines=(),
    ),
    PackDefinition(
        code=PACK_CODE_OPO,
        name="ОПО",
        description="Документы для опасных производственных объектов",
        scenario=PackScenario.OPO,
        module=DocumentPackModule.OT,
        scenario_type=DocumentPackScenario.DOCUMENT_BATCH,
        templates=(
            PackTemplateSpec(
                code="pack_opo_passport",
                name="Паспорт",
                description="Паспорт опасного объекта",
                category="passport",
                builder=_opo_passport,
            ),
            PackTemplateSpec(
                code="pack_opo_emergency_plan",
                name="ПЛА",
                description="План ликвидации аварий",
                category="plan",
                builder=_opo_emergency_plan,
            ),
            PackTemplateSpec(
                code="pack_opo_training_matrix",
                name="Матрица обучения",
                description="Матрица обучения и допусков",
                category="training",
                builder=_opo_training_matrix,
            ),
        ),
        item_order=(
            "pack_opo_passport",
            "pack_opo_emergency_plan",
            "pack_opo_training_matrix",
        ),
        metadata={
            "logo": build_inline_image_descriptor(DEFAULT_LOGO_BYTES)["data"],
            "stamp": build_inline_image_descriptor(DEFAULT_STAMP_BYTES)["data"],
            "discipline": "Промышленная безопасность",
        },
        disciplines=(Discipline.INDUSTRIAL_SAFETY,),
    ),
    PackDefinition(
        code=PACK_CODE_WASTE,
        name="Отходы/экология",
        description="Минимальный пакет для обращения с отходами",
        scenario=PackScenario.WASTE,
        module=DocumentPackModule.HEALTH,
        scenario_type=DocumentPackScenario.DOCUMENT_BATCH,
        templates=(
            PackTemplateSpec(
                code="pack_waste_register",
                name="Реестр",
                description="Реестр отходов",
                category="register",
                builder=_waste_register,
            ),
            PackTemplateSpec(
                code="pack_waste_contract",
                name="Договор",
                description="Договор с подрядчиком",
                category="contract",
                builder=_waste_contract,
            ),
            PackTemplateSpec(
                code="pack_waste_briefing",
                name="Памятка",
                description="Инструктаж по обращению с отходами",
                category="instruction",
                builder=_waste_briefing,
            ),
        ),
        item_order=(
            "pack_waste_register",
            "pack_waste_contract",
            "pack_waste_briefing",
        ),
        metadata={
            "logo": build_inline_image_descriptor(DEFAULT_LOGO_BYTES)["data"],
            "stamp": build_inline_image_descriptor(DEFAULT_STAMP_BYTES)["data"],
            "discipline": "Экология",
        },
        disciplines=(Discipline.ECOLOGY,),
    ),
    # Срез-8 экологии (разд. 55.3): отчётные формы. Заготовки с
    # плейсхолдерами, как весь каталог, а не копии госформ: значения
    # специалист берёт из реестров экологии (срезы 2–7), суммы и массы —
    # ответы мастера, не вычисления платформы.
    PackDefinition(
        code=PACK_CODE_ECO_REPORTS,
        name="Экологическая отчётность",
        description="Формы отчётности: декларация о плате, 2-ТП, отчёт по ПЭК, объект НВОС",
        scenario=PackScenario.ECOLOGY_REPORTING,
        module=DocumentPackModule.HEALTH,
        scenario_type=DocumentPackScenario.DOCUMENT_BATCH,
        templates=(
            PackTemplateSpec(
                code="pack_eco_fee_declaration",
                name="Декларация о плате за НВОС",
                description="Плата по выбросам, сбросам и отходам с авансовыми платежами",
                category="report",
                builder=_eco_fee_declaration,
            ),
            PackTemplateSpec(
                code="pack_eco_2tp_waste",
                name="2-ТП (отходы)",
                description="Образование и обращение с отходами за год",
                category="report",
                builder=_eco_2tp_waste,
            ),
            PackTemplateSpec(
                code="pack_eco_2tp_air",
                name="2-ТП (воздух)",
                description="Охрана атмосферного воздуха за год",
                category="report",
                builder=_eco_2tp_air,
            ),
            PackTemplateSpec(
                code="pack_eco_2tp_water",
                name="2-ТП (водхоз)",
                description="Использование воды за год",
                category="report",
                builder=_eco_2tp_water,
            ),
            PackTemplateSpec(
                code="pack_eco_pek_report",
                name="Отчёт по ПЭК",
                description="Организация и результаты производственного экологического контроля",
                category="report",
                builder=_eco_pek_report,
            ),
            PackTemplateSpec(
                code="pack_eco_nvos_report",
                name="Отчётность по объектам НВОС",
                description="Сведения по объекту государственного реестра НВОС",
                category="report",
                builder=_eco_nvos_report,
            ),
        ),
        item_order=(
            "pack_eco_fee_declaration",
            "pack_eco_2tp_waste",
            "pack_eco_2tp_air",
            "pack_eco_2tp_water",
            "pack_eco_pek_report",
            "pack_eco_nvos_report",
        ),
        metadata={
            "logo": build_inline_image_descriptor(DEFAULT_LOGO_BYTES)["data"],
            "stamp": build_inline_image_descriptor(DEFAULT_STAMP_BYTES)["data"],
            "discipline": "Экология",
        },
        disciplines=(Discipline.ECOLOGY,),
    ),
    PackDefinition(
        code=PACK_CODE_CEO_SHIELD,
        name="Щит генерального директора",
        description="Базовый пакет для распределения ответственности и контроля",
        scenario=PackScenario.CEO_SHIELD,
        module=DocumentPackModule.FIRE_SAFETY,
        scenario_type=DocumentPackScenario.DOCUMENT_BATCH,
        templates=(
            PackTemplateSpec(
                code="pack_ceo_order",
                name="Приказ",
                description="Приказ генерального директора",
                category="order",
                builder=_ceo_order,
            ),
            PackTemplateSpec(
                code="pack_ceo_delegation",
                name="Делегирование",
                description="Делегирование полномочий",
                category="delegation",
                builder=_ceo_delegation,
            ),
            PackTemplateSpec(
                code="pack_ceo_control_card",
                name="Контроль",
                description="Карта контроля директора",
                category="control",
                builder=_ceo_control_card,
            ),
        ),
        item_order=(
            "pack_ceo_order",
            "pack_ceo_delegation",
            "pack_ceo_control_card",
        ),
        metadata={
            "logo": build_inline_image_descriptor(DEFAULT_LOGO_BYTES)["data"],
            "stamp": build_inline_image_descriptor(DEFAULT_STAMP_BYTES)["data"],
            "discipline": "Общее руководство",
        },
        disciplines=(),
    ),
    PackDefinition(
        code=PACK_CODE_NEW_EMPLOYEE,
        name="Приём нового сотрудника",
        description="Комплект документов на приём работника: инструктажи, медосмотр, СИЗ, ознакомления",
        scenario=PackScenario.NEW_EMPLOYEE,
        module=DocumentPackModule.OT,
        scenario_type=DocumentPackScenario.DOCUMENT_BATCH,
        templates=(
            PackTemplateSpec(
                code="pack_ot_new_employee_intro",
                name="Вводный инструктаж",
                description="Протокол вводного инструктажа по охране труда",
                category="instruction",
                builder=_new_employee_intro_briefing,
            ),
            PackTemplateSpec(
                code="pack_ot_new_employee_primary",
                name="Первичный инструктаж",
                description="Первичный инструктаж на рабочем месте",
                category="instruction",
                builder=_new_employee_primary_briefing,
            ),
            PackTemplateSpec(
                code="pack_ot_new_employee_medical",
                name="Направление на медосмотр",
                description="Направление на предварительный медицинский осмотр",
                category="referral",
                builder=_new_employee_medical_referral,
            ),
            PackTemplateSpec(
                code="pack_ot_new_employee_ppe_card",
                name="Карточка СИЗ",
                description="Личная карточка учёта выдачи СИЗ",
                category="ppe_card",
                builder=_new_employee_ppe_card,
            ),
            PackTemplateSpec(
                code="pack_ot_new_employee_ack",
                name="Лист ознакомления",
                description="Ознакомление с локальными нормативными актами",
                category="acknowledgement",
                builder=_new_employee_acknowledgement,
            ),
        ),
        item_order=(
            "pack_ot_new_employee_intro",
            "pack_ot_new_employee_primary",
            "pack_ot_new_employee_medical",
            "pack_ot_new_employee_ppe_card",
            "pack_ot_new_employee_ack",
        ),
        metadata={
            "logo": build_inline_image_descriptor(DEFAULT_LOGO_BYTES)["data"],
            "stamp": build_inline_image_descriptor(DEFAULT_STAMP_BYTES)["data"],
            "discipline": "Охрана труда",
        },
        disciplines=(Discipline.MEDICAL, Discipline.PPE, Discipline.TRAINING),
    ),
    PackDefinition(
        code=PACK_CODE_CONTRACTOR,
        name="Новый подрядчик",
        description="Анкета, проверка готовности, приказ о взаимодействии и инструктажи персонала подрядчика",
        scenario=PackScenario.CONTRACTOR,
        module=DocumentPackModule.OT,
        scenario_type=DocumentPackScenario.DOCUMENT_BATCH,
        templates=(
            PackTemplateSpec(
                code="pack_ot_contractor_questionnaire",
                name="Анкета подрядчика",
                description="Сведения о подрядной организации и её работах",
                category="questionnaire",
                builder=_contractor_questionnaire,
            ),
            PackTemplateSpec(
                code="pack_ot_contractor_readiness",
                name="Проверка готовности",
                description="Контроль обучения, медосмотров, СИЗ и допусков",
                category="checklist",
                builder=_contractor_readiness_check,
            ),
            PackTemplateSpec(
                code="pack_ot_contractor_order",
                name="Приказ о взаимодействии",
                description="Разграничение зон ответственности заказчика и подрядчика",
                category="order",
                builder=_contractor_interaction_order,
            ),
            PackTemplateSpec(
                code="pack_ot_contractor_journal",
                name="Журнал инструктажей",
                description="Инструктажи персонала подрядчика на объекте",
                category="journal",
                builder=_contractor_briefing_log,
            ),
        ),
        item_order=(
            "pack_ot_contractor_questionnaire",
            "pack_ot_contractor_readiness",
            "pack_ot_contractor_order",
            "pack_ot_contractor_journal",
        ),
        metadata={
            "logo": build_inline_image_descriptor(DEFAULT_LOGO_BYTES)["data"],
            "stamp": build_inline_image_descriptor(DEFAULT_STAMP_BYTES)["data"],
            "discipline": "Охрана труда и промышленная безопасность",
        },
        disciplines=(),
    ),
    PackDefinition(
        code=PACK_CODE_FIRE_INSPECTION,
        name="Пожарная проверка объекта",
        description="Приказы, инструкции, журналы и программа тренировки по пожарной безопасности",
        scenario=PackScenario.FIRE_INSPECTION,
        module=DocumentPackModule.FIRE_SAFETY,
        scenario_type=DocumentPackScenario.DOCUMENT_BATCH,
        templates=(
            PackTemplateSpec(
                code="pack_pb_inspection_checklist",
                name="Чек-лист проверки",
                description="Противопожарное состояние объекта",
                category="checklist",
                builder=_fire_inspection_checklist,
            ),
            PackTemplateSpec(
                code="pack_pb_inspection_order",
                name="Приказ о противопожарном режиме",
                description="Режим на объекте и ответственные",
                category="order",
                builder=_fire_inspection_order,
            ),
            PackTemplateSpec(
                code="pack_pb_inspection_instruction",
                name="Инструкция о мерах ПБ",
                description="Меры пожарной безопасности и действия при пожаре",
                category="instruction",
                builder=_fire_inspection_instruction,
            ),
            PackTemplateSpec(
                code="pack_pb_inspection_drill",
                name="Программа тренировки",
                description="Практическая тренировка по эвакуации",
                category="program",
                builder=_fire_inspection_drill_program,
            ),
            PackTemplateSpec(
                code="pack_pb_inspection_journal",
                name="Журнал систем защиты",
                description="Эксплуатация систем противопожарной защиты",
                category="journal",
                builder=_fire_inspection_journal,
            ),
        ),
        item_order=(
            "pack_pb_inspection_checklist",
            "pack_pb_inspection_order",
            "pack_pb_inspection_instruction",
            "pack_pb_inspection_drill",
            "pack_pb_inspection_journal",
        ),
        metadata={
            "logo": build_inline_image_descriptor(DEFAULT_LOGO_BYTES)["data"],
            "stamp": build_inline_image_descriptor(DEFAULT_STAMP_BYTES)["data"],
            "discipline": "Пожарная безопасность",
        },
        disciplines=(Discipline.FIRE_SAFETY,),
    ),
    PackDefinition(
        code=PACK_CODE_CIVIL_DEFENCE,
        name="Пакет ГО и ЧС",
        description="План действий, приказы, состав комиссии, программа учений и журнал занятий",
        scenario=PackScenario.CIVIL_DEFENCE,
        module=DocumentPackModule.CUSTOM,
        scenario_type=DocumentPackScenario.DOCUMENT_BATCH,
        templates=(
            PackTemplateSpec(
                code="pack_gochs_plan",
                name="План действий при ЧС",
                description="Предупреждение и ликвидация чрезвычайных ситуаций",
                category="plan",
                builder=_civil_defence_plan,
            ),
            PackTemplateSpec(
                code="pack_gochs_order",
                name="Приказ об организации ГО",
                description="Ответственные и категория объекта",
                category="order",
                builder=_civil_defence_order,
            ),
            PackTemplateSpec(
                code="pack_gochs_commission",
                name="Состав комиссии",
                description="Комиссия по предупреждению и ликвидации ЧС",
                category="commission",
                builder=_civil_defence_commission,
            ),
            PackTemplateSpec(
                code="pack_gochs_drill",
                name="Программа учений",
                description="Учения и тренировки по гражданской обороне",
                category="program",
                builder=_civil_defence_drill_program,
            ),
            PackTemplateSpec(
                code="pack_gochs_journal",
                name="Журнал занятий",
                description="Учёт занятий по гражданской обороне",
                category="journal",
                builder=_civil_defence_journal,
            ),
        ),
        item_order=(
            "pack_gochs_plan",
            "pack_gochs_order",
            "pack_gochs_commission",
            "pack_gochs_drill",
            "pack_gochs_journal",
        ),
        metadata={
            "logo": build_inline_image_descriptor(DEFAULT_LOGO_BYTES)["data"],
            "stamp": build_inline_image_descriptor(DEFAULT_STAMP_BYTES)["data"],
            "discipline": "Гражданская оборона и ЧС",
        },
        disciplines=(Discipline.CIVIL_DEFENSE,),
    ),
    PackDefinition(
        code=PACK_CODE_ROAD_SAFETY,
        name="Безопасность дорожного движения",
        description="Базовый комплект БДД: ответственный, инструкция, план мероприятий, инструктаж водителей",
        scenario=PackScenario.ROAD_SAFETY,
        module=DocumentPackModule.OT,
        scenario_type=DocumentPackScenario.DOCUMENT_BATCH,
        templates=(
            PackTemplateSpec(
                code="pack_bdd_order",
                name="Приказ об ответственном за БДД",
                description="Назначение ответственного за безопасность дорожного движения",
                category="order",
                builder=_road_safety_order,
            ),
            PackTemplateSpec(
                code="pack_bdd_instruction",
                name="Инструкция по БДД",
                description="Порядок выпуска транспорта, предрейсовые осмотры и контроль",
                category="instruction",
                builder=_road_safety_instruction,
            ),
            PackTemplateSpec(
                code="pack_bdd_action_plan",
                name="План мероприятий по предупреждению ДТП",
                description="Профилактика аварийности с ответственными и сроками",
                category="plan",
                builder=_road_safety_action_plan,
            ),
            PackTemplateSpec(
                code="pack_bdd_briefing",
                name="Программа инструктажа водителей",
                description="Темы и продолжительность инструктажа по БДД",
                category="training",
                builder=_road_safety_briefing_program,
            ),
        ),
        item_order=(
            "pack_bdd_order",
            "pack_bdd_instruction",
            "pack_bdd_action_plan",
            "pack_bdd_briefing",
        ),
        metadata={
            "logo": build_inline_image_descriptor(DEFAULT_LOGO_BYTES)["data"],
            "stamp": build_inline_image_descriptor(DEFAULT_STAMP_BYTES)["data"],
            "discipline": "Безопасность дорожного движения",
        },
        disciplines=(Discipline.ROAD_SAFETY,),
    ),
    PackDefinition(
        code=PACK_CODE_BDD_REPORTS,
        name="Отчётность по БДД",
        description="Положение о системе управления БДД, отчёты об аварийности и мероприятиях, приказ о закреплении ТС",
        scenario=PackScenario.ROAD_SAFETY_REPORTING,
        module=DocumentPackModule.OT,
        scenario_type=DocumentPackScenario.DOCUMENT_BATCH,
        templates=(
            PackTemplateSpec(
                code="pack_bdd_regulation",
                name="Положение о системе управления БДД",
                description="Цели, порядок контроля и разбора происшествий",
                category="regulation",
                builder=_bdd_management_regulation,
            ),
            PackTemplateSpec(
                code="pack_bdd_accident_report",
                name="Отчёт о состоянии аварийности",
                description="Парк, водители, ДТП, пострадавшие и нарушения за период",
                category="report",
                builder=_bdd_accident_report,
            ),
            PackTemplateSpec(
                code="pack_bdd_measures_report",
                name="Отчёт о выполнении мероприятий по БДД",
                description="План и факт мероприятий по предупреждению аварийности",
                category="report",
                builder=_bdd_measures_report,
            ),
            PackTemplateSpec(
                code="pack_bdd_assignment_order",
                name="Приказ о закреплении ТС за водителями",
                description="Кто на чём ездит — по внесённому",
                category="order",
                builder=_bdd_vehicle_assignment_order,
            ),
        ),
        item_order=(
            "pack_bdd_regulation",
            "pack_bdd_accident_report",
            "pack_bdd_measures_report",
            "pack_bdd_assignment_order",
        ),
        metadata={
            "logo": build_inline_image_descriptor(DEFAULT_LOGO_BYTES)["data"],
            "stamp": build_inline_image_descriptor(DEFAULT_STAMP_BYTES)["data"],
            "discipline": "Безопасность дорожного движения",
        },
        disciplines=(Discipline.ROAD_SAFETY,),
    ),
    PackDefinition(
        code=PACK_CODE_GOCHS_REPORTS,
        name="Отчётность ГО и ЧС",
        description="Положение об объектовом звене РСЧС, инструкция о действиях при ЧС, донесение о ЧС, сведения о силах и средствах",
        scenario=PackScenario.CIVIL_DEFENCE_REPORTING,
        module=DocumentPackModule.OT,
        scenario_type=DocumentPackScenario.DOCUMENT_BATCH,
        templates=(
            PackTemplateSpec(
                code="pack_gochs_regulation",
                name="Положение об объектовом звене РСЧС",
                description="Состав, задачи и режимы функционирования звена",
                category="regulation",
                builder=_gochs_rsches_regulation,
            ),
            PackTemplateSpec(
                code="pack_gochs_instruction",
                name="Инструкция о действиях при ЧС",
                description="Оповещение, эвакуация, укрытие, действия персонала",
                category="instruction",
                builder=_gochs_emergency_instruction,
            ),
            PackTemplateSpec(
                code="pack_gochs_emergency_report",
                name="Донесение о чрезвычайной ситуации",
                description="Заготовка донесения в орган управления ГОЧС",
                category="report",
                builder=_gochs_emergency_report,
            ),
            PackTemplateSpec(
                code="pack_gochs_forces_report",
                name="Сведения о силах и средствах ГО",
                description="Формирования, личный состав, СИЗ и оповещение за период",
                category="report",
                builder=_gochs_forces_report,
            ),
        ),
        item_order=(
            "pack_gochs_regulation",
            "pack_gochs_instruction",
            "pack_gochs_emergency_report",
            "pack_gochs_forces_report",
        ),
        metadata={
            "logo": build_inline_image_descriptor(DEFAULT_LOGO_BYTES)["data"],
            "stamp": build_inline_image_descriptor(DEFAULT_STAMP_BYTES)["data"],
            "discipline": "Гражданская оборона и ЧС",
        },
        disciplines=(Discipline.CIVIL_DEFENSE,),
    ),
    PackDefinition(
        code=PACK_CODE_OPO_REPORTS,
        name="Отчётность по промышленной безопасности",
        description="Положение о производственном контроле, приказ о назначении ответственного, сведения об организации ПК и об инцидентах для Ростехнадзора",
        scenario=PackScenario.OPO_REPORTING,
        module=DocumentPackModule.OT,
        scenario_type=DocumentPackScenario.DOCUMENT_BATCH,
        templates=(
            PackTemplateSpec(
                code="pack_opo_pc_regulation",
                name="Положение о производственном контроле",
                description="Задачи ПК, порядок проверок, учёта и отчётности",
                category="regulation",
                builder=_opo_pc_regulation,
            ),
            PackTemplateSpec(
                code="pack_opo_pc_order",
                name="Приказ о назначении ответственного за ПК",
                description="Ответственный за производственный контроль и за эксплуатацию ОПО",
                category="order",
                builder=_opo_pc_order,
            ),
            PackTemplateSpec(
                code="pack_opo_pc_report",
                name="Сведения об организации производственного контроля",
                description="Годовой отчёт в Ростехнадзор за прошедший год",
                category="report",
                builder=_opo_pc_report,
            ),
            PackTemplateSpec(
                code="pack_opo_incident_report",
                name="Сведения об инцидентах на ОПО",
                description="Заготовка периодических сведений об инцидентах для Ростехнадзора",
                category="report",
                builder=_opo_incident_report,
            ),
        ),
        item_order=(
            "pack_opo_pc_regulation",
            "pack_opo_pc_order",
            "pack_opo_pc_report",
            "pack_opo_incident_report",
        ),
        metadata={
            "logo": build_inline_image_descriptor(DEFAULT_LOGO_BYTES)["data"],
            "stamp": build_inline_image_descriptor(DEFAULT_STAMP_BYTES)["data"],
            "discipline": "Промышленная безопасность",
        },
        disciplines=(Discipline.INDUSTRIAL_SAFETY,),
    ),
)

PACK_DEFINITIONS_BY_CODE: dict[str, PackDefinition] = {pack.code: pack for pack in DEFAULT_PACKS}

#: Комплекты БЕЗ дисциплины словаря — с причиной у каждого. Пустая разметка и
#: забытая разметка выглядят одинаково, а стоят разного: первое — решение,
#: второе — дыра в приёмке §58.3, которую никто не заметит.
PACKS_WITHOUT_DISCIPLINE: dict[str, str] = {
    PACK_CODE_SITE_ACCESS: (
        "общая охрана труда: допуск бригады нужен на любом объекте, "
        "а промышленная безопасность следует из признака ОПО у площадки, "
        "а не из самого комплекта"
    ),
    PACK_CODE_INCIDENT: (
        "общая охрана труда: расследование несчастного случая ведётся "
        "одинаково независимо от дисциплины"
    ),
    PACK_CODE_INSPECTION_PREP: (
        "дисциплина задаётся НАДЗОРНЫМ ОРГАНОМ конкретной проверки, а он "
        "хранится свободной строкой — приписать комплекту одну дисциплину "
        "значило бы выбрать её за специалиста"
    ),
    PACK_CODE_CEO_SHIELD: (
        "комплект руководителя: распределение ответственности не относится " "к одной дисциплине"
    ),
    PACK_CODE_CONTRACTOR: (
        "общая охрана труда: допуск подрядчика одинаков для всех дисциплин, "
        "конкретные требования приходят из его работ"
    ),
}


def packs_for(discipline: Discipline) -> tuple[PackDefinition, ...]:
    """Комплекты одной дисциплины (в объявленном порядке каталога)."""

    return tuple(pack for pack in DEFAULT_PACKS if discipline in pack.disciplines)


def discipline_pack_coverage() -> list[dict[str, object]]:
    """Покрытие ВСЕХ дисциплин ТЗ сценарными комплектами (приёмка §58.3).

    Возвращает по строке на дисциплину: сколько комплектов и какие. Требование
    «минимум один комплект на дисциплину» проверяется по этому списку, а не по
    словам в описании.
    """

    return [
        {
            "discipline": discipline.value,
            "packs": [pack.code for pack in packs_for(discipline)],
        }
        for discipline in Discipline
    ]
