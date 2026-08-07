"""Юниты чистого сборщика печатных форм СОУТ (без БД, без async)."""

from io import BytesIO

from docx import Document

from app.domains.sout.print_form import (
    SoutCardFactorRow,
    SoutCardGuaranteeRow,
    SoutCardPrintData,
    SoutSummaryPrintData,
    SoutSummaryRow,
    build_sout_card_docx,
    build_summary_sheet_docx,
    class_counts,
    harmful_factor_count,
)


def _text(docx_bytes: bytes) -> str:
    doc = Document(BytesIO(docx_bytes))
    parts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _card(**over) -> SoutCardPrintData:
    base = dict(
        org_header="ООО Ромашка",
        generated_at="2026-06-27",
        expert_org_name="ЭкспертОрг",
        report_number="СОУТ-1",
        report_date="2026-06-01",
        workplace_code="РМ-01",
        position_name="Сварщик",
        assessed_class="harmful_3_2",
        assessment_date="2026-05-20",
        next_assessment_date="2031-05-20",
        factors=[SoutCardFactorRow(code="4.50", name="Шум", measured_class="harmful_3_1")],
        guarantees=[SoutCardGuaranteeRow(kind="milk", detail="0.5 л/смена")],
    )
    base.update(over)
    return SoutCardPrintData(**base)


def test_card_contains_identity_and_class_labels() -> None:
    out = build_sout_card_docx(_card())
    txt = _text(out)
    assert "РМ-01" in txt and "Сварщик" in txt
    assert "ЭкспертОрг" in txt
    assert "3.2" in txt  # итоговый класс РМ (RU-метка)
    assert "Шум" in txt and "3.1" in txt
    assert "Молоко" in txt  # GUARANTEE_KIND_LABELS["milk"]


def test_card_empty_sections_render() -> None:
    out = build_sout_card_docx(_card(factors=[], guarantees=[], assessed_class=None))
    txt = _text(out)
    assert "РМ-01" in txt
    assert "—" in txt  # отсутствующий класс / пустые секции


def test_harmful_factor_count_counts_3_1_and_above() -> None:
    rows = [
        SoutCardFactorRow(code=None, name="a", measured_class="acceptable"),
        SoutCardFactorRow(code=None, name="b", measured_class="harmful_3_1"),
        SoutCardFactorRow(code=None, name="c", measured_class="dangerous"),
        SoutCardFactorRow(code=None, name="d", measured_class=None),
    ]
    assert harmful_factor_count(rows) == 2


def test_class_counts_canonical_order_and_tally() -> None:
    counts = class_counts(["optimal", "harmful_3_1", "harmful_3_1", None, "dangerous"])
    as_dict = dict(counts)
    assert as_dict["3.1"] == 2
    assert as_dict["4 (опасный)"] == 1
    # канонический порядок: оптимальный раньше опасного
    labels = [label for label, _ in counts]
    assert labels.index("1 (оптимальный)") < labels.index("4 (опасный)")


def test_summary_lists_workplaces_and_stats() -> None:
    data = SoutSummaryPrintData(
        org_header="ООО Ромашка",
        generated_at="2026-06-27",
        campaign_name="СОУТ 2026",
        expert_org_name="ЭкспертОрг",
        report_number="СОУТ-1",
        report_date="2026-06-01",
        rows=[
            SoutSummaryRow(
                workplace_code="РМ-01",
                position_name="Сварщик",
                assessed_class="harmful_3_2",
                harmful_factor_count=2,
                next_assessment_date="2031-05-20",
            ),
            SoutSummaryRow(
                workplace_code="РМ-02",
                position_name="Слесарь",
                assessed_class="acceptable",
                harmful_factor_count=0,
                next_assessment_date=None,
            ),
        ],
        class_counts=class_counts(["harmful_3_2", "acceptable"]),
    )
    txt = _text(build_summary_sheet_docx(data))
    assert "РМ-01" in txt and "РМ-02" in txt
    assert "Сварщик" in txt and "Слесарь" in txt
    assert "СОУТ 2026" in txt
    assert "3.2" in txt  # секция «Итоги по классам» (class_counts)


def test_summary_empty_campaign_renders() -> None:
    data = SoutSummaryPrintData(
        org_header="ООО Ромашка",
        generated_at="2026-06-27",
        campaign_name="Пустая",
        expert_org_name=None,
        report_number=None,
        report_date=None,
        rows=[],
        class_counts=[],
    )
    txt = _text(build_summary_sheet_docx(data))
    assert "Пустая" in txt
