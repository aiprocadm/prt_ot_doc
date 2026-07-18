"""Юниты чистого домена декларации СОУТ (без БД, без async)."""
from io import BytesIO

from docx import Document

from app.domains.sout.declaration import (
    DeclarationPrintData,
    DeclarationRow,
    build_declaration_docx,
    evaluate_eligibility,
)
from app.domains.sout.print_form import SoutCardFactorRow


def _text(docx_bytes: bytes) -> str:
    doc = Document(BytesIO(docx_bytes))
    parts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _f(cls: str | None) -> SoutCardFactorRow:
    return SoutCardFactorRow(code=None, name="x", measured_class=cls)


def test_optimal_class_no_factors_is_eligible() -> None:
    assert evaluate_eligibility("optimal", []) == (True, None)


def test_acceptable_class_with_acceptable_factor_is_eligible() -> None:
    assert evaluate_eligibility("acceptable", [_f("acceptable")]) == (True, None)


def test_harmful_class_is_ineligible_with_reason() -> None:
    eligible, reason = evaluate_eligibility("harmful_3_1", [])
    assert eligible is False
    assert "3.1" in reason


def test_missing_class_is_ineligible() -> None:
    eligible, reason = evaluate_eligibility(None, [])
    assert eligible is False
    assert "не проставлен" in reason


def test_class_2_but_harmful_factor_is_ineligible() -> None:
    # Ключевой анти-грабли: класс проставлен 2, но есть фактор 3.1 → дисквалификация.
    eligible, reason = evaluate_eligibility("acceptable", [_f("harmful_3_1")])
    assert eligible is False
    assert "вредный фактор" in reason


def test_build_docx_lists_eligible_rows_and_employer_placeholders() -> None:
    data = DeclarationPrintData(
        org_header="ООО Ромашка",
        generated_at="2026-06-28",
        campaign_name="СОУТ 2026",
        report_number="СОУТ-1",
        report_date="2026-06-01",
        rows=[
            DeclarationRow(
                workplace_code="РМ-02", position_name="Слесарь",
                assessed_class="acceptable", headcount="1",
                report_ref="СОУТ-1 от 2026-06-01", eligible=True, ineligible_reason=None,
            ),
        ],
    )
    txt = _text(build_declaration_docx(data))
    assert "Декларация соответствия условий труда" in txt
    assert "ООО Ромашка" in txt
    assert "ИНН" in txt and "ОГРН" in txt   # плейсхолдеры реквизитов работодателя
    assert "РМ-02" in txt and "Слесарь" in txt
    assert "СОУТ-1" in txt                   # основание (отчёт СОУТ)


def test_build_docx_empty_rows_renders_placeholder() -> None:
    data = DeclarationPrintData(
        org_header="ООО Ромашка", generated_at="2026-06-28",
        campaign_name="Пустая", report_number=None, report_date=None, rows=[],
    )
    txt = _text(build_declaration_docx(data))
    assert "Нет рабочих мест" in txt
