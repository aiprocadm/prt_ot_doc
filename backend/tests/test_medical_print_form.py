"""Чистый сборщик печатных форм 29н (контингент + поименный список), python-docx.

Без I/O / без sqlalchemy — зеркало domains/work_permits/print_form.py.
Покрываем: перевод меток (вид осмотра / статус), наличие данных в DOCX-тексте,
пустой набор строк → не падает.
"""

from __future__ import annotations

from io import BytesIO

from docx import Document

from app.domains.medical.print_form import (
    NamedListPrintData,
    NamedListPrintRow,
    ReferralPrintData,
    RegisterPrintData,
    RegisterPrintRow,
    build_contingent_register_docx,
    build_named_list_docx,
    build_referral_docx,
    contingent_status_label,
    exam_kind_label,
)


def _docx_text(b: bytes) -> str:
    doc = Document(BytesIO(b))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


class TestLabelHelpers:
    def test_exam_kind_label_known(self):
        assert exam_kind_label("periodic") == "Периодический"
        assert exam_kind_label("psychiatric") == "Психиатрическое освидетельствование"

    def test_exam_kind_label_unknown_falls_back_to_code(self):
        assert exam_kind_label("weird_kind") == "weird_kind"

    def test_contingent_status_label_known_and_fallback(self):
        assert contingent_status_label("overdue") == "Просрочен"
        assert contingent_status_label("ok") == "В норме"
        assert contingent_status_label("mystery") == "mystery"


class TestContingentRegisterDocx:
    def _data(self, **kw) -> RegisterPrintData:
        defaults = dict(
            org_header="ООО Тест",
            generated_at="2026-06-24",
            rows=[
                RegisterPrintRow(
                    position_name="Сварщик",
                    factors=[("4.4", "Шум")],
                    headcount=3,
                    exam_kinds=["periodic"],
                    periodicity_months=12,
                )
            ],
        )
        defaults.update(kw)
        return RegisterPrintData(**defaults)

    def test_heading_mentions_contingent(self):
        text = _docx_text(build_contingent_register_docx(self._data()))
        assert "контингент" in text.lower()

    def test_org_header_present(self):
        text = _docx_text(build_contingent_register_docx(self._data()))
        assert "ООО Тест" in text

    def test_row_data_present_with_labels(self):
        text = _docx_text(build_contingent_register_docx(self._data()))
        assert "Сварщик" in text
        assert "4.4" in text and "Шум" in text
        assert "3" in text  # headcount
        assert "Периодический" in text  # exam_kind label, not raw "periodic"
        assert "periodic" not in text

    def test_empty_rows_does_not_crash(self):
        text = _docx_text(build_contingent_register_docx(self._data(rows=[])))
        assert "контингент" in text.lower()  # heading still rendered


class TestNamedListDocx:
    def _data(self, **kw) -> NamedListPrintData:
        defaults = dict(
            org_header="ООО Тест",
            generated_at="2026-06-24",
            rows=[
                NamedListPrintRow(
                    full_name="Иванов Иван",
                    position_name="Сварщик",
                    department="Цех №1",
                    factors=[("4.4", "Шум")],
                    required_kinds=["periodic"],
                    last_exam_date="2025-01-01",
                    next_due_date="2026-01-01",
                    status="overdue",
                )
            ],
        )
        defaults.update(kw)
        return NamedListPrintData(**defaults)

    def test_heading_mentions_named_list(self):
        text = _docx_text(build_named_list_docx(self._data()))
        assert "поименный список" in text.lower()

    def test_person_row_with_labels(self):
        text = _docx_text(build_named_list_docx(self._data()))
        assert "Иванов Иван" in text
        assert "Цех №1" in text
        assert "Сварщик" in text
        assert "Периодический" in text  # required_kind label
        assert "Просрочен" in text  # status label, not raw "overdue"
        assert "overdue" not in text

    def test_empty_rows_does_not_crash(self):
        text = _docx_text(build_named_list_docx(self._data(rows=[])))
        assert "поименный список" in text.lower()


class TestReferralDocx:
    def _data(self, **kw) -> ReferralPrintData:
        defaults = dict(
            org_header="ООО Тест",
            generated_at="2026-06-24",
            full_name="Иванов Иван Иванович",
            birth_date="1990-05-01",
            position_name="Сварщик",
            department="Цех №1",
            exam_kind="periodic",
            medical_org_name="Клиника №5",
            due_date="2026-07-01",
            snils="123-456-789 00",
            factors=[("4.4", "Шум")],
        )
        defaults.update(kw)
        return ReferralPrintData(**defaults)

    def test_heading_mentions_referral(self):
        text = _docx_text(build_referral_docx(self._data()))
        assert "направление" in text.lower()

    def test_person_and_details_present_with_labels(self):
        text = _docx_text(build_referral_docx(self._data()))
        assert "Иванов Иван Иванович" in text
        assert "Сварщик" in text
        assert "Цех №1" in text
        assert "Клиника №5" in text
        assert "123-456-789 00" in text
        assert "4.4" in text and "Шум" in text
        assert "Периодический" in text  # exam_kind label, not raw "periodic"
        assert "periodic" not in text

    def test_missing_optionals_render_dash_and_do_not_crash(self):
        text = _docx_text(
            build_referral_docx(
                self._data(
                    birth_date=None,
                    position_name=None,
                    department=None,
                    medical_org_name=None,
                    due_date=None,
                    snils=None,
                    factors=[],
                )
            )
        )
        assert "направление" in text.lower()
        assert "Иванов Иван Иванович" in text
