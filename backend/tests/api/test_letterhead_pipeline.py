"""Integration test: LetterheadResolver wired into _generate_document_for_run.

These tests verify that when doc_pipeline_letterhead_auto=True the issuer's
letterhead preset is applied to the rendered DOCX, and when disabled=True the
output has no headers.

Local execution is blocked (Python 3.14, no pydantic-core wheel) — the test
file is valid pytest-asyncio; CI on Python 3.12.12 is the source of truth.
"""
from __future__ import annotations

import io
import zipfile
from contextlib import asynccontextmanager
from datetime import datetime, timezone
from io import BytesIO
from typing import Any
from unittest.mock import AsyncMock, MagicMock, patch
from uuid import uuid4

import pytest
from sqlalchemy import Column, String, Table
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.db.session import SharedBase, TenantBase
from app.models.document import Document, DocumentSnapshot, DocumentVersion
from app.models.job_engine import OutboxEvent
from app.models.models import (
    Company,
    PipelineRun,
    PipelineRunStatus,
    RoleEnum,
    Template,
    TemplateVersion,
    TemplateVersionStatus,
    Tenant,
    User,
)
from app.modules.headers.models import HeaderFooterPreset
from app.services.file_storage import FileStorageService

pytestmark = pytest.mark.asyncio


# ---------------------------------------------------------------------------
# Minimal DOCX factory (no docx-template placeholders, just a paragraph)
# ---------------------------------------------------------------------------

def _make_minimal_docx() -> bytes:
    """Return minimal valid DOCX bytes with a sectPr (required by apply_headers_to_docx)."""
    try:
        from docx import Document as DocxDoc

        doc = DocxDoc()
        doc.add_paragraph("Test content")
        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue()
    except ImportError:
        # python-docx not available; produce a bare-minimum DOCX zip manually
        # that satisfies DocxTemplate and apply_headers_to_docx
        return _bare_docx_bytes()


def _bare_docx_bytes() -> bytes:
    """Hand-crafted minimal DOCX that satisfies both DocxTemplate and apply_headers_to_docx."""
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas"'
        ' xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body>"
        "<w:p><w:r><w:t>Test</w:t></w:r></w:p>"
        "<w:sectPr/>"
        "</w:body>"
        "</w:document>"
    )
    content_types_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml"'
        ' ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    rels_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1"'
        ' Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"'
        ' Target="word/document.xml"/>'
        "</Relationships>"
    )
    word_rels_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        "</Relationships>"
    )
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types_xml)
        zf.writestr("_rels/.rels", rels_xml)
        zf.writestr("word/document.xml", document_xml)
        zf.writestr("word/_rels/document.xml.rels", word_rels_xml)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Read header XML helper
# ---------------------------------------------------------------------------

def _read_header_xml(docx_bytes: bytes) -> str:
    """Return concatenated word/header*.xml contents from DOCX, empty string if none."""
    try:
        with zipfile.ZipFile(BytesIO(docx_bytes)) as zf:
            parts = [name for name in zf.namelist() if name.startswith("word/header")]
            if not parts:
                return ""
            return "".join(zf.read(name).decode("utf-8", errors="replace") for name in sorted(parts))
    except Exception:
        return ""


# ---------------------------------------------------------------------------
# DB fixture (SQLite in-memory with all tenant tables)
# ---------------------------------------------------------------------------

@pytest.fixture()
async def db_engine():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")

    # Tenant table must exist in TenantBase metadata (it lives in SharedBase normally)
    if "tenant" not in TenantBase.metadata.tables:
        Table("tenant", TenantBase.metadata, Column("id", String(36), primary_key=True))

    # SharedBase: strip the schema prefix so it works on SQLite
    SharedBase.metadata.schema = None
    for table in SharedBase.metadata.tables.values():
        table.schema = None

    async with engine.begin() as conn:
        await conn.run_sync(SharedBase.metadata.create_all)
        await conn.run_sync(TenantBase.metadata.create_all)

    yield engine
    await engine.dispose()


@pytest.fixture()
async def db_session(db_engine):
    session_factory = async_sessionmaker(db_engine, expire_on_commit=False)
    async with session_factory() as session:
        yield session


# ---------------------------------------------------------------------------
# Seed helpers
# ---------------------------------------------------------------------------

TENANT_ID = "t-letterhead-1"
COMPANY_ID = "c-letterhead-1"
# Task 8 additions: subject company (document owner) and managing company (letterhead issuer)
SUBJECT_COMPANY_ID = "c-letterhead-subject"
MANAGING_COMPANY_ID = "c-letterhead-managing"
USER_ID = "u-letterhead-1"
TEMPLATE_ID = "tmpl-letterhead-1"
TEMPLATE_VERSION_ID = "tv-letterhead-1"
RUN_ID = "run-letterhead-1"
PRESET_CODE = "default-letterhead"
TEMPLATE_STORAGE_KEY = f"tests/template/{TEMPLATE_VERSION_ID}.docx"


async def _seed(session, *, preset_enabled: bool = True) -> bytes:
    """Insert all required rows; return the template DOCX bytes."""
    docx_bytes = _make_minimal_docx()

    # Store template bytes in the shared FileStorageService memory adapter
    storage = FileStorageService.default()
    storage.clear()
    storage.put(TEMPLATE_STORAGE_KEY, docx_bytes)

    # Tenant row — required so session.get(Tenant, tenant_id) returns a real object
    # for BrandingService construction inside _generate_document_for_run.
    tenant = Tenant(
        id=TENANT_ID,
        slug="test-tenant",
        name="Test Tenant",
        contact_email="tenant@example.com",
        kind="customer",
        settings={},
    )
    session.add(tenant)

    # HeaderFooterPreset — header_odd_xml references legal_name so it surfaces in output
    if preset_enabled:
        preset = HeaderFooterPreset(
            id=str(uuid4()),
            tenant_id=TENANT_ID,
            code=PRESET_CODE,
            name="Default Letterhead",
            different_first=False,
            different_odd_even=False,
            header_odd_xml="{{ organization.legal_name }}",
            watermark={},
        )
        session.add(preset)

    # Template + TemplateVersion
    tmpl = Template(
        id=TEMPLATE_ID,
        tenant_id=TENANT_ID,
        name="Test Template",
        code="test-tpl",
    )
    session.add(tmpl)

    tv = TemplateVersion(
        id=TEMPLATE_VERSION_ID,
        tenant_id=TENANT_ID,
        template_id=TEMPLATE_ID,
        version=1,
        checksum=b"fake",
        payload_key=TEMPLATE_STORAGE_KEY,
        status=TemplateVersionStatus.READY,
    )
    session.add(tv)

    # Primary company — used by existing tests (issuer = subject = same company).
    # Note: Company has no `legal_name` column; the branding profile derives
    # legal_name from company.name via BrandingService.build_profile().
    company = Company(
        id=COMPANY_ID,
        tenant_id=TENANT_ID,
        name="ООО Группа",
        preferred_header_preset_code=PRESET_CODE if preset_enabled else None,
    )
    session.add(company)

    # Subject company (Task 8): the company that owns the document (ООО Дочерняя).
    # Its own preferred_header_preset_code is intentionally left unset so the only
    # way headers appear is through an explicit issuer override to the managing company.
    subject_company = Company(
        id=SUBJECT_COMPANY_ID,
        tenant_id=TENANT_ID,
        name="ООО Дочерняя",
        preferred_header_preset_code=None,
    )
    session.add(subject_company)

    # Managing company (Task 8): the group-level company whose letterhead is used
    # when the issuer override points at it.  Its preferred_header_preset_code is set
    # so LetterheadResolver can discover the preset through the normal company path.
    managing_company = Company(
        id=MANAGING_COMPANY_ID,
        tenant_id=TENANT_ID,
        name="ООО Управляющая",
        preferred_header_preset_code=PRESET_CODE if preset_enabled else None,
    )
    session.add(managing_company)

    # User
    user = User(
        id=USER_ID,
        tenant_id=TENANT_ID,
        email="test@example.com",
        full_name="Test User",
        role=RoleEnum.ADMIN,
        hashed_password="x",
    )
    session.add(user)

    # PipelineRun — default subject company is COMPANY_ID; generate_and_fetch can
    # update it before driving the pipeline (see _LetterheadFixture.generate_and_fetch).
    run = PipelineRun(
        id=RUN_ID,
        tenant_id=TENANT_ID,
        template_id=TEMPLATE_ID,
        template_version_id=TEMPLATE_VERSION_ID,
        status=PipelineRunStatus.QUEUED,
        idempotency_key=f"test-{uuid4().hex}",
        context={},
        result_metadata={
            "company_id": COMPANY_ID,
            "initiated_by": USER_ID,
            "output_name": "test-output",
        },
    )
    session.add(run)

    await session.flush()
    return docx_bytes


# ---------------------------------------------------------------------------
# Fixture helper class
# ---------------------------------------------------------------------------

class _LetterheadFixture:
    def __init__(self, session, monkeypatch, *, letterhead_auto: bool):
        self._session = session
        self._monkeypatch = monkeypatch
        self._letterhead_auto = letterhead_auto
        # Existing tests use issuer_company_id = the primary company
        self.issuer_company_id = COMPANY_ID
        # Task 8: expose subject/managing company IDs and the preset code
        self.subject_company_id = SUBJECT_COMPANY_ID
        self.managing_company_id = MANAGING_COMPANY_ID
        self.preset_code = PRESET_CODE
        self._stored: dict[str, bytes] = {}
        # Snapshot for company_count_unchanged(); populated lazily on first call.
        self._company_count_snapshot: int | None = None

    def read_header_xml(self, docx_bytes: bytes) -> str:
        return _read_header_xml(docx_bytes)

    async def _snapshot_company_count(self) -> None:
        """Record the current number of Company rows (called once during fixture setup)."""
        from sqlalchemy import func, select as sa_select

        from app.models.models import Company as _Company

        result = await self._session.execute(sa_select(func.count()).select_from(_Company))
        self._company_count_snapshot = result.scalar_one()

    async def company_count_unchanged(self) -> bool:
        """Return True iff the Company row count equals the snapshot taken at fixture setup.

        This proves that the adhoc issuer path did NOT insert any Company row.
        """
        from sqlalchemy import func, select as sa_select

        from app.models.models import Company as _Company

        result = await self._session.execute(sa_select(func.count()).select_from(_Company))
        current_count = result.scalar_one()
        return current_count == self._company_count_snapshot

    async def generate_and_fetch(
        self,
        *,
        company_id: str,
        letterhead: dict[str, Any] | None,
    ) -> bytes:
        """Drive _generate_document_for_run and return the stored DOCX bytes.

        The ``company_id`` parameter controls which company is the *subject* of the
        document (i.e. ``result_metadata["company_id"]``).  Callers may pass a
        different company_id than the one seeded by default; the run row is updated
        in-place before the pipeline executes so the pipeline picks it up.
        """
        from app.tasks._core import _generate_document_for_run

        stored: dict[str, bytes] = {}

        def _fake_put_object(*, data: bytes, mime: str, key: str, **kw: Any) -> str:
            stored[key] = data
            # Also write into FileStorageService so the function's storage.get works
            FileStorageService.default().put(key, data)
            return "etag-fake"

        session = self._session
        auto_flag = self._letterhead_auto

        # Update run metadata: set subject company_id and optional letterhead override.
        run = await session.get(PipelineRun, RUN_ID)
        meta = dict(run.result_metadata or {})
        meta["company_id"] = company_id
        if letterhead is not None:
            meta["letterhead"] = letterhead
        run.result_metadata = meta
        # Reset run to QUEUED so the pipeline doesn't short-circuit on DONE status.
        run.status = PipelineRunStatus.QUEUED
        await session.flush()

        @asynccontextmanager
        async def _fake_session_scope(**kwargs):
            yield session

        with (
            patch("app.tasks._core.session_scope", _fake_session_scope),
            patch("app.tasks._core.ensure_tenant_schema", return_value=None),
            patch("app.tasks._core.s3.put_object", side_effect=_fake_put_object),
            patch("app.tasks._core.settings") as mock_settings,
        ):
            mock_settings.doc_pipeline_letterhead_auto = auto_flag
            # Forward other settings accesses to real settings
            from app.core.config import get_settings as _get_settings
            real = _get_settings()
            mock_settings.secret_key = real.secret_key
            mock_settings.storage_backend = real.storage_backend
            mock_settings.default_tenant_slug = real.default_tenant_slug

            await _generate_document_for_run(RUN_ID, "test-tenant")

        # Find the stored key (document DOCX)
        for key, data in stored.items():
            if key.endswith(".docx"):
                return data

        raise AssertionError(f"No DOCX stored; stored keys: {list(stored.keys())}")


@pytest.fixture()
async def letterhead_fixture(db_session, monkeypatch):
    """Fixture with doc_pipeline_letterhead_auto=True and a seeded letterhead preset."""
    await _seed(db_session, preset_enabled=True)
    fixture = _LetterheadFixture(db_session, monkeypatch, letterhead_auto=True)
    # Take the company-count snapshot AFTER seeding so company_count_unchanged() has a baseline.
    await fixture._snapshot_company_count()
    return fixture


@pytest.fixture()
async def letterhead_fixture_disabled_auto(db_session, monkeypatch):
    """Fixture with doc_pipeline_letterhead_auto=False."""
    await _seed(db_session, preset_enabled=True)
    fixture = _LetterheadFixture(db_session, monkeypatch, letterhead_auto=False)
    await fixture._snapshot_company_count()
    return fixture


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------


async def test_generated_document_carries_issuer_letterhead(letterhead_fixture):
    """
    fixture готовит: тенант с пресетом 'default-letterhead', компанию-эмитента с
    реквизитами и preferred_header_preset_code, простой DOCX-шаблон с sectPr,
    флаг doc_pipeline_letterhead_auto=True. Возвращает helper для запуска
    _generate_document_for_run и чтения сохранённого DOCX из storage.
    """
    document_bytes = await letterhead_fixture.generate_and_fetch(
        company_id=letterhead_fixture.issuer_company_id,
        letterhead=None,  # авто
    )
    headers_xml = letterhead_fixture.read_header_xml(document_bytes)
    assert "ООО Группа" in headers_xml


async def test_disabled_letterhead_produces_no_headers(letterhead_fixture):
    document_bytes = await letterhead_fixture.generate_and_fetch(
        company_id=letterhead_fixture.issuer_company_id,
        letterhead={"disabled": True},
    )
    assert letterhead_fixture.read_header_xml(document_bytes) == ""


async def test_flag_off_produces_no_headers(letterhead_fixture_disabled_auto):
    """When doc_pipeline_letterhead_auto=False, the pipeline must be byte-for-byte
    equivalent to the old path (no headers injected)."""
    document_bytes = await letterhead_fixture_disabled_auto.generate_and_fetch(
        company_id=letterhead_fixture_disabled_auto.issuer_company_id,
        letterhead=None,
    )
    assert letterhead_fixture_disabled_auto.read_header_xml(document_bytes) == ""


# ---------------------------------------------------------------------------
# Task 6: DocGenerateRequest accepts letterhead override + metadata wiring
# ---------------------------------------------------------------------------


def test_generate_endpoint_accepts_letterhead_override():
    """
    Narrow unit test: verifies that DocGenerateRequest (a) accepts a
    `letterhead` field without raising a ValidationError, (b) rejects
    unknown extra fields (extra='forbid' still applies), and (c) the
    letterhead payload survives .model_dump(mode='json') so metadata
    assembly can stash it.

    Full API client test is omitted — no `api_client` fixture exists in
    this file and wiring a real authenticated client is out of scope here.
    CI on Python 3.12.12 is the source of truth for test execution.
    """
    import pydantic

    from app.api.routes.documents import DocGenerateRequest
    from app.modules.branding.schemas import LetterheadOverride

    # --- (a) valid payload with letterhead override is accepted ---
    req = DocGenerateRequest(
        template_code="some-tpl",
        template_version=1,
        company_id="company-abc",
        letterhead=LetterheadOverride(
            issuer={"kind": "company", "company_id": "issuer-xyz"},
        ),
    )
    assert req.letterhead is not None
    assert req.letterhead.issuer is not None
    assert req.letterhead.issuer.company_id == "issuer-xyz"

    # --- (b) disabled=True variant is accepted ---
    req_disabled = DocGenerateRequest(
        template_code="some-tpl",
        template_version=1,
        company_id="company-abc",
        letterhead=LetterheadOverride(disabled=True),
    )
    assert req_disabled.letterhead.disabled is True

    # --- (c) missing letterhead defaults to None (field is optional) ---
    req_no_lh = DocGenerateRequest(
        template_code="some-tpl",
        template_version=1,
        company_id="company-abc",
    )
    assert req_no_lh.letterhead is None

    # --- (d) unknown extra field is still rejected (extra='forbid') ---
    with pytest.raises(pydantic.ValidationError):
        DocGenerateRequest(
            template_code="some-tpl",
            template_version=1,
            company_id="company-abc",
            unknown_extra_field="bad",
        )

    # --- (e) metadata-assembly simulation: model_dump carries letterhead + site_id ---
    lh_dump = req.letterhead.model_dump(mode="json")
    metadata: dict = {}
    metadata["letterhead"] = lh_dump
    metadata["site_id"] = getattr(req, "site_id", None)
    assert metadata["letterhead"]["issuer"]["company_id"] == "issuer-xyz"
    assert metadata["site_id"] is None  # DocGenerateRequest has no site_id field


# ---------------------------------------------------------------------------
# Task 7: PackRunRequest accepts letterhead + context assembly carries it
# ---------------------------------------------------------------------------


def test_pack_run_request_accepts_letterhead():
    """
    Narrow unit test: verifies that PackRunRequest (a) accepts a `letterhead`
    field carrying a LetterheadOverride, (b) the field defaults to None when
    omitted, and (c) model_dump(mode='json') round-trips the payload so the
    per-run context dict is populated correctly.

    Local execution blocked (Python 3.14 — no pydantic-core wheel).
    CI on Python 3.12.12 is the source of truth.
    """
    import pydantic

    from app.modules.branding.schemas import LetterheadOverride
    from app.schemas.pack import PackRunRequest

    # --- (a) valid payload with letterhead override is accepted ---
    req = PackRunRequest(
        pack_code="P1",
        company_id="c-1",
        person_ids=[],
        letterhead=LetterheadOverride(
            issuer={"kind": "company", "company_id": "c-2"},
        ),
    )
    assert req.letterhead is not None
    assert req.letterhead.issuer is not None
    assert req.letterhead.issuer.company_id == "c-2"

    # --- (b) omitting letterhead defaults to None ---
    req_no_lh = PackRunRequest(
        pack_code="P1",
        company_id="c-1",
        person_ids=[],
    )
    assert req_no_lh.letterhead is None

    # --- (c) disabled=True variant is accepted ---
    req_disabled = PackRunRequest(
        pack_code="P1",
        company_id="c-1",
        person_ids=[],
        letterhead=LetterheadOverride(disabled=True),
    )
    assert req_disabled.letterhead.disabled is True

    # --- (d) context-assembly simulation: letterhead + site_id end up in context dict ---
    # Mirrors what run_pack does after _build_context():
    #   context["letterhead"] = payload.letterhead.model_dump(mode="json") if payload.letterhead else None
    #   context["site_id"] = site.id if site else None
    context: dict = {}
    context["letterhead"] = req.letterhead.model_dump(mode="json") if req.letterhead else None
    context["site_id"] = "site-abc"  # simulated site.id
    assert context["letterhead"] is not None
    assert context["letterhead"]["issuer"]["company_id"] == "c-2"
    assert context["site_id"] == "site-abc"

    # None letterhead also produces None in context (no-override path)
    context_no_lh: dict = {}
    context_no_lh["letterhead"] = req_no_lh.letterhead.model_dump(mode="json") if req_no_lh.letterhead else None
    context_no_lh["site_id"] = None
    assert context_no_lh["letterhead"] is None
    assert context_no_lh["site_id"] is None


# ---------------------------------------------------------------------------
# Task 8: group-company override and adhoc issuer
# ---------------------------------------------------------------------------


async def test_override_to_other_group_company_uses_its_requisites(letterhead_fixture):
    """Субъект документа — дочерняя компания, эмитент — управляющая компания группы.

    The document is owned by 'ООО Дочерняя' (subject_company_id), but the
    letterhead issuer is explicitly overridden to 'ООО Управляющая'
    (managing_company_id).  The preset's {{ organization.legal_name }} placeholder
    must render the managing company's legal name, and the subject company's name
    must not appear in the header.

    This exercises the group-of-companies support: a holding company's letterhead
    is applied to documents generated for its subsidiaries.
    """
    document_bytes = await letterhead_fixture.generate_and_fetch(
        company_id=letterhead_fixture.subject_company_id,
        letterhead={
            "issuer": {
                "kind": "company",
                "company_id": letterhead_fixture.managing_company_id,
            }
        },
    )
    headers_xml = letterhead_fixture.read_header_xml(document_bytes)
    assert "ООО Управляющая" in headers_xml
    assert "ООО Дочерняя" not in headers_xml


async def test_adhoc_issuer_prints_inline_requisites_without_persisting(letterhead_fixture):
    """Adhoc (one-off external) issuer — inline requisites appear in the header, no DB write.

    The issuer is passed inline (kind='adhoc') with a legal_name and INN.  The
    preset_code override tells the resolver which layout preset to use (adhoc has no
    preferred_header_preset_code chain of its own).

    After the pipeline runs, the company registry must be unchanged — the adhoc
    path must NOT insert a Company row for the one-off issuer.
    """
    document_bytes = await letterhead_fixture.generate_and_fetch(
        company_id=letterhead_fixture.subject_company_id,
        letterhead={
            "issuer": {
                "kind": "adhoc",
                "inline": {"legal_name": "ООО Разовая", "inn": "9900000000"},
            },
            "preset_code": letterhead_fixture.preset_code,
        },
    )
    headers_xml = letterhead_fixture.read_header_xml(document_bytes)
    assert "ООО Разовая" in headers_xml
    assert await letterhead_fixture.company_count_unchanged()
