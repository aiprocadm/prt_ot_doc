"""Document generation API endpoints."""

from __future__ import annotations

import csv
import hashlib
import json
from io import StringIO
from typing import Annotated, Any

from fastapi import (
    APIRouter,
    Depends,
    Form,
    Header,
    HTTPException,
    Request,
    Response,
    UploadFile,
    status,
)
from uuid import UUID
from pydantic import BaseModel, ConfigDict, Field, model_validator
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.dependencies import get_session, get_tenant_record
from app.core.payload_constraints import PayloadConstraintError, enforce_mapping_constraints
from app.core.idempotency import compute_request_hash
from app.core.tracing import get_trace_id
from app.core.security import AccessContext, abac, rbac
from app.models.document import (
    DocumentBatchItem,
    DocumentBatchItemStatus,
    DocumentBatchRun,
    DocumentBatchStatus,
)
from app.models.models import (
    Company,
    Person,
    PipelineRun,
    PipelineRunStatus,
    Template,
    TemplateVersion,
    TemplateVersionStatus,
    Tenant,
    User,
)
from app.schemas.document import DocumentBatchRunRead, DocumentRead, DocumentStatusUpdate
from app.schemas.task import TaskAcceptedResponse, TaskStatusResponse
from app.services.audit import AuditService
from app.services.documents import (
    DocumentNotFoundError,
    DocumentWorkflowService,
    InvalidStatusTransitionError,
)
from app.services.idempotency import IdempotencyService, normalize_idempotency_key
from app.tasks import generate_document_task, generate_document_batch_item_task

router = APIRouter()

SessionDep = Depends(get_session)
TenantDep = Depends(get_tenant_record)


def _tenant_resource_id(tenant: Tenant = Depends(get_tenant_record)) -> UUID | None:
    return getattr(tenant, "id", None)


_DOCUMENT_RUN_ROLES = ["admin", "employee", "client_admin"]
_DOCUMENT_STATUS_ROLES = ["admin"]

AccessDep = Depends(
    abac(
        _tenant_resource_id,
        required_roles=_DOCUMENT_RUN_ROLES,
        action="manage documents",
    )
)
StatusAccessDep = Depends(rbac(_DOCUMENT_STATUS_ROLES))


class DocGenerateRequest(BaseModel):
    """Incoming payload for document generation requests."""

    model_config = ConfigDict(extra="forbid")

    template_code: str | None = Field(default=None, min_length=1, max_length=255)
    template_id: str | None = Field(default=None, min_length=1)
    template_version: int | None = Field(default=None, ge=1)
    company_id: str = Field(..., min_length=1)
    person_id: str | None = Field(default=None)
    data: dict[str, Any] = Field(default_factory=dict)

    @model_validator(mode="after")
    def _ensure_identifier(self) -> "DocGenerateRequest":
        if not self.template_code and not self.template_id:
            raise ValueError("Either template_id or template_code must be provided")
        if self.template_version is None:
            raise ValueError("template_version is required to select a template")
        return self

def _serialize_payload(payload: dict[str, Any]) -> str:
    try:
        return json.dumps(payload, sort_keys=True, separators=(",", ":"))
    except (TypeError, ValueError) as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "data must be JSON serializable") from exc


def _hash_payload(payload: dict[str, Any]) -> str:
    serialized = _serialize_payload(payload)
    digest = hashlib.sha256(serialized.encode("utf-8"))
    return digest.hexdigest()


def _parse_csv_payload(file: UploadFile) -> list[dict[str, Any]]:
    raw = file.file.read()
    content = raw.decode("utf-8")
    reader = csv.DictReader(StringIO(content))
    return [dict(row) for row in reader if any(row.values())]


def _parse_xlsx_payload(file: UploadFile) -> list[dict[str, Any]]:
    from io import BytesIO

    from openpyxl import load_workbook

    raw = file.file.read()
    workbook = load_workbook(BytesIO(raw), read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    data_rows = []
    for row in rows[1:]:
        entry = {
            header: value for header, value in zip(headers, row) if header
        }
        if any(value is not None and value != "" for value in entry.values()):
            data_rows.append(entry)
    return data_rows


def _apply_naming_pattern(pattern: str | None, row: dict[str, Any], row_index: int) -> str | None:
    if not pattern:
        return None

    class _SafeDict(dict):
        def __missing__(self, key: str) -> str:
            return ""

    return pattern.format_map(_SafeDict(row_index=row_index, **row))

def _extract_document_version_id(run: PipelineRun) -> str | None:
    metadata = run.result_metadata or {}
    outputs = run.outputs or {}
    return metadata.get("document_version_id") or outputs.get("document_version_id")


async def _fetch_template(
    session: AsyncSession,
    tenant: Tenant,
    *,
    template_code: str | None = None,
    template_id: str | None = None,
    template_version: int | None = None,
) -> tuple[Template, TemplateVersion]:
    tenant_slug = tenant.slug
    filters: list[Any] = [
        Template.tenant_id == tenant_slug,
        TemplateVersion.tenant_id == tenant_slug,
        TemplateVersion.status == TemplateVersionStatus.ACTIVE,
    ]
    if template_id:
        filters.append(Template.id == template_id)
    if template_code:
        filters.append(Template.name == template_code)
    if template_version is not None:
        filters.append(TemplateVersion.version == template_version)
    if len(filters) == 2:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "template_id or template_code must be provided",
        )
    if template_version is None:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "template_version is required for template selection",
        )

    stmt = (
        select(Template, TemplateVersion)
        .join(TemplateVersion, TemplateVersion.template_id == Template.id)
        .where(*filters)
        .limit(1)
    )
    row = (await session.execute(stmt)).first()
    if row is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Template not found")
    return row


async def _ensure_company(session: AsyncSession, tenant: Tenant, company_id: str) -> Company:
    company = await session.get(Company, company_id)
    if company is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Company not found")
    if str(company.tenant_id) != str(tenant.id):
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Tenant mismatch for company resource")
    return company


async def _ensure_person(
    session: AsyncSession, person_id: str | None, company: Company
) -> Person | None:
    if person_id is None:
        return None
    person = await session.get(Person, person_id)
    if person is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Person not found")
    if person.company_id != company.id:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "person_id does not belong to the provided company",
        )
    return person


async def _resolve_run(
    session: AsyncSession,
    *,
    idempotency_key: str,
    tenant: Tenant,
    template: Template,
    template_version: TemplateVersion,
    company: Company,
    person: Person | None,
    payload_hash: str,
    current_user: User,
    context: dict[str, Any],
) -> tuple[PipelineRun, bool]:
    stmt = select(PipelineRun).where(
        PipelineRun.tenant_id == tenant.id,
        PipelineRun.idempotency_key == idempotency_key,
    )
    existing = (await session.execute(stmt)).scalar_one_or_none()
    if existing is not None:
        metadata = existing.result_metadata or {}
        if existing.template_id != template.id:
            raise HTTPException(status.HTTP_409_CONFLICT, "Idempotency key already used")
        if existing.template_version_id != template_version.id:
            raise HTTPException(status.HTTP_409_CONFLICT, "Idempotency key already used")
        if metadata.get("company_id") != company.id:
            raise HTTPException(status.HTTP_409_CONFLICT, "Idempotency key already used")
        expected_person = person.id if person else None
        if metadata.get("person_id") != expected_person:
            raise HTTPException(status.HTTP_409_CONFLICT, "Idempotency key already used")
        if metadata.get("payload_hash") != payload_hash:
            raise HTTPException(status.HTTP_409_CONFLICT, "Idempotency key already used")
        if existing.context != context:
            raise HTTPException(status.HTTP_409_CONFLICT, "Idempotency key already used")
        if existing.status is PipelineRunStatus.ERROR:
            raise HTTPException(
                status.HTTP_409_CONFLICT,
                "Idempotency key refers to a failed generation task",
            )
        return existing, False

    run = PipelineRun(
        tenant_id=tenant.id,
        template_id=template.id,
        template_version_id=template_version.id,
        status=PipelineRunStatus.QUEUED,
        context=dict(context),
        idempotency_key=idempotency_key,
        result_metadata={
            "company_id": company.id,
            "person_id": person.id if person else None,
            "initiated_by": current_user.id,
            "payload_hash": payload_hash,
        },
    )
    session.add(run)
    await session.flush()
    return run, True


@router.post(
    "/generate",
    response_model=TaskAcceptedResponse,
    status_code=status.HTTP_202_ACCEPTED,
)
async def generate_document(
    payload: DocGenerateRequest,
    request: Request,
    response: Response,
    idempotency_key: Annotated[str | None, Header(alias="Idempotency-Key")] = None,
    session: AsyncSession = SessionDep,
    tenant: Tenant = TenantDep,
    access: AccessContext = AccessDep,
) -> TaskAcceptedResponse:
    try:
        enforce_mapping_constraints(payload.data, field="data")
    except PayloadConstraintError as exc:
        raise HTTPException(exc.status_code, str(exc)) from exc

    normalized_key = normalize_idempotency_key(idempotency_key)
    idem_state = getattr(request.state, "idempotency", {})
    request_hash = idem_state.get("fingerprint")
    if request_hash is None:
        request_hash = compute_request_hash(payload)
        idem_state = {"key": normalized_key, "fingerprint": request_hash}
        request.state.idempotency = idem_state
    idempotency = IdempotencyService(
        session=session,
        tenant_id=str(tenant.id),
        endpoint="documents.generate",
    )
    record, created_record = await idempotency.acquire(
        key=normalized_key,
        request_hash=request_hash,
        method=request.method.upper(),
        path=request.url.path,
    )

    created_run = False
    try:
        payload_hash = _hash_payload(payload.data)

        template, template_version = await _fetch_template(
            session,
            tenant,
            template_code=payload.template_code,
            template_id=payload.template_id,
            template_version=payload.template_version,
        )
        company = await _ensure_company(session, tenant, payload.company_id)
        person = await _ensure_person(session, payload.person_id, company)
        current_user: User = access.user
        if access.role in {"client_admin", "client_user"}:
            access.ensure_company_access(company.id, action="generate documents")

        run, created_run = await _resolve_run(
            session,
            idempotency_key=normalized_key,
            tenant=tenant,
            template=template,
            template_version=template_version,
            company=company,
            person=person,
            payload_hash=payload_hash,
            current_user=current_user,
            context=payload.data,
        )

        if created_run:
            audit_service = AuditService(session)
            ip = request.client.host if request.client else "unknown"
            user_agent = request.headers.get("user-agent")
            await audit_service.log_event(
                tenant_id=str(tenant.id),
                action="render_start",
                object_type="pipeline_run",
                object_id=run.id,
                user_id=current_user.id,
                ip=ip,
                user_agent=user_agent,
                details={
                    "template_id": template.id,
                    "company_id": company.id,
                    "person_id": person.id if person else None,
                },
            )

        if not created_record:
            document_version_id = _extract_document_version_id(run)
            if document_version_id:
                await idempotency.update_document_version_id(
                    key=normalized_key,
                    document_version_id=document_version_id,
                )
            await session.commit()
            return await idempotency.respond_from_store(
                record, model=TaskAcceptedResponse, response=response
            )

        status_url = f"/api/v1/documents/tasks/{run.id}"
        metadata = run.result_metadata or {}
        result = TaskAcceptedResponse(
            task_id=run.id,
            status_url=status_url,
            document_version_id=metadata.get("document_version_id"),
        )
        await idempotency.store_success(
            record,
            status_code=status.HTTP_202_ACCEPTED,
            body=result.model_dump(mode="json"),
        )
        await session.commit()
    except HTTPException as exc:
        await idempotency.store_failure(
            record,
            status_code=exc.status_code,
            detail={"detail": exc.detail},
        )
        await session.commit()
        raise
    except Exception as exc:
        await idempotency.store_failure(
            record,
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail={"detail": str(exc)},
        )
        await session.commit()
        raise

    if created_run:
        try:
            trace_id = get_trace_id()
            generate_document_task.apply_async(
                args=[run.id],
                kwargs={"tenant_slug": tenant.slug},
                task_id=run.id,
                headers={"trace_id": trace_id},
            )
        except Exception as exc:  # pragma: no cover - defensive logging
            stored = await idempotency.get(key=normalized_key)
            if stored is not None:
                await idempotency.store_failure(
                    stored,
                    status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                    detail={"detail": str(exc)},
                )
                await session.commit()
            raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR, str(exc)) from exc

    response.status_code = status.HTTP_202_ACCEPTED
    return result


@router.post(
    "/batch",
    response_model=DocumentBatchRunRead,
    status_code=status.HTTP_202_ACCEPTED,
)
async def generate_document_batch(
    file: UploadFile,
    request: Request,
    template_code: str | None = Form(default=None),
    template_id: str | None = Form(default=None),
    template_version: int | None = Form(default=None),
    company_id: str | None = Form(default=None),
    naming_pattern: str | None = Form(default=None),
    session: AsyncSession = SessionDep,
    tenant: Tenant = TenantDep,
    access: AccessContext = AccessDep,
) -> DocumentBatchRunRead:
    if template_version is None or not company_id:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "template_version and company_id required")

    filename = (file.filename or "").lower()
    if filename.endswith(".csv"):
        rows = _parse_csv_payload(file)
    elif filename.endswith(".xlsx"):
        rows = _parse_xlsx_payload(file)
    else:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Unsupported batch file type")

    if not rows:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Batch file is empty")

    template, template_version_row = await _fetch_template(
        session,
        tenant,
        template_code=template_code,
        template_id=template_id,
        template_version=template_version,
    )
    company = await _ensure_company(session, tenant, company_id)
    if access.role in {"client_admin", "client_user"}:
        access.ensure_company_access(company.id, action="generate documents")

    current_user: User = access.user
    batch = DocumentBatchRun(
        tenant_id=tenant.id,
        template_id=template.id,
        template_version_id=template_version_row.id,
        company_id=company.id,
        naming_pattern=naming_pattern,
        status=DocumentBatchStatus.RUNNING,
        total=len(rows),
        processed=0,
        succeeded=0,
        failed=0,
        created_by=current_user.id,
        started_at=datetime.now(tz=timezone.utc),
    )
    session.add(batch)
    await session.flush()

    items: list[DocumentBatchItem] = []
    for index, row in enumerate(rows, start=1):
        row_payload = {k: v for k, v in row.items() if k not in {"person_id"}}
        try:
            enforce_mapping_constraints(row_payload, field="data")
        except PayloadConstraintError as exc:
            raise HTTPException(exc.status_code, str(exc)) from exc
        person_id = row.get("person_id") if isinstance(row, dict) else None
        person = await _ensure_person(session, person_id, company)
        output_name = _apply_naming_pattern(naming_pattern, row_payload, index)
        payload_hash = _hash_payload(row_payload)
        run = PipelineRun(
            tenant_id=tenant.id,
            template_id=template.id,
            template_version_id=template_version_row.id,
            status=PipelineRunStatus.QUEUED,
            context=dict(row_payload),
            idempotency_key=f"batch-{batch.id}-{index}",
            result_metadata={
                "company_id": company.id,
                "person_id": person.id if person else None,
                "initiated_by": current_user.id,
                "payload_hash": payload_hash,
                "batch_id": batch.id,
                "row_index": index,
                "output_name": output_name,
            },
        )
        session.add(run)
        item = DocumentBatchItem(
            tenant_id=tenant.id,
            batch_id=batch.id,
            row_index=index,
            payload=row_payload,
            person_id=person.id if person else None,
            output_name=output_name,
            pipeline_run_id=run.id,
            status=DocumentBatchItemStatus.PENDING,
        )
        session.add(item)
        items.append(item)

    await session.commit()

    trace_id = get_trace_id()
    for item in items:
        generate_document_batch_item_task.apply_async(
            args=[batch.id, item.id],
            kwargs={"tenant_slug": tenant.slug},
            headers={"trace_id": trace_id},
        )

    return DocumentBatchRunRead(
        id=batch.id,
        status=batch.status.value,
        total=batch.total,
        processed=batch.processed,
        succeeded=batch.succeeded,
        failed=batch.failed,
        items=[
            {
                "id": item.id,
                "row_index": item.row_index,
                "status": item.status.value,
                "output_name": item.output_name,
            }
            for item in items
        ],
    )


@router.get(
    "/batch/{batch_id}",
    response_model=DocumentBatchRunRead,
)
async def get_document_batch(
    batch_id: str,
    session: AsyncSession = SessionDep,
    tenant: Tenant = TenantDep,
    access: AccessContext = AccessDep,
) -> DocumentBatchRunRead:
    access.ensure_tenant_access(tenant.id, action="read document batch")
    batch = await session.get(DocumentBatchRun, batch_id)
    if batch is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Batch not found")
    items = batch.items
    return DocumentBatchRunRead(
        id=batch.id,
        status=batch.status.value,
        total=batch.total,
        processed=batch.processed,
        succeeded=batch.succeeded,
        failed=batch.failed,
        items=[
            {
                "id": item.id,
                "row_index": item.row_index,
                "status": item.status.value,
                "error": item.error,
                "document_id": item.document_id,
                "document_version_id": item.document_version_id,
                "output_name": item.output_name,
            }
            for item in items
        ],
    )


@router.get("/tasks/{task_id}", response_model=TaskStatusResponse)
async def get_generation_task_status(
    task_id: str,
    session: AsyncSession = SessionDep,
    tenant: Tenant = TenantDep,
    access: AccessContext = AccessDep,
) -> TaskStatusResponse:
    stmt = select(PipelineRun).where(
        PipelineRun.id == task_id,
        PipelineRun.tenant_id == tenant.id,
    )
    run = (await session.execute(stmt)).scalar_one_or_none()
    if run is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Task not found")

    metadata = dict(run.result_metadata or {})
    outputs = dict(run.outputs or {})
    if outputs:
        metadata.setdefault("outputs", outputs)

    pipeline_status = (
        run.status.value if hasattr(run.status, "value") else str(run.status)
    )
    metadata.setdefault("pipeline_status", pipeline_status)

    document_id = metadata.get("document_id") or outputs.get("document_id")
    document_version_id = metadata.get("document_version_id") or outputs.get(
        "document_version_id"
    )

    return TaskStatusResponse(
        task_id=run.id,
        status=pipeline_status,
        document_id=document_id,
        document_version_id=document_version_id,
        error=run.error,
        metadata=metadata or None,
    )


@router.patch("/{document_id}/status", response_model=DocumentRead)
async def update_document_status(
    document_id: str,
    payload: DocumentStatusUpdate,
    request: Request,
    session: AsyncSession = SessionDep,
    tenant: Tenant = TenantDep,
    access: AccessContext = StatusAccessDep,
) -> DocumentRead:
    access.ensure_tenant_access(tenant.id, action="manage document status")
    service = DocumentWorkflowService(session=session)
    actor_id = getattr(access.user, "id", None)
    ip = request.client.host if request.client else "unknown"
    user_agent = request.headers.get("user-agent")
    try:
        document = await service.change_status(
            document_id=document_id,
            tenant_id=str(tenant.id),
            new_status=payload.to,
            actor_id=actor_id,
            ip=ip,
            user_agent=user_agent,
        )
    except DocumentNotFoundError as exc:
        raise HTTPException(status.HTTP_404_NOT_FOUND, str(exc)) from exc
    except InvalidStatusTransitionError as exc:
        await session.commit()
        raise HTTPException(status.HTTP_409_CONFLICT, str(exc)) from exc

    await session.commit()
    await session.refresh(document)
    return DocumentRead.model_validate(document)
